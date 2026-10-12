"""Regression tests for bugs found sweeping the install, command line, exports and data handling. Each test
reproduces the bug as it was found, so it is not made again."""

from __future__ import annotations

import csv
import dataclasses
import importlib.util
import io
import json
import os
import sqlite3
import sys
from email.message import EmailMessage
from pathlib import Path
from types import SimpleNamespace

import pytest
from openpyxl import load_workbook

from ap_coder import cli, envfile, jde, labels, paths
from ap_coder.config import DocumentIntelligenceSettings, Settings
from ap_coder.doctor import run_checks
from ap_coder.extraction import DocumentExtractor
from ap_coder.store import Store

from .conftest import ROOT, SAMPLE_STEM, SAMPLES
from .test_jde import AN8, ON, _files, _invoice

# --- JD Edwards: dates and field widths ------------------------------------------------------------------------


def test_jde_a_date_e1_cannot_hold_is_left_out_not_a_crash():
    """An invoice dated 1899 (a typo for 1999 / 2026) passed the checks, then the batch's ZIP could not be built
    (julian() raised): the whole batch, every invoice in it, could not be downloaded."""
    settings = jde.JdeSettings()
    for changes in ({"invoice_date": "1899-05-01"}, {"due_date": "0026-05-01"}):
        inv = _invoice(ON, **changes)
        problems = jde.validate([inv], settings, AN8)
        assert problems and "JD Edwards" in " ".join(problems[1]), changes
    good = _invoice(ON, inv_id=2)
    zipped, _, _ = jde.build_zip([good], settings, AN8, batch=1)  # the others still go out
    assert zipped
    credit = _invoice(ON, inv_id=3, due_date="0026-05-01", grand_total=-10.0)  # a credit's due date isn't sent
    assert all("due date" not in p for p in jde.validate([credit], settings, AN8).get(3, []))


def test_jde_text_made_formula_safe_still_fits_its_column():
    """A description starting with "-" (a discount line) gets an apostrophe in front so a spreadsheet does not run
    it: a 30-character VNEXA became 31, and a 25-character VLVINV 26, longer than the Z-table columns."""
    inv = _invoice(ON, invoice_number="=" + "1" * 24)
    inv["final_output"]["gl_distribution"][0]["description"] = "-" + "x" * 40
    files = _files([inv])
    line = files[jde.DIST_FILE][0]
    assert line["VNEXA"].startswith("'-") and len(line["VNEXA"]) == 30
    assert len(line["VNEXR"]) <= 30
    header = files[jde.HEADER_FILE][0]
    assert header["VLVINV"].startswith("'=") and len(header["VLVINV"]) == 25


# --- Labels workbook (python -m ap_coder labels) ---------------------------------------------------------------


def _prediction(tmp_path, **changes):
    out = tmp_path / "output"
    out.mkdir()
    doc = {
        "vendor_name": "Fournitures Laval", "invoice_number": "A-1", "invoice_date": "2026-01-05", "currency": "CAD",
        "subtotal": 10, "tax_total": 0, "grand_total": 10,
        "line_items": [{"line_number": 1, "description": "Paper", "quantity": 1, "unit_price": 10, "amount": 10}],
        **changes,
    }  # fmt: skip
    (out / "inv.json").write_text(json.dumps(doc), encoding="utf-8")
    return out


def test_labels_workbook_with_a_damaged_pdfs_control_characters(tmp_path):
    """Text read from a damaged PDF can hold control characters: openpyxl refused them and the labels command
    stopped with a traceback (the export workbook already dropped them)."""
    out = _prediction(tmp_path, invoice_number="A\x0c1")
    out_doc = json.loads((out / "inv.json").read_text())
    out_doc["line_items"][0]["description"] = "Pa\x01per"
    (out / "inv.json").write_text(json.dumps(out_doc))
    assert labels.export_labels(out, tmp_path / "labels.xlsx") == 1
    wb = load_workbook(tmp_path / "labels.xlsx")
    values = [c.value for row in wb["Labels"].iter_rows(min_row=2) for c in row]
    assert "A1" in values and "Paper" in values


def test_labels_csv_is_formula_safe_and_reads_back_the_same(tmp_path):
    """The CSV labels file is opened in Excel by the AP team: a vendor name from a document such as
    =HYPERLINK(...) ran as a formula there. Neutralised with an apostrophe, and read back without it."""
    evil = '=HYPERLINK("http://evil.example","x")'
    out = _prediction(tmp_path, vendor_name=evil)
    dest = tmp_path / "labels.csv"
    labels.export_labels(out, dest)
    text = dest.read_text(encoding="utf-8-sig")
    assert '"=HYPERLINK' not in text and "\"'=HYPERLINK" in text
    rows = list(csv.DictReader(io.StringIO(text)))
    assert rows[0]["amount"] == "10"  # numbers untouched
    for row in rows:
        row["reviewed"] = "Y"
    with dest.open("w", newline="", encoding="utf-8-sig") as fh:
        writer = csv.DictWriter(fh, fieldnames=list(rows[0]))
        writer.writeheader()
        writer.writerows(rows)
    docs, unreviewed = labels.load_labels(dest)
    assert not unreviewed and docs["inv"]["vendor_name"] == evil


# --- The .env settings file -------------------------------------------------------------------------------------


def test_an_env_file_saved_as_ansi_by_notepad_is_read(tmp_path, monkeypatch):
    """Older Notepad saves "ANSI" (Windows-1252): one accented letter in a comment (Paramètres) made every AP Coder
    command, the dashboard and the installer stop with UnicodeDecodeError."""
    env = tmp_path / ".env"
    env.write_bytes("# Paramètres de Zoé\nAP_REVIEWER=Zoé Tremblay\nAP_LLM_MODEL=qwen\n".encode("cp1252"))
    monkeypatch.delenv("AP_REVIEWER", raising=False)
    monkeypatch.setenv("AP_LLM_MODEL", "")  # restored after the test (load_dotenv writes os.environ)
    monkeypatch.delenv("AP_LLM_MODEL")
    assert envfile.read_env(env) == {"AP_REVIEWER": "Zoé Tremblay", "AP_LLM_MODEL": "qwen"}
    Settings.from_env(env)
    assert os.environ["AP_REVIEWER"] == "Zoé Tremblay"
    envfile.write_env(env, {"AP_LLM_MODEL": "other"})
    assert envfile.read_env(env)["AP_REVIEWER"] == "Zoé Tremblay"  # the rest kept (now saved as UTF-8)


def test_no_env_file_is_taken_from_a_folder_above_ap_coder(tmp_path, monkeypatch):
    """With no .env in the data folder, the current folder or AP Coder's folder, python-dotenv went looking in every
    folder above AP Coder's code and loaded the first .env it found there (another program's settings)."""
    import dotenv.main

    stray = tmp_path / "Downloads" / ".env"
    stray.parent.mkdir()
    stray.write_text("AP_STRAY_SETTING=from-another-program\n", encoding="utf-8")
    monkeypatch.setattr(dotenv.main, "find_dotenv", lambda *a, **k: str(stray))
    monkeypatch.setattr(paths, "env_file", lambda: None)
    monkeypatch.setenv("AP_STRAY_SETTING", "")
    monkeypatch.delenv("AP_STRAY_SETTING")
    Settings.from_env()
    assert "AP_STRAY_SETTING" not in os.environ


def test_folders_set_in_the_environment_with_quotes_or_a_tilde(tmp_path, monkeypatch):
    """`set AP_REFERENCE_DIR="C:\\AP Reference"` in cmd keeps the quotes in the value: AP_PRIVATE_DIR and
    AP_ENV_FILE dropped them, AP_REFERENCE_DIR and AP_DB_PATH did not (the GL accounts "not found", a database
    made in a folder named with quotes). A ~ (home folder) was not expanded: the data went to a folder named "~"
    inside the folder AP Coder was started from."""
    reference = tmp_path / "AP Reference"
    reference.mkdir()
    monkeypatch.setenv("AP_REFERENCE_DIR", f'"{reference}"')
    assert cli.reference_dir() == (reference, False)
    monkeypatch.setenv("AP_DB_PATH", f'"{tmp_path / "My Data" / "ap.db"}"')
    assert Path(cli.build_parser().parse_args(["schema"]).db) == tmp_path / "My Data" / "ap.db"
    monkeypatch.setenv("HOME", str(tmp_path / "home"))
    monkeypatch.setenv("USERPROFILE", str(tmp_path / "home"))
    monkeypatch.setenv("AP_PRIVATE_DIR", "~/APCoder")
    assert paths.private_dir() == tmp_path / "home" / "APCoder"
    monkeypatch.delenv("AP_PRIVATE_DIR")
    paths.write_user_settings(data_dir="~/AP Data")  # edited by hand
    assert paths.private_dir() == tmp_path / "home" / "AP Data"


# --- Command line ------------------------------------------------------------------------------------------------


def _cli(tmp_path, *args):
    return cli.main(["--env-file", str(tmp_path / "none.env"), "--db", str(tmp_path / "a.db"), *args])


def _offline(monkeypatch):
    for key in [k for k in os.environ if k.startswith(("AZURE_", "AP_LLM_"))]:
        monkeypatch.delenv(key)


def test_process_codes_each_invoice_once(tmp_path, monkeypatch):
    """`process folder folder/a.md` (the folder, and a file in it again), or the same invoice saved twice under two
    names: each copy was coded and added to the review queue, so AP saw the invoice two or three times."""
    _offline(monkeypatch)
    folder = tmp_path / "in put"
    folder.mkdir()
    text = (SAMPLES / f"{SAMPLE_STEM}.md").read_text(encoding="utf-8")
    (folder / "a.md").write_text(text, encoding="utf-8")
    (folder / "copy of a.md").write_text(text, encoding="utf-8")
    assert _cli(tmp_path, "process", str(folder), str(folder / "a.md"), "-o", str(tmp_path / "out"),
                "--cache-dir", "") == 0  # fmt: skip
    assert len(Store(tmp_path / "a.db").list_invoices()) == 1


def _email(path, attachment: Path) -> None:
    msg = EmailMessage()
    msg["Subject"], msg["From"] = "Invoice", "billing@supplier.example"
    msg.set_content("Please find our invoice attached.")
    msg.add_attachment(attachment.read_bytes(), maintype="application", subtype="pdf", filename="facture.pdf")
    path.write_bytes(bytes(msg))


def test_process_takes_the_invoice_out_of_a_saved_email(tmp_path, monkeypatch, capsys):
    """The dashboard and `watch` take the invoices out of saved emails (.eml); `process mail.eml` tried to read the
    email itself as a PDF: a traceback, and a failed "mail.eml" invoice in the review queue."""
    _offline(monkeypatch)
    folder = tmp_path / "mail"
    folder.mkdir()
    _email(folder / "supplier mail.eml", SAMPLES / f"{SAMPLE_STEM}.pdf")
    assert _cli(tmp_path, "process", str(folder / "supplier mail.eml"), "-o", str(tmp_path / "out"),
                "--cache-dir", "") == 0  # fmt: skip
    invoices = Store(tmp_path / "a.db").list_invoices()
    assert [(i["file_name"], i["status"]) for i in invoices] == [("supplier mail - facture.pdf", "review")]
    assert (folder / "emails" / "supplier mail.eml").exists()
    assert "Traceback" not in capsys.readouterr().err


def test_command_line_output_on_a_windows_1252_console(tmp_path, monkeypatch):
    """A Windows console or a log file (Task Scheduler) in Windows-1252: printing a Japanese file name or vendor
    stopped the command with UnicodeEncodeError, after its work was done (exit code 1, a traceback)."""
    predictions = tmp_path / "output"
    predictions.mkdir()
    (predictions / "inv.json").write_text(json.dumps({"vendor_name": "東京商事", "line_items": []}), encoding="utf-8")
    console = io.TextIOWrapper(io.BytesIO(), encoding="cp1252")
    monkeypatch.setattr(sys, "stdout", console)
    assert _cli(tmp_path, "labels", "--predictions", str(predictions), "-o", str(tmp_path / "ラベル.csv")) == 0
    console.flush()
    assert b"Wrote 1 invoice(s)" in console.buffer.getvalue()


def test_a_file_open_in_excel_is_one_plain_line(tmp_path, monkeypatch, capsys):
    """Windows locks a workbook open in Excel: `labels --force` (or share-report, export-training) writing over it
    stopped with a PermissionError traceback."""
    from openpyxl import Workbook

    predictions = tmp_path / "output"
    predictions.mkdir()
    (predictions / "inv.json").write_text(json.dumps({"vendor_name": "Laval", "line_items": []}), encoding="utf-8")
    dest = tmp_path / "labels.xlsx"
    dest.write_bytes(b"open in Excel")

    def locked(self, filename):
        raise PermissionError(13, "Permission denied", str(filename))

    monkeypatch.setattr(Workbook, "save", locked)
    assert _cli(tmp_path, "labels", "--predictions", str(predictions), "-o", str(dest), "--force") == 1
    err = capsys.readouterr().err
    assert "labels.xlsx" in err and "open in another program" in err


def test_export_training_says_why_bulk_approvals_are_left_out(tmp_path, capsys, ground_truth):
    """Every approval a bulk approval: the command said "No approved invoices to export yet", and with some
    checked ones it listed what was left out without the bulk approvals."""
    store = Store(tmp_path / "a.db")
    pdf = SAMPLES / f"{SAMPLE_STEM}.pdf"
    for number in ("A-1", "A-2"):
        invoice_id = store.add_invoice(pdf, {**ground_truth, "invoice_number": number}, {"requires_review": False})
        store.approve_invoice(invoice_id, {**ground_truth, "invoice_number": number}, "Ann", bulk=True)
    assert _cli(tmp_path, "export-training", "-o", str(tmp_path / "t.zip")) == 1
    err = capsys.readouterr().err
    assert "No approved invoices" not in err and "left out: 2 (bulk-approved" in err
    invoice_id = store.add_invoice(pdf, ground_truth, {"requires_review": False})
    store.approve_invoice(invoice_id, ground_truth, "Ann")
    assert _cli(tmp_path, "export-training", "-o", str(tmp_path / "t.zip")) == 0
    assert "left out: 2 (bulk-approved" in capsys.readouterr().err


# --- Nothing reaches the internet unless AP_ALLOW_INTERNET=1 ----------------------------------------------------


def _azure_di(monkeypatch) -> list[str]:
    """Stand-ins for Azure Document Intelligence's packages and clients: each client made is recorded (a call to
    Azure). Stand-ins whether or not the packages are installed: they are an optional extra (``.[azure]``)."""
    from types import ModuleType

    made: list[str] = []
    modules = {name: ModuleType(name) for name in ("azure", "azure.core", "azure.core.credentials", "azure.identity",
                                                    "azure.ai", "azure.ai.documentintelligence")}  # fmt: skip
    modules["azure.core.credentials"].AzureKeyCredential = lambda key: ("key", key)
    modules["azure.identity"].DefaultAzureCredential = lambda: "entra"
    di = modules["azure.ai.documentintelligence"]
    di.DocumentIntelligenceClient = lambda *a, **k: made.append("client")
    di.DocumentIntelligenceAdministrationClient = lambda *a, **k: made.append("admin")
    for name, module in modules.items():
        monkeypatch.setitem(sys.modules, name, module)
    return made


def _with_di_endpoint() -> Settings:
    di = DocumentIntelligenceSettings(endpoint="https://di-prod.cognitiveservices.azure.com/", api_key="k" * 32)
    return dataclasses.replace(Settings(), document_intelligence=di)


def test_doctor_online_never_calls_azure_in_the_offline_build(monkeypatch, reference):
    """An Azure endpoint left in the .env (an older install): `doctor --online` said "set, but not used" and then
    called Azure Document Intelligence anyway."""
    made = _azure_di(monkeypatch)
    checks = {c.area: c for c in run_checks(_with_di_endpoint(), lambda: reference, online=True)}
    assert made == [] and checks["DI connectivity"].status == "SKIP"


def test_extract_never_sends_an_invoice_to_azure_in_the_offline_build(monkeypatch):
    """`python -m ap_coder extract` (and anything else building the extractor's client) sent the invoice to Azure
    Document Intelligence whenever its endpoint was set, internet allowed or not."""
    made = _azure_di(monkeypatch)
    extractor = DocumentExtractor(_with_di_endpoint().document_intelligence)
    with pytest.raises(RuntimeError, match="AP_ALLOW_INTERNET"):
        extractor.extract(SAMPLES / f"{SAMPLE_STEM}.pdf")
    assert made == []
    monkeypatch.setenv("AP_ALLOW_INTERNET", "1")  # a developer opted in: then it may
    assert extractor.client is None and made == ["client"]


# --- Backups ------------------------------------------------------------------------------------------------------


def test_a_backup_is_one_self_contained_file(tmp_path):
    """Backups were copies in WAL mode: opening one needs -wal / -shm files beside it, which a read-only folder
    (a network share, a USB stick) cannot hold, so such a backup could not be restored ("damaged")."""
    store = Store(tmp_path / "data" / "ap_coder.db")
    store.set_setting("x", "1")
    made = store.backup_now("manual")
    assert made.read_bytes()[18:20] == b"\x01\x01"  # file format 1 = rollback journal, not WAL (2)


def test_checking_and_restoring_a_backup_writes_nothing_next_to_it(tmp_path):
    """Checking a backup (and restoring it) opened it in place: SQLite wrote -wal / -shm files next to an older,
    WAL-format backup; in a read-only folder that failed and the backup was called damaged."""
    store = Store(tmp_path / "data" / "ap_coder.db")
    store.set_setting("x", "old")
    share = tmp_path / "share"
    share.mkdir()
    old = share / "ap_coder-20260101-120000.db"
    with sqlite3.connect(old) as conn:  # a backup made by an older version: WAL format
        conn.execute("PRAGMA journal_mode=WAL")
    conn.close()
    store.backup_to(tmp_path / "plain.db")
    with sqlite3.connect(old) as target, sqlite3.connect(tmp_path / "plain.db") as source:
        source.backup(target)
        target.execute("PRAGMA journal_mode=WAL")
    target.close()
    source.close()
    assert old.read_bytes()[18:20] == b"\x02\x02"
    before = sorted(p.name for p in share.iterdir())
    Store.check_backup(old)
    assert sorted(p.name for p in share.iterdir()) == before
    store.set_setting("x", "new")
    store.restore_from(old)
    assert store.get_setting("x") == "old"
    assert sorted(p.name for p in share.iterdir()) == before


# --- Launcher: what the offline bundle has for this computer ---------------------------------------------------


@pytest.fixture
def launch(tmp_path, monkeypatch):
    spec = importlib.util.spec_from_file_location("ap_script_launch_sweep", ROOT / "scripts" / "launch.py")
    module = importlib.util.module_from_spec(spec)
    sys.modules[spec.name] = module
    spec.loader.exec_module(module)
    monkeypatch.setattr(module, "ROOT", tmp_path)
    monkeypatch.setattr(module, "VENV", tmp_path / ".venv")
    monkeypatch.setattr(module, "VENV_PY", tmp_path / ".venv" / "bin" / "python")
    return module


def _bundle(launch, *targets: str) -> None:
    (launch.ROOT / "bundle_manifest.json").write_text(json.dumps({"targets": list(targets)}), encoding="utf-8")


def test_a_bundle_without_packages_for_this_kind_of_computer_says_so(launch, monkeypatch):
    """A Windows-only bundle (built on Linux) unzipped on a Mac: the launcher said it was built for Python 3.11 to
    3.13 and to "install Python 3.12", which cannot help: it has no Mac packages at all."""
    _bundle(launch, "this computer (linux-x86_64, Python 3.12)", "win_amd64, Python 3.11", "win_amd64, Python 3.12",
            "win_amd64, Python 3.13")  # fmt: skip
    monkeypatch.setattr(launch, "wheel_platform", lambda: "macosx_15_0_arm64")
    assert launch.bundle_pythons() == []
    hint = launch.install_hint(["--no-index"], ["fastapi>=0.115,<1"])
    assert "no packages for this kind of computer (macosx_15_0_arm64)" in hint
    assert "--platform macosx_15_0_arm64" in hint and "install Python 3.12" not in hint
    monkeypatch.setattr(launch, "wheel_platform", lambda: "win_amd64")  # a Windows laptop: those are its versions
    assert launch.bundle_pythons() == ["3.11", "3.12", "3.13"]
    monkeypatch.setattr(launch, "wheel_platform", lambda: "manylinux2014_x86_64")  # the builder's own kind
    assert launch.bundle_pythons() == ["3.12"]
    _bundle(launch, "macosx_14_0_arm64, Python 3.12", "macosx_13_0_x86_64, Python 3.13")
    monkeypatch.setattr(launch, "wheel_platform", lambda: "macosx_15_0_arm64")
    assert launch.bundle_pythons() == ["3.12"]


def test_installing_the_advised_python_is_not_undone_by_the_existing_venv(launch, monkeypatch):
    """Inside a .venv made with Python 3.14 and a bundle for 3.12, the advice was "install Python 3.12": done, the
    double-click still started the working 3.14 .venv, with the same failure. The .venv must go too."""
    _bundle(launch, "win_amd64, Python 3.12")
    monkeypatch.setattr(launch, "wheel_platform", lambda: "win_amd64")
    monkeypatch.setattr(launch, "sys", SimpleNamespace(version_info=(3, 14, 0)))
    hint = launch.install_hint(["--no-index"], ["fastapi>=0.115,<1"])
    assert "install Python 3.12 on this computer and delete the .venv folder" in hint
