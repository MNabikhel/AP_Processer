"""Spend analysis and the full data workbook."""

import datetime as dt
import io

from openpyxl import load_workbook

from ap_coder import spend
from ap_coder.demo import load_demo
from ap_coder.store import Store

YEAR = (dt.date(2020, 1, 1), dt.date(2030, 12, 31))


def test_spend_from_the_demo(tmp_path):
    store = Store(tmp_path / "s.db")
    load_demo(store)
    rows = spend.lines(store, *YEAR)
    assert rows and {r.currency for r in rows} == {"CAD"}  # only approved invoices: the demo's two are in CAD
    approved = [i for i in store.list_invoices("approved")]
    assert {r.invoice_id for r in rows} == {i["id"] for i in approved}
    # Spend is net of recoverable tax: the expense postings, not the invoice totals.
    expected = sum(
        e["amount"]
        for i in approved
        for e in store.get_invoice(i["id"])["final_output"]["gl_distribution"]
        if e["kind"] == "expense"
    )
    assert round(sum(r.amount for r in rows), 2) == round(expected, 2) < sum(i["grand_total"] for i in approved)
    by_gl = spend.total_by(rows, "gl_code", "CAD")
    assert by_gl == sorted(by_gl, key=lambda x: -x[1]) and sum(n for _, _, n in by_gl) >= len(approved)
    chart = spend.by_month_and_category(rows, "CAD", {})
    assert {c["category"] for c in chart} == {"Uncategorised"}
    assert spend.lines(store, dt.date(2030, 1, 1), dt.date(2030, 12, 31)) == []


def test_currencies_are_kept_apart():
    rows = [
        spend.SpendLine(1, "2026-09", "A", "6000", "", 100.0, "CAD"),
        spend.SpendLine(2, "2026-09", "B", "6000", "", 500.0, "USD"),
        spend.SpendLine(3, "2026-09", "A", "6000", "", -20.0, "CAD"),  # a credit note
    ]
    assert spend.currencies(rows) == ["USD", "CAD"]
    assert spend.total_by(rows, "vendor", "CAD") == [("A", 80.0, 2)]


def test_default_period_is_the_last_twelve_months():
    assert spend.default_period(dt.date(2026, 10, 7)) == (dt.date(2025, 11, 1), dt.date(2026, 10, 7))
    assert spend.default_period(dt.date(2026, 12, 31)) == (dt.date(2026, 1, 1), dt.date(2026, 12, 31))


def test_workbook(tmp_path):
    store = Store(tmp_path / "s.db")
    load_demo(store)
    first = store.list_invoices()[0]["id"]
    wb = load_workbook(io.BytesIO(spend.workbook(store, store.reference_data())))
    assert wb.sheetnames == ["Invoices", "Lines", "Posting"]
    invoices = list(wb["Invoices"].values)
    assert invoices[0][0] == "AP Coder #" and len(invoices) == len(store.list_invoices()) + 1
    assert first in [r[0] for r in invoices[1:]]
    lines = list(wb["Lines"].values)
    assert lines[0][10] == "GL name" and any(r[10] for r in lines[1:])  # GL names filled in from the setup
    assert wb["Posting"].tables and wb["Lines"].tables


def test_workbook_text_cannot_run_as_a_formula(tmp_path):
    store = Store(tmp_path / "s.db")
    load_demo(store)
    inv = store.list_invoices()[0]
    doc = store.get_invoice(inv["id"])["ai_output"]
    store.add_invoice(tmp_path / "x.pdf", {**doc, "vendor_name": '=HYPERLINK("http://x")'}, {})
    wb = load_workbook(io.BytesIO(spend.workbook(store, None)))
    cell = next(c for row in wb["Invoices"].iter_rows() for c in row if str(c.value).startswith("=HYPER"))
    assert cell.data_type == "s" and cell.quotePrefix


def test_non_recoverable_tax_on_its_own_account_is_spend(tmp_path):
    import json

    from ap_coder.config import Settings
    from ap_coder.pipeline import finalise_coding
    from ap_coder.schema import InvoiceCoding
    from ap_coder.store import load_sample_setup
    from ap_coder.tax import EXPENSE_SEPARATE

    from .conftest import SAMPLES

    store = Store(tmp_path / "s.db")
    load_sample_setup(store)
    store.set_tax_treatment("PST", EXPENSE_SEPARATE, "5900")
    doc = json.loads((SAMPLES / "ground_truth" / "pacific_BC_GST_PST_PO-77120.json").read_text())
    output, _ = finalise_coding(InvoiceCoding.model_validate(doc), store.reference_data(), Settings(), store=store)
    store.approve_invoice(store.add_invoice(tmp_path / "p.pdf", output, {}), output, "Jane")
    pst = sum(t["tax_amount"] for t in doc["tax_lines"] if t["tax_type"] == "PST")
    rows = spend.lines(store, *YEAR)
    assert round(sum(r.amount for r in rows), 2) == round(doc["subtotal"] + pst, 2)
    assert any(r.gl_code == "5900" for r in rows)


def test_control_characters_do_not_break_the_workbooks(tmp_path):
    from ap_coder import exports

    store = Store(tmp_path / "s.db")
    load_demo(store)
    inv = store.list_invoices()[0]
    doc = store.get_invoice(inv["id"])["ai_output"]
    doc = {**doc, "vendor_name": "Bad\x0bVendor\x01", "line_items": [{**doc["line_items"][0], "description": "a\x0cb"}]}
    invoice_id = store.add_invoice(tmp_path / "x.pdf", doc, {})
    store.approve_invoice(invoice_id, doc, "Jane")
    wb = load_workbook(io.BytesIO(spend.workbook(store, None)))
    assert "BadVendor" in [r[2] for r in wb["Invoices"].values]
    assert exports.build_xlsx([store.get_invoice(invoice_id)])[:2] == b"PK"


def test_exchange_rates_and_all_in_cad(tmp_path):
    from ap_coder.store import parse_fx_rates

    assert parse_fx_rates("USD=1.37, eur 1,50; GBP: 1.85, XXX=0") == {"CAD": 1.0, "USD": 1.37, "EUR": 1.5, "GBP": 1.85}
    assert parse_fx_rates("") == {"CAD": 1.0}
    rows = [spend.SpendLine(1, "2026-09", "A", "6000", "", 100.0, "USD"),
            spend.SpendLine(2, "2026-09", "B", "6000", "", 50.0, "CAD"),
            spend.SpendLine(3, "2026-09", "C", "6000", "", 10.0, "JPY")]  # fmt: skip
    converted = spend.in_cad(rows, {"CAD": 1.0, "USD": 1.37})
    assert [(r.amount, r.currency) for r in converted] == [(137.0, "CAD"), (50.0, "CAD")]  # no JPY rate: left out
    store = Store(tmp_path / "s.db")
    store.set_setting("fx_rates", "USD=1.37")
    assert store.fx_rates()["USD"] == 1.37
