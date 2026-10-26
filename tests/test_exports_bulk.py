"""ERP export batches (each invoice once, re-downloadable, undoable) and bulk approval of clean invoices."""

import csv
import io
import json

import pytest
from openpyxl import load_workbook

from ap_coder.coding.bulk import bulk_approve, clean_candidates
from ap_coder.core.config import Settings
from ap_coder.core.memory import vendor_key
from ap_coder.core.store import APPROVED, REVIEW, Store
from ap_coder.erp import exports

from .conftest import SAMPLE_STEM, SAMPLES, copy_database


@pytest.fixture
def demo_store(tmp_path, demo_template):
    return copy_database(demo_template, tmp_path / "ap.db")


def test_export_batch_lifecycle(demo_store):
    store = demo_store
    ready = store.unexported_approved()
    assert len(ready) == 2  # the two demo invoices approved "last week"
    batch = store.create_export_batch([i["id"] for i in ready], "xlsx", actor="jane")
    assert store.unexported_approved() == []
    assert sorted(store.batch_invoice_ids(batch)) == sorted(i["id"] for i in ready)
    with pytest.raises(ValueError):
        store.create_export_batch([i["id"] for i in ready], "csv")  # never exported twice
    assert len(store.events(actions=["exported"])) == 2
    assert store.undo_export_batch(batch, actor="jane") == 2
    assert len(store.unexported_approved()) == 2
    assert store.export_batches()[0]["undone_by"] == "jane"


def test_a_batch_of_more_invoices_than_sqlite_takes_variables(demo_store, monkeypatch, few_sql_variables):
    """Ticking every invoice: more ids than an older SQLite build takes as bound variables (999)."""
    ready = [i["id"] for i in demo_store.unexported_approved()]
    batch = demo_store.create_export_batch([*ready, *range(10_000, 11_500)], "csv")
    assert sorted(demo_store.batch_invoice_ids(batch)) == sorted(ready)


def test_export_files_balance_to_the_invoices(demo_store):
    store = demo_store
    invoices = [store.get_invoice(i["id"]) for i in store.unexported_approved()]
    data, name, _ = exports.build("xlsx", invoices, {"6010": "IT Hardware"}, batch=7)
    assert name == "ap_coder_export_7.xlsx"
    wb = load_workbook(io.BytesIO(data))
    assert wb.sheetnames == ["Invoices", "GL lines", "By GL account"]
    header = [c.value for c in wb["GL lines"][1]]
    amount_col = header.index("Amount")
    lines_total = round(sum(r[amount_col] for r in wb["GL lines"].iter_rows(min_row=2, values_only=True)), 2)
    assert lines_total == round(sum(i["final_output"]["grand_total"] for i in invoices), 2)
    csv_data, csv_name, _ = exports.build("csv", invoices, batch=7)
    rows = list(csv.DictReader(io.StringIO(csv_data.decode("utf-8-sig"))))
    assert csv_name.endswith(".csv") and len(rows) == wb["GL lines"].max_row - 1
    assert {r["Batch"] for r in rows} == {"7"}


def test_bulk_approve_only_takes_invoices_that_are_still_clean(demo_store):
    store = demo_store
    candidates = clean_candidates(store)
    assert candidates and all(not c["requires_review"] for c in candidates)
    # One of them gets a new problem after processing: its vendor goes on hold.
    victim = candidates[0]
    store.save_vendor(vendor_key(victim["vendor_name"]), victim["vendor_name"], "on_hold")
    result = bulk_approve(store, store.reference_data(), Settings(), [c["id"] for c in candidates], "jane")
    assert result["skipped"] and result["skipped"][0][0] == victim["id"] and "on hold" in result["skipped"][0][1]
    assert len(result["approved"]) == len(candidates) - 1
    for invoice_id in result["approved"]:
        assert store.get_invoice(invoice_id)["status"] == APPROVED
    assert store.get_invoice(victim["id"])["status"] == REVIEW
    approved_event = store.events(result["approved"][0], actions=["approved"])[0]
    assert approved_event["detail"]["bulk"] is True and approved_event["actor"] == "jane"


def test_bulk_approve_skips_an_invoice_a_colleague_parks_meanwhile(demo_store, monkeypatch):
    """One invoice parked (or approved) by someone else while the loop checks the others is skipped with the
    reason; the rest are still approved and the result is not lost."""
    from ap_coder.coding import bulk

    store = demo_store
    ids = [c["id"] for c in clean_candidates(store)]
    assert len(ids) >= 2
    real = bulk.finalise_coding

    def finalise(coding, *args, **kwargs):
        out = real(coding, *args, **kwargs)
        if kwargs.get("exclude_invoice_id") == ids[1]:  # parked between its check and its approval
            store.park_invoice(ids[1], "bob", "waiting on buyer")
        return out

    monkeypatch.setattr(bulk, "finalise_coding", finalise)
    result = bulk_approve(store, store.reference_data(), Settings(), ids, "jane")
    assert [i for i, _ in result["skipped"]] == [ids[1]] and "no longer in the queue" in result["skipped"][0][1]
    assert sorted(result["approved"]) == sorted(i for i in ids if i != ids[1])
    assert store.get_invoice(ids[1])["status"] != APPROVED


def test_bulk_approve_skips_an_invoice_flagged_after_the_page_listed_it(demo_store):
    """The Azure check (or the page reader) flags an invoice after the queue page listed it as clean: a warning only
    the stored checks hold (finalise_coding does not raise it again). Bulk approval leaves it for a person."""
    store = demo_store
    ids = [c["id"] for c in clean_candidates(store)]
    victim = ids[0]
    issue = {"severity": "warning", "code": "AZURE_DI_DISAGREES", "message": "Azure read the total as 1,234.00"}
    assert store.update_after_check(victim, None, issue, ("AZURE_DI_DISAGREES",))
    result = bulk_approve(store, store.reference_data(), Settings(), ids, "jane")
    assert [i for i, _ in result["skipped"]] == [victim] and "Azure read" in result["skipped"][0][1]
    assert store.get_invoice(victim)["status"] == REVIEW


def test_bulk_approve_never_takes_a_check_that_always_needs_a_person(tmp_path):
    """S2-T3: a new vendor's first invoice (VENDOR_NEW, info only) is not "clean": a person opens it."""
    import json

    from ap_coder.ai.pipeline import finalise_coding
    from ap_coder.core.schema import InvoiceCoding
    from ap_coder.core.store import load_sample_setup

    from .conftest import SAMPLES

    store = Store(tmp_path / "ap.db")
    load_sample_setup(store)
    ref = store.reference_data()
    known = json.loads((SAMPLES / "ground_truth" / "northwind_ON_HST_NW-2026-0912.json").read_text())
    out, _ = finalise_coding(InvoiceCoding.model_validate(known), ref, Settings(), store=store)
    store.approve_invoice(store.add_invoice(tmp_path / "a.pdf", out, {}), out, "jane")
    doc = json.loads((SAMPLES / "ground_truth" / "chinook_AB_GST_CCO-26-10418.json").read_text())
    doc = {**doc, "vendor_name": "Totally New Supplies Inc", "confidence_score": 0.99}
    out, report = finalise_coding(InvoiceCoding.model_validate(doc), ref, Settings(), store=store)
    assert [(i.severity, i.code) for i in report.issues] == [("info", "VENDOR_NEW")] and not report.requires_review
    invoice_id = store.add_invoice(tmp_path / "b.pdf", out, report.to_dict())
    assert invoice_id not in [c["id"] for c in clean_candidates(store)]
    result = bulk_approve(store, ref, Settings(), [invoice_id], "jane")
    assert result["approved"] == [] and "First invoice from this vendor" in result["skipped"][0][1]
    assert store.get_invoice(invoice_id)["status"] == REVIEW


def _negated(doc):
    """A credit note for ``doc``: every amount with its sign turned."""
    d = json.loads(json.dumps(doc))
    for k in ("subtotal", "tax_total", "grand_total"):
        d[k] = -d[k]
    for li in d["line_items"]:
        li["amount"], li["unit_price"] = -li["amount"], -li["unit_price"]
    for tl in d["tax_lines"]:
        tl["taxable_amount"], tl["tax_amount"] = -tl["taxable_amount"], -tl["tax_amount"]
    return d


def test_bulk_approval_never_takes_a_credit_note(tmp_path):
    """B-132: a clean credit note (it credits an approved invoice, no warning) is applied by a person, as the
    touchless rules say: bulk approval does not list it nor approve it. Titled a credit note but printed with a
    positive total is one too (what capture read)."""
    from ap_coder.ai.pipeline import finalise_coding
    from ap_coder.core.schema import InvoiceCoding
    from ap_coder.core.store import load_sample_setup

    store = Store(tmp_path / "ap.db")
    load_sample_setup(store)
    ref = store.reference_data()
    known = json.loads((SAMPLES / "ground_truth" / "northwind_ON_HST_NW-2026-0912.json").read_text())
    out, _ = finalise_coding(InvoiceCoding.model_validate(known), ref, Settings(), store=store)
    store.approve_invoice(store.add_invoice(tmp_path / "a.pdf", out, {}), out, "jane")
    credit = {**_negated(known), "invoice_number": "CN-1", "original_invoice_number": known["invoice_number"],
              "confidence_score": 0.99, "due_date": ""}  # fmt: skip
    out, report = finalise_coding(InvoiceCoding.model_validate(credit), ref, Settings(), store=store)
    assert not [i for i in report.issues if i.severity != "info"] and not report.requires_review
    credit_id = store.add_invoice(tmp_path / "b.pdf", out, report.to_dict())
    # Titled a credit note, with its total printed positive (without the delivery line: not the same amount).
    tax = {**known["tax_lines"][0], "taxable_amount": 15850.0, "tax_amount": 2060.5}
    titled = {**known, "invoice_number": "CN-2", "confidence_score": 0.99, "line_items": known["line_items"][:4],
              "subtotal": 15850.0, "tax_lines": [tax], "tax_total": 2060.5, "grand_total": 17910.5}  # fmt: skip
    out, report = finalise_coding(InvoiceCoding.model_validate(titled), ref, Settings(), store=store)
    assert not [i for i in report.issues if i.severity != "info"] and not report.requires_review
    titled_id = store.add_invoice(tmp_path / "c.pdf", out, report.to_dict())
    store.save_capture(titled_id, {"credit_note": True})
    listed = [c["id"] for c in clean_candidates(store)]
    assert credit_id not in listed and titled_id not in listed
    result = bulk_approve(store, ref, Settings(), [credit_id, titled_id], "jane")
    assert result["approved"] == [] and all("credit note" in why for _, why in result["skipped"])
    assert store.get_invoice(credit_id)["status"] == REVIEW and store.get_invoice(titled_id)["status"] == REVIEW


def test_every_list_query_takes_more_ids_than_sqlite_takes_variables(demo_store, monkeypatch, few_sql_variables):
    """Each store call that takes a list of ids or keys: past an older SQLite build's 999 bound variables."""
    s = demo_store
    real = [i["id"] for i in s.list_invoices()]
    many = [*real, *range(10_000, 11_500)]
    assert sorted(i["id"] for i in s.get_invoices(many)) == sorted(real)
    assert sorted(r["id"] for r in s.invoice_columns(("id",), ids=many)) == sorted(real)
    assert isinstance(s.latest_events("approved", many), dict)
    assert isinstance(s.supplier_keys_for_invoices(many), list)
    keys = [f"none-{n}" for n in range(1_500)]
    assert s.prune_supplier_profiles(keys) == 0 and s.delete_vendors(keys) == 0
    s.delete_purchase_orders(keys)


def _loader_invoice():
    """An approved invoice whose invoice number and texts start with formula characters."""
    gt = json.loads((SAMPLES / "ground_truth" / f"{SAMPLE_STEM}.json").read_text())
    gt["invoice_number"] = "+A-100"
    gt["gl_distribution"] = [
        {"kind": "expense", "line_number": 1, "gl_code": "6010", "cost_center": "", "net_amount": 4350.0,
         "non_recoverable_tax": 0.0, "amount": 4350.0, "description": "- Labour"},
        {"kind": "expense", "line_number": 2, "gl_code": "6010", "cost_center": "", "net_amount": 1.0,
         "non_recoverable_tax": 0.0, "amount": 1.0, "description": "=1+1"},
    ]  # fmt: skip
    return {"id": 1, "final_output": gt, "reviewer": "Jane", "reviewed_at": "2026-10-01T10:00:00", "file_name": "a.pdf"}


def test_erp_csv_keeps_text_exactly_as_typed():
    """The CSV file is for an ERP loader (E2-2): an apostrophe made invoice +A-100 into '+A-100, and the ERP's
    supplier + invoice number duplicate check missed it."""
    data, _, _ = exports.build("csv", [_loader_invoice()], {}, 1)
    cells = [v for r in csv.DictReader(io.StringIO(data.decode("utf-8-sig"))) for v in r.values()]
    assert "+A-100" in cells and "- Labour" in cells and "=1+1" in cells
    assert not [v for v in cells if v.startswith("'")]


def test_xlsx_export_still_neutralises_formulas():
    """The xlsx export is opened in Excel by people: formula-looking text is stored as text."""
    data, _, _ = exports.build("xlsx", [_loader_invoice()], {}, 1)
    cells = [c for ws in load_workbook(io.BytesIO(data)).worksheets for row in ws.iter_rows() for c in row]
    assert not [c.coordinate for c in cells if c.data_type == "f"]
    assert any(c.value == "=1+1" and c.data_type == "s" for c in cells)


def test_an_invoice_approved_without_a_stored_distribution_still_has_its_lines_in_the_export():
    """An invoice approved by an older version has no gl_distribution stored: its posting lines are built again
    from its coding (as for JD Edwards), with the tax setup's accounts, instead of the file leaving it out."""
    from ap_coder.coding.tax import load_tax_mapping

    from .conftest import DATA

    doc = json.loads((SAMPLES / "ground_truth" / "northwind_ON_HST_NW-2026-0912.json").read_text())
    assert "gl_distribution" not in doc
    inv = {"id": 5, "final_output": doc}
    treatments = load_tax_mapping(DATA / "tax_gl_mapping.csv")
    for fmt in ("csv", "custom"):
        data, _, _ = exports.build(fmt, [inv], {}, 3, {}, tax_treatments=treatments)
        rows = list(csv.DictReader(io.StringIO(data.decode("utf-8-sig"))))
        assert len(rows) == len(doc["line_items"]) + 1, fmt  # each line and the HST
        assert round(sum(float(r["Amount"]) for r in rows), 2) == doc["grand_total"]
        assert "2310" in {r.get("GL account") or r.get("GLAccount") for r in rows}  # the HST account
    (lines,) = [ws for ws in load_workbook(io.BytesIO(exports.build("xlsx", [inv], {}, 3)[0])).worksheets
                if ws.title == "GL lines"]  # fmt: skip
    assert lines.max_row == len(doc["line_items"]) + 2  # the header, each line and the HST


def test_the_exports_say_which_bank_account_pays(demo_store):
    """D-66: the workbook and the CSV have Pay from (our bank account's name) and Pay from GL (its cash GL); the
    custom layout offers both as fields, and a saved layout is unchanged."""
    from ap_coder.coding import banks
    from ap_coder.erp import export_layout

    rows, _ = banks.clean_rows([{"code": "CAD-OP", "name": "Sample Bank CAD operating", "currency": "CAD",
                                 "gl_account": "1010", "is_default": "yes"}])  # fmt: skip
    demo_store.save_bank_accounts(rows)
    inv = demo_store.list_invoices(REVIEW)[0]
    final = {**demo_store.get_invoice(inv["id"])["ai_output"], "pay_from_account": "CAD-OP"}
    demo_store.approve_invoice(inv["id"], final, "Jordan Lee")
    invoices = demo_store.get_invoices([inv["id"]])
    accounts = {a["code"]: a for a in demo_store.bank_accounts()}
    data, _, _ = exports.build("csv", invoices, bank_accounts=accounts)
    lines = list(csv.DictReader(io.StringIO(data.decode("utf-8-sig"))))
    assert {(r["Pay from"], r["Pay from GL"]) for r in lines} == {("Sample Bank CAD operating", "1010")}
    book = load_workbook(io.BytesIO(exports.build("xlsx", invoices, bank_accounts=accounts)[0]))
    head = [c.value for c in book["Invoices"][1]]
    values = dict(zip(head, [c.value for c in book["Invoices"][2]], strict=True))
    assert (values["Pay from"], values["Pay from GL"]) == ("Sample Bank CAD operating", "1010")
    assert {"pay_from", "pay_from_gl"} <= set(export_layout.FIELDS)
    assert "pay_from" not in {c.field for c in export_layout.default_columns()}
    assert exports.pay_from({}, accounts) == {"pay_from": "", "pay_from_gl": ""}


def test_the_queue_says_clean_and_ready_by_todays_checks_after_a_setting_changes(demo_store):
    """B-815: JD Edwards turned on (no Default BU) after the invoices were read: the queue offered them as clean
    and Ready from the checks stored at intake, and bulk approval then approved none ("have no JDE account"). The
    queue re-checks what a settings change can change; back as before, they are clean again. A finding made after
    processing (the Azure check) still holds its invoice back."""
    from ap_coder.erp import jde
    from ap_coder.web.pages.review import attention

    store = demo_store
    ids = [c["id"] for c in clean_candidates(store)]
    assert len(ids) >= 2
    pending = [i for i in store.list_invoices(REVIEW) if i["id"] in ids]
    assert not set(attention(store, store.reference_data(), pending)) & set(ids)  # Ready
    store.set_setting(jde.SHOW_SETTING, "1")
    jde.save(store, jde.JdeSettings(use_default_rule=False, bu_from_cost_center=False))
    assert clean_candidates(store) == []
    reasons = attention(store, store.reference_data(), pending)
    assert set(ids) <= set(reasons) and all("error" in reasons[i] for i in ids)
    assert bulk_approve(store, store.reference_data(), Settings(), ids, "jane")["approved"] == []
    store.set_setting(jde.SHOW_SETTING, "0")
    assert sorted(c["id"] for c in clean_candidates(store)) == sorted(ids)
    assert not set(attention(store, store.reference_data(), pending)) & set(ids)
    issue = {"severity": "warning", "code": "AZURE_DI_DISAGREES", "message": "Azure read the total as 1,234.00"}
    assert store.update_after_check(ids[0], None, issue, ("AZURE_DI_DISAGREES",))
    store.set_setting(jde.SHOW_SETTING, "")
    assert ids[0] not in [c["id"] for c in clean_candidates(store)]


def test_an_invoice_picked_for_the_audit_sample_is_never_bulk_approved(demo_store):
    """The audit sample is an invoice touchless processing passed that a person must check: bulk approval offered it
    and approved it unseen (B-952)."""
    from ap_coder.web.pages.review import left_out

    store = demo_store
    picked = clean_candidates(store)[0]["id"]
    meta = store.get_invoice(picked)["meta"]
    meta.setdefault("capture", {}).update(audit=True, auto=False, reason="picked for the audit sample")
    with store._conn() as conn:
        conn.execute("UPDATE invoices SET meta = ? WHERE id = ?", (json.dumps(meta), picked))
    assert picked not in [c["id"] for c in clean_candidates(store)]
    from ap_coder.coding.bulk import sort_for_bulk

    left = sort_for_bulk(store)[1]
    assert left[picked] == "audit"
    assert left_out([{"id": picked}], {}, left) == "1 picked for the audit sample left for a person"
    result = bulk_approve(store, store.reference_data(), Settings(), [picked], "jane")
    assert result["approved"] == []
    assert result["skipped"] == [(picked, "Picked for the audit sample: a person checks it")]
    assert store.get_invoice(picked)["status"] == REVIEW
