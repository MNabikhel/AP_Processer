"""Tests for the tooling that keeps enterprise data local: aliases, labels, share report, doctor."""

import json

import pytest

from ap_coder import cli
from ap_coder.config import Settings
from ap_coder.doctor import FAIL, PASS, SKIP, WARN, format_checks, run_checks
from ap_coder.evaluation import evaluate
from ap_coder.labels import export_labels, load_labels
from ap_coder.reference_data import load_table
from ap_coder.share_report import build_share_report

from .conftest import SAMPLE_STEM, SAMPLES, FakeOpenAI, make_completion

# --- Reference data aliases ----------------------------------------------------------


def test_erp_column_aliases(tmp_path):
    p = tmp_path / "coa.csv"
    p.write_text("Main Account,Account Name,Is Active\n610100,Office supplies,Yes\n699900,Old,No\n")
    table = load_table(p, "GL account", "gl_code")
    assert table.codes == ["610100"]
    assert table.rows[0] == {"gl_code": "610100", "Account Name": "Office supplies"}


def test_canonical_column_wins_over_alias(tmp_path):
    p = tmp_path / "coa.csv"
    p.write_text("gl_code,account,name\n6000,ignored,Office\n")
    assert load_table(p, "GL account", "gl_code").codes == ["6000"]


def test_missing_key_column_lists_found_columns(tmp_path):
    p = tmp_path / "coa.csv"
    p.write_text("Ledger,Name\n6000,A\n")
    with pytest.raises(ValueError, match="Found columns: Ledger, Name"):
        load_table(p, "GL account", "gl_code")


# --- Labels round trip ------------------------------------------------------------------


def _write_prediction(out_dir, stem, doc, requires_review=False):
    out_dir.mkdir(parents=True, exist_ok=True)
    (out_dir / f"{stem}.json").write_text(json.dumps(doc))
    meta = {"source": f"/x/{stem}.pdf", "status": "ok", "validation": {"requires_review": requires_review}}
    (out_dir / f"{stem}.validation.json").write_text(json.dumps(meta))


@pytest.mark.parametrize("suffix", [".xlsx", ".csv"])
def test_labels_round_trip(tmp_path, ground_truth, reference, suffix):
    out = tmp_path / "out"
    _write_prediction(out, "inv", ground_truth, requires_review=True)
    labels = tmp_path / f"labels{suffix}"
    assert export_labels(out, labels, reference=reference) == 1

    docs, unreviewed = load_labels(labels)
    assert docs == {} and unreviewed == ["inv"], "nothing counts until rows are marked reviewed"

    _mark_reviewed(labels, fix_gl={1: "6020"})
    docs, unreviewed = load_labels(labels)
    assert unreviewed == []
    doc = docs["inv"]
    assert doc["invoice_date"] == "2026-09-14" and doc["grand_total"] == 18017.85
    assert [li["predicted_gl_code"] for li in doc["line_items"]][:2] == ["6020", "1500"]

    report = evaluate(out, labels)
    assert report.gl_accuracy == round(4 / 5, 4)


def _mark_reviewed(path, fix_gl):
    if path.suffix == ".csv":
        import csv

        with path.open(encoding="utf-8-sig", newline="") as fh:
            rows = list(csv.DictReader(fh))
        for r in rows:
            r["reviewed"] = "Y"
            if int(r["line_number"]) in fix_gl:
                r["gl_code"] = fix_gl[int(r["line_number"])]
        with path.open("w", encoding="utf-8-sig", newline="") as fh:
            w = csv.DictWriter(fh, fieldnames=list(rows[0]))
            w.writeheader()
            w.writerows(rows)
        return
    from openpyxl import load_workbook

    wb = load_workbook(path)
    ws = wb["Labels"]
    header = [c.value for c in ws[1]]
    for row in ws.iter_rows(min_row=2):
        cells = dict(zip(header, row, strict=False))
        cells["reviewed"].value = "y"
        if int(cells["line_number"].value) in fix_gl:
            cells["gl_code"].value = fix_gl[int(cells["line_number"].value)]
    wb.save(path)


def test_labels_workbook_has_dropdowns_and_text_cells(tmp_path, ground_truth, reference):
    from openpyxl import load_workbook

    gt = dict(ground_truth, invoice_number="000123")
    _write_prediction(tmp_path / "out", "inv", gt)
    export_labels(tmp_path / "out", tmp_path / "l.xlsx", reference=reference)
    wb = load_workbook(tmp_path / "l.xlsx")
    assert wb.sheetnames == ["Labels", "GL Accounts", "Cost Centers", "Instructions"]
    ws = wb["Labels"]
    assert len(ws.data_validations.dataValidation) == 2
    header = [c.value for c in ws[1]]
    assert ws.cell(2, header.index("invoice_number") + 1).value == "000123"
    assert ws.cell(3, header.index("vendor_name") + 1).value in (None, ""), "header only on first row"


def test_blind_labels_hide_model_codes(tmp_path, ground_truth):
    _write_prediction(tmp_path / "out", "inv", ground_truth)
    export_labels(tmp_path / "out", tmp_path / "l.csv", blind=True)
    text = (tmp_path / "l.csv").read_text(encoding="utf-8-sig")
    assert "6030" not in text and "Production cloud" not in text


def test_inserted_row_inherits_document(tmp_path):
    p = tmp_path / "l.csv"
    p.write_text(
        "reviewed,document,vendor_name,invoice_number,invoice_date,currency,subtotal,tax_total,grand_total,"
        "line_number,description,quantity,unit_price,amount,gl_code,cost_center\n"
        "Y,inv,V,1,2026-01-01,USD,15,0,15,1,a,1,10,10,6000,CC100\n"
        "Y,,,,,,,,,2,added line,1,5,5,6800,CC100\n"
    )
    docs, _ = load_labels(p)
    assert [li["line_number"] for li in docs["inv"]["line_items"]] == [1, 2]


# --- Share report -----------------------------------------------------------------------


def test_share_report_is_redacted(tmp_path, settings, reference, ground_truth, sample_markdown_path):
    from ap_coder.inference import InvoiceCoder
    from ap_coder.pipeline import InvoicePipeline, write_outputs

    pipe = InvoicePipeline(
        settings,
        reference,
        coder=InvoiceCoder(settings, reference, client=FakeOpenAI(make_completion(json.dumps(ground_truth)))),
    )
    write_outputs(pipe.process(sample_markdown_path), tmp_path / "out")

    key = tmp_path / "key.csv"
    gt_dir = tmp_path / "gt"
    gt_dir.mkdir()
    (gt_dir / f"{SAMPLE_STEM}.json").write_text((SAMPLES / "ground_truth" / f"{SAMPLE_STEM}.json").read_text())
    text = build_share_report(tmp_path / "out", gt_dir, key_file=key)
    for secret in ("Northwind", "18017", "18,017", "NW-2026", "Latitude", SAMPLE_STEM, "6010", "CC400"):
        assert secret not in text, secret
    assert "doc-01" in text and "GL code accuracy: 100.0%" in text
    assert SAMPLE_STEM in key.read_text()

    with_codes = build_share_report(tmp_path / "out", gt_dir, include_codes=True)
    assert "codes included: yes" in with_codes


def test_share_report_failure_category_hides_message(tmp_path):
    out = tmp_path / "out"
    out.mkdir()
    meta = {"source": "/x/acme.pdf", "status": "failed", "error": "CodingError: Model refused: ACME secret 123"}
    (out / "acme.validation.json").write_text(json.dumps(meta))
    text = build_share_report(out)
    assert "CodingError/model_refusal" in text
    assert "ACME" not in text and "acme" not in text


def test_share_report_on_a_database_with_demo_invoices(tmp_path):
    """Demo invoices recorded no table count: a KeyError half-way, after the key file was already written."""
    from ap_coder.demo import load_demo
    from ap_coder.store import Store, load_sample_setup

    store = Store(tmp_path / "ap.db")
    load_sample_setup(store)
    load_demo(store)
    with store._conn() as conn:  # a database from before demo invoices recorded one
        conn.execute("UPDATE invoices SET meta = json_remove(meta, '$.extraction.table_count')")
    key = tmp_path / "share_key.csv"
    text = build_share_report(tmp_path / "none", key_file=key, db_path=tmp_path / "ap.db")
    assert "tables per invoice: n/a" in text and "doc-01" in text
    assert key.read_text().startswith("alias,file") and not list(tmp_path.glob("*.partial"))


def test_share_report_that_fails_leaves_no_key_file(tmp_path, monkeypatch):
    from ap_coder import share_report

    out = tmp_path / "out"
    out.mkdir()
    (out / "a.validation.json").write_text(json.dumps({"source": "/x/a.pdf", "status": "failed", "error": "X"}))
    monkeypatch.setattr(share_report, "_failure_category", lambda error: 1 / 0)  # any failure while building
    key = tmp_path / "share_key.csv"
    with pytest.raises(ZeroDivisionError):
        build_share_report(out, key_file=key)
    assert not key.exists() and not list(tmp_path.glob("*.partial"))


# --- Doctor -----------------------------------------------------------------------------


def test_doctor_offline_never_prints_secrets(reference):
    from ap_coder.config import DocumentIntelligenceSettings, OpenAISettings

    s = Settings(
        document_intelligence=DocumentIntelligenceSettings(endpoint="https://secret-di.example", api_key="k-di-123"),
        openai=OpenAISettings(endpoint="https://secret-aoai.example", api_key="k-oai-456", model_name="gpt-4o"),
    )
    checks = run_checks(s, lambda: reference)
    text = format_checks(checks)
    for secret in ("secret-di", "secret-aoai", "k-di-123", "k-oai-456"):
        assert secret not in text
    by_area = {c.area: c for c in checks}
    assert by_area["GL accounts"].status == PASS and "enforced as schema enum" in by_area["GL accounts"].detail
    assert by_area["AOAI connectivity"].status == SKIP
    assert not any(c.status == FAIL for c in checks)


def test_doctor_flags_missing_config_and_bad_reference():
    def broken():
        raise ValueError("chart_of_accounts.csv: missing required column 'gl_code'")

    checks = run_checks(Settings(), broken)
    by_area = {c.area: c for c in checks}
    # Azure is optional: without it invoices are read and coded on this computer.
    assert by_area["DI endpoint"].status == SKIP
    assert by_area["AOAI endpoint"].status == SKIP
    assert by_area["reference data"].status == FAIL
    assert "AOAI model" not in by_area


def test_doctor_large_chart_falls_back_from_enums(tmp_path):
    coa = tmp_path / "coa.csv"
    coa.write_text("gl_code,name\n" + "".join(f"{i},Account {i}\n" for i in range(600)))
    cc = tmp_path / "cc.csv"
    cc.write_text("cost_center,name\nCC1,Ops\n")
    from ap_coder.reference_data import load_reference_data

    checks = run_checks(Settings(), lambda: load_reference_data(coa, cc))
    coa_check = next(c for c in checks if c.area == "GL accounts")
    assert coa_check.status == WARN and "free-text" in coa_check.detail


def test_cli_doctor_uses_sample_data_by_default(capsys, monkeypatch, tmp_path):
    monkeypatch.chdir(tmp_path)
    monkeypatch.delenv("AP_REFERENCE_DIR", raising=False)
    assert cli.main(["doctor"]) == 0  # no Azure and no model: still ready (read and coded on this computer)
    out = capsys.readouterr().out
    assert "BUNDLED SAMPLE DATA" in out


def test_cli_prefers_private_reference(capsys, monkeypatch, tmp_path):
    monkeypatch.chdir(tmp_path)
    monkeypatch.delenv("AP_REFERENCE_DIR", raising=False)
    ref = tmp_path / "private" / "reference"
    ref.mkdir(parents=True)
    (ref / "chart_of_accounts.csv").write_text("Account,Description\n7777,Widgets\n")
    (ref / "cost_centers.csv").write_text("Cost Centre,Name\nZ1,Zed\n")
    assert cli.main(["schema"]) == 0
    schema = json.loads(capsys.readouterr().out)
    line = schema["properties"]["line_items"]["items"]["properties"]
    assert line["predicted_gl_code"]["enum"] == ["7777", "UNASSIGNED"]
    assert line["predicted_cost_center"]["enum"] == ["Z1", "UNASSIGNED"]


def test_cli_labels_refuses_to_overwrite(tmp_path, ground_truth, capsys):
    _write_prediction(tmp_path / "out", "inv", ground_truth)
    dest = tmp_path / "labels.csv"
    args = ["labels", "--predictions", str(tmp_path / "out"), "-o", str(dest)]
    assert cli.main(args) == 0
    dest.write_text("corrected by AP team")
    assert cli.main(args) == 2
    assert dest.read_text() == "corrected by AP team"
    assert cli.main([*args, "--force"]) == 0
