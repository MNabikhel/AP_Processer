"""Outlook .msg emails and email triage: every email from the AP mailbox sorted, no invoice ever lost.

The triage cases here are written by hand, apart from the benchmark's generated emails (``bench triage``), and aim
at the mistakes that would lose an invoice: a bill that calls itself a statement, an invoice that mentions a quote,
a scan with nothing to read, a bank-detail change hidden in an invoice's footer."""

from __future__ import annotations

import json
import time
from email.message import EmailMessage

import pymupdf
import pytest

from ap_coder.bench.msgwriter import Attachment, Message, build_msg
from ap_coder.core.store import Store, load_sample_setup
from ap_coder.intake import inbox, triage
from ap_coder.intake.mailbox import PERSON_FOLDER, STATEMENTS_FOLDER, read_msg_mail, triage_log, unpack, unpack_folder
from ap_coder.intake.msgfile import MsgError, read_msg

from .webclient import client_for, follow, post


def _pdf(*lines: str) -> bytes:
    doc = pymupdf.open()
    page = doc.new_page()
    for n, line in enumerate(lines):
        page.insert_text((50, 60 + 16 * n), line, fontsize=10)
    return doc.tobytes()


INVOICE = _pdf("Northwind Supplies", "INVOICE", "Invoice No. INV-4471", "Subtotal 100.00", "GST 5.00", "Total 105.00")
STATEMENT = _pdf("Northwind Supplies", "STATEMENT OF ACCOUNT", "Invoice INV-4471 105.00", "Invoice INV-4502 88.10",
                 "Invoice INV-4519 12.00", "Current 31-60 61-90 Over 90")  # fmt: skip
QUOTE = _pdf("Northwind Supplies", "QUOTATION", "Valid for 30 days", "Rebar 15M 40 1,200.00")
LOGO = b"\x89PNG\r\n\x1a\n" + b"0" * 4000


def _kind(subject: str, body: str, *docs: tuple[str, bytes]) -> triage.EmailTriage:
    return triage.classify_email(subject, body, [triage.document_for(n, c) for n, c in docs])


# --- .msg -------------------------------------------------------------------------------------------------------


def test_a_msg_email_is_read_like_an_eml(tmp_path):
    inner = Message(subject="Fwd: older invoice", body="see below", attachments=[Attachment("INV-0.pdf", INVOICE)])
    msg = Message(subject="Invoice INV-4471", sender_name="Northwind AR", sender_email="ar@northwind.example",
                  body="Please find our invoice attached.",
                  attachments=[Attachment("INV-4471.pdf", INVOICE),
                               Attachment("image001.png", LOGO, "image/png", "logo"),
                               Attachment("terms.docx", b"PK doc", "application/msword")],
                  attached_emails=[inner])  # fmt: skip
    mail = read_msg_mail(build_msg(msg))
    assert mail.subject == "Invoice INV-4471" and mail.sender == "Northwind AR <ar@northwind.example>"
    assert "invoice attached" in mail.body
    assert [n for n, _ in mail.found] == ["INV-4471.pdf", "INV-0.pdf"]  # the forwarded email's invoice too
    assert dict(mail.found)["INV-4471.pdf"] == INVOICE
    assert any("image001.png" in s and "logo" in s for s in mail.skipped)  # the signature logo
    assert any("terms.docx" in s for s in mail.skipped)

    path = tmp_path / "Invoice from Northwind.msg"
    path.write_bytes(build_msg(msg))
    result = unpack(path)
    assert [p.name for p in result.saved] == [
        "Invoice from Northwind - INV-4471.pdf",
        "Invoice from Northwind - INV-0.pdf",
    ]
    # terms.docx is not read: the email waits for a person, in case it is an invoice (B-140)
    assert (tmp_path / "emails" / PERSON_FOLDER / "Invoice from Northwind.msg").exists()
    assert result.triage.kind == triage.INVOICE and any("terms.docx" in r for r in result.triage.reasons)


def test_a_large_attachment_outside_the_mini_stream(tmp_path):
    big = INVOICE + b"%" + b"x" * 200_000  # read through the file's own sectors, not the mini stream
    assert (
        read_msg(build_msg(Message(subject="big", attachments=[Attachment("big.pdf", big)]))).attachments[0].data == big
    )


@pytest.mark.parametrize("damage", ["not an ole file", "truncated", "loop"])
def test_a_damaged_msg_is_filed_away_with_the_reason(tmp_path, damage):
    raw = bytearray(build_msg(Message(subject="x", attachments=[Attachment("a.pdf", INVOICE)])))
    if damage == "not an ole file":
        raw = bytearray(b"From: x\n\nnot an outlook file at all" * 30)
    elif damage == "truncated":
        raw = raw[:1500]
    else:  # the directory's sector chain points at itself
        first_dir = int.from_bytes(raw[0x30:0x34], "little")
        fat_sector = int.from_bytes(raw[0x4C:0x50], "little")
        off = (fat_sector + 1) * 512 + first_dir * 4
        raw[off : off + 4] = first_dir.to_bytes(4, "little")
    with pytest.raises(MsgError):
        read_msg(bytes(raw)) if damage != "loop" else _read_all(bytes(raw))
    (tmp_path / "odd.msg").write_bytes(bytes(raw))
    (result,) = unpack_folder(tmp_path, now=time.time() + 60)
    assert (
        result.error.startswith("could not be read") and (tmp_path / "emails" / "could not read" / "odd.msg").exists()
    )


def _read_all(raw: bytes) -> None:
    read_msg(raw)
    raise MsgError("read")  # a loop that was cut short still reads as a damaged file in unpack


def _entries_named(raw: bytes, name: str) -> list[int]:
    """The offsets of the directory entries called ``name`` in a .msg made by ``build_msg``."""
    key, out, at = name.encode("utf-16-le") + b"\x00\x00", [], 0
    while (at := raw.find(key, at)) != -1:
        out.append(at)
        at += 1
    return out


def test_attachments_sharing_one_stream_cannot_fill_the_memory():
    """Each stream is checked against the file's size, but many attachments can point at the same big stream: a
    small .msg must not read as many times its size."""
    big = INVOICE + b"%" + b"x" * 40_000
    small = [Attachment(f"s{n}.pdf", b"%PDF-1.4 s") for n in range(10)]
    raw = bytearray(build_msg(Message(subject="x", attachments=[Attachment("big.pdf", big), *small])))
    offsets = _entries_named(bytes(raw), "__substg1.0_37010102")
    sizes = {off: int.from_bytes(raw[off + 120 : off + 128], "little") for off in offsets}
    target = max(sizes, key=sizes.get)
    for off in offsets:  # every attachment's data now points at the big one's sectors
        raw[off + 116 : off + 128] = raw[target + 116 : target + 128]
    with pytest.raises(MsgError):
        read_msg(bytes(raw))


def test_a_directory_walk_reads_each_entry_once():
    """Storages that all claim the same chain of children: each entry belongs to one storage, so the walk stays as
    large as the directory instead of growing with its square."""
    from ap_coder.bench.msgwriter import _Node, compound_file
    from ap_coder.intake.msgfile import CompoundFile

    n = 200
    raw = bytearray(compound_file([_Node(f"s{i}") for i in range(n)] + [_Node(f"t{i}", b"x") for i in range(n)]))
    index = {e.name: i for i, e in enumerate(CompoundFile(bytes(raw)).entries)}
    at = {name: _entries_named(bytes(raw), name)[0] for name in index if name[:1] in "st"}
    free = (0xFFFFFFFF).to_bytes(4, "little")
    for i in range(n):  # t0 -> t1 -> ... as left siblings; every storage's child is t0
        nxt = index[f"t{i + 1}"].to_bytes(4, "little") if i + 1 < n else free
        raw[at[f"t{i}"] + 68 : at[f"t{i}"] + 76] = nxt + free
        raw[at[f"s{i}"] + 76 : at[f"s{i}"] + 80] = index["t0"].to_bytes(4, "little")
    cf = CompoundFile(bytes(raw))
    assert sum(len(e.children) for e in cf.entries) <= len(cf.entries)


def _renumbered(raw: bytes) -> bytes:
    """The same .msg with its directory entries in reverse order (the root stays first): a storage now comes before
    the storage that holds it. The file format puts no order on entries, so this is still a valid file."""
    from ap_coder.intake.msgfile import CompoundFile

    cf, out = CompoundFile(raw), bytearray(raw)
    sectors, s = [], int.from_bytes(raw[0x30:0x34], "little")
    while s not in (0xFFFFFFFE, 0xFFFFFFFF):
        sectors.append(s)
        s = cf.fat[s]
    slots = [(sec + 1) * cf.sector + k * 128 for sec in sectors for k in range(cf.sector // 128)]
    order = [0, *range(len(cf.entries) - 1, 0, -1)]  # new position -> old index
    new_of = {old: new for new, old in enumerate(order)}
    for new, old in enumerate(order):
        entry = bytearray(raw[slots[old] : slots[old] + 128])
        links = [int.from_bytes(entry[at : at + 4], "little") for at in (68, 72, 76)]
        entry[68:80] = b"".join(new_of.get(x, x).to_bytes(4, "little") for x in links)
        out[slots[new] : slots[new] + 128] = entry
    return bytes(out)


def test_the_order_of_directory_entries_never_drops_an_attachment():
    inner = Message(subject="Fwd", body="older", attachments=[Attachment("INV-0.pdf", INVOICE)])
    msg = Message(subject="Invoices", body="attached", attachments=[Attachment("INV-1.pdf", INVOICE)],
                  attached_emails=[inner])  # fmt: skip
    out = read_msg(_renumbered(build_msg(msg)))
    assert [a.name for a in out.attachments] == ["INV-1.pdf", "INV-0.pdf"] and "older" in out.body


# --- triage: no invoice lost --------------------------------------------------------------------------------------


@pytest.mark.parametrize(
    "subject, doc",
    [
        ("Invoice", _pdf("ACME", "INVOICE 991", "Quote reference Q-5521", "Total 1,200.00")),  # a quote mentioned
        ("Your hydro bill", _pdf("City Hydro", "Billing statement", "Previous balance 120.00", "Billing period Sep",
                                 "Usage 400 kWh", "Amount due 98.40")),  # a utility "statement" is a bill
        ("Statement", _pdf("Atlas", "Statement date 2026-09-30", "INVOICE 12", "Amount due 50.00")),
        ("Scan", b"%PDF-1.4 a scan with no text layer"),  # nothing to read: the queue, never set aside
        ("", _pdf("ACME", "FACTURE no 55", "Total 10,00 $")),
        ("Packing slip and invoice", _pdf("ACME", "INVOICE 77", "Packing slip enclosed", "Total 80.00")),
        ("Order confirmation", _pdf("ACME", "Order 4471 confirmed", "Total 1,200.00")),  # says nothing: the queue
    ],
)  # fmt: skip
def test_anything_that_may_be_an_invoice_goes_to_the_queue(subject, doc):
    t = _kind(subject, "", ("doc.pdf", doc))
    assert [d.kind for _, d in t.documents] in (["invoice"], ["credit_note"]) and t.kind in triage.QUEUE_KINDS


def test_a_picture_or_a_scan_named_statement_still_goes_to_the_queue():
    t = _kind("Statement", "", ("statement.pdf", b"%PDF-1.4 scan"), ("photo.jpg", b"\xff\xd8\xff" + b"0" * 20_000))
    assert [d.kind for _, d in t.documents] == ["invoice", "invoice"]


@pytest.mark.parametrize("title", ["AVOIR", "Facture d'avoir", "Credit Invoice", "Avis de crédit", "Credit Note",
                                   "Note de crédit", "Credit Memo"])  # fmt: skip
def test_triage_and_capture_know_the_same_credit_note_titles(title):
    """Email triage and capture read a document's title from one list: a title triage sends to the queue as a credit
    note is a credit for capture too (its amounts posted negative, and a person applies it)."""
    from ap_coder.capture.reader import _CREDIT_TITLE as capture_title
    from ap_coder.capture.reader import plain

    doc = triage.Document("doc.pdf", f"{title}\nNo 55\nTotal 10,00")
    assert triage.classify_document(doc).kind == triage.CREDIT
    assert capture_title.match(plain(title))


def test_kinds_in_english_and_french():
    assert _kind("Statement", "", ("soa.pdf", STATEMENT)).kind == triage.STATEMENT
    assert (
        _kind(
            "Relevé", "", ("releve.pdf", _pdf("RELEVÉ DE COMPTE", "Facture 1 10", "Facture 2 20", "Facture 3 5"))
        ).kind
        == "statement"
    )
    assert _kind("Avoir", "", ("nc.pdf", _pdf("NOTE DE CRÉDIT", "No 55", "Total -10,00"))).kind == triage.CREDIT
    assert _kind("Credit memo", "", ("cm.pdf", _pdf("CREDIT MEMO 12", "Total -50.00"))).kind == triage.CREDIT
    assert _kind("Rappel de paiement", "Votre facture 123 est en souffrance.").kind == triage.REMINDER
    assert _kind("PAST DUE", "Invoice 4471 is 45 days overdue.").kind == triage.REMINDER
    assert _kind("Quote", "", ("q.pdf", QUOTE)).kind == triage.NOT_FOR_AP
    assert _kind("Soumission", "", ("s.pdf", _pdf("SOUMISSION", "Valide 30 jours"))).kind == triage.NOT_FOR_AP
    auto = _kind("Automatic reply: Invoice attached", "I am out of the office. > Please find the invoice attached.")
    assert auto.kind == triage.NOT_FOR_AP and "automatic reply" in auto.reasons[0]
    portal = _kind("Your invoice is ready", "View it online at https://portal.example/inv/4471")
    assert portal.kind == triage.NOT_FOR_AP and "has none attached" in portal.reasons[0] and portal.needs_person


@pytest.mark.parametrize(
    "subject, body, docs",
    [
        ("Change of banking information", "Please update your records with our new banking information.", ()),
        ("Mise à jour", "Veuillez noter nos nouvelles coordonnées bancaires pour les paiements futurs.", ()),
        ("Invoice 4471", "Invoice attached. Please note our bank details have changed.", (("i.pdf", INVOICE),)),
        ("Invoice 4471", "", (("i.pdf", _pdf("INVOICE 4471", "Total 10.00", "Please note our new bank details:",
                                             "Transit 12345 Account 7654321")),)),  # in the invoice's footer
        ("Letter", "", (("l.pdf", _pdf("Notice", "Changement de coordonnées bancaires", "Compte 123")),)),
    ],
)  # fmt: skip
def test_every_bank_detail_change_is_flagged(subject, body, docs):
    t = _kind(subject, body, *docs)
    assert t.possible_fraud and t.needs_person and "bank details" in " ".join(t.reasons)
    for doc, dec in t.documents:  # an invoice in such an email is still read
        if b"INVOICE" in dict(docs)[doc.name]:
            assert dec.kind == triage.INVOICE


def test_a_reminder_with_the_invoice_attached_sends_the_invoice_to_the_queue():
    t = _kind("PAST DUE: invoice 4471", "Copy attached.", ("inv.pdf", INVOICE))
    assert t.kind == triage.INVOICE and t.documents[0][1].kind == triage.INVOICE
    assert "duplicate check" in t.reasons[0] and not t.needs_person


def test_a_spreadsheet_is_a_statement_only_when_it_says_so():
    sheet = b"Invoice,Amount,Balance\nINV-1,10,10\nINV-2,20,30\n"
    assert _kind("Statement of account", "", ("soa.csv", sheet)).kind == triage.STATEMENT
    other = _kind("Price list", "", ("prices.csv", b"Item,Price\nRebar,10\n"))
    assert other.kind == triage.NOT_FOR_AP and other.needs_person


def test_a_spreadsheet_that_unzips_to_far_more_is_not_opened(monkeypatch):
    """A small .xlsx can hold hundreds of MB of XML that openpyxl reads in full: its sizes are checked first."""
    import io
    import zipfile

    import openpyxl

    wb = openpyxl.Workbook()
    wb.active.append(["Statement of account"])
    good = io.BytesIO()
    wb.save(good)
    assert "Statement of account" in triage.sheet_text("soa.xlsx", good.getvalue())
    bomb = io.BytesIO()
    with zipfile.ZipFile(bomb, "w", zipfile.ZIP_DEFLATED) as z, zipfile.ZipFile(good) as src:
        for item in src.infolist():
            if item.filename != "xl/sharedStrings.xml":
                z.writestr(item, src.read(item))
        with z.open("xl/sharedStrings.xml", "w") as fh:
            for _ in range(25):  # about 25 MB once unzipped, from under 100 KB
                fh.write(b"<si><t>a</t></si>" * 60_000)
    assert len(bomb.getvalue()) < 200_000
    opened: list[int] = []

    def load_workbook(*args, **kwargs):
        opened.append(1)
        raise ValueError("opened")

    monkeypatch.setattr(openpyxl, "load_workbook", load_workbook)
    assert triage.sheet_text("soa.xlsx", bomb.getvalue()) == "" and opened == []


# --- where unpack puts things -------------------------------------------------------------------------------------


def _eml(path, subject, body, *attachments):
    m = EmailMessage()
    m["Subject"], m["From"] = subject, "ar@vendor.example"
    m.set_content(body)
    for name, content in attachments:
        sub = "pdf" if name.endswith(".pdf") else "octet-stream"
        m.add_attachment(content, maintype="application", subtype=sub, filename=name)
    path.write_bytes(m.as_bytes())
    return path


def test_unpack_sorts_attachments_and_logs_each_email(tmp_path):
    _eml(tmp_path / "a.eml", "Invoice and slip", "attached", ("inv.pdf", INVOICE), ("slip.pdf", _pdf("PACKING SLIP")))
    _eml(tmp_path / "b.eml", "Statement", "attached", ("soa.pdf", STATEMENT))
    _eml(tmp_path / "c.eml", "New bank details", "We have changed banks: please update your records.")
    a, b, c = unpack_folder(tmp_path, now=time.time() + 60)
    assert [p.name for p in a.saved] == ["a - inv.pdf"] and [p.parent.name for p in a.held] == [PERSON_FOLDER]
    assert [p.parent.name for p in b.statements] == [STATEMENTS_FOLDER] and b.saved == []
    assert c.triage.kind == triage.BANK_CHANGE and (tmp_path / "emails" / PERSON_FOLDER / "c.eml").exists()
    assert (tmp_path / "emails" / "b.eml").exists()  # nothing for a person: filed as before
    log = triage_log(tmp_path)
    assert [e["kind"] for e in log] == ["bank_change", "statement", "invoice"]  # newest first
    slip = log[2]["documents"][1]
    assert (slip["name"], slip["kind"], slip["file"]) == (
        "slip.pdf",
        "not_for_ap",
        "emails/needs a person/a - slip.pdf",
    )
    (tmp_path / "emails" / "triage.jsonl").open("a").write("not json\n")  # a damaged line is skipped
    assert len(triage_log(tmp_path)) == 3


def test_what_a_person_does_with_an_email_set_aside(tmp_path):
    store = Store(tmp_path / "ap.db")
    folder = tmp_path / "invoices"
    folder.mkdir()
    _eml(folder / "q.eml", "Quote for your project", "see attached", ("quote.pdf", QUOTE))
    _eml(folder / "x.eml", "New bank account", "Please note our new bank details.")
    unpack_folder(folder, now=time.time() + 60)
    rows = inbox.needs_person(store, folder)
    assert [r["kind"] for r in rows] == ["bank_change", "not_for_ap"]  # possible fraud first
    quote = rows[1]
    assert quote["docs"][0]["there"] and inbox.file_of(folder, quote["id"], -1).name == "q.eml"
    assert inbox.file_of(folder, quote["id"], 5) is None and inbox.file_of(folder, "nope", 0) is None
    moved = inbox.to_queue(store, folder, quote["id"], 0, "clerk")  # it is an invoice after all
    assert moved.parent == folder and moved.read_bytes() == QUOTE
    assert inbox.to_queue(store, folder, quote["id"], 0) is None  # already moved
    assert inbox.mark_done(store, folder, quote["id"], "clerk") and not inbox.mark_done(store, folder, "nope")
    assert [r["kind"] for r in inbox.needs_person(store, folder)] == ["bank_change"]
    actions = [e["action"] for e in store.events()]
    assert "email_triage_done" in actions and "email_attachment_to_queue" in actions


def test_an_open_email_stays_in_needs_a_person_however_many_were_sorted_after_it(tmp_path):
    store = Store(tmp_path / "ap.db")
    bank = {"id": "fraud01", "kind": "bank_change", "possible_fraud": True, "needs_person": True, "documents": []}
    newer = [{"id": f"inv{n}", "kind": "invoice", "needs_person": False, "documents": []} for n in range(6000)]
    (tmp_path / "emails").mkdir()
    (tmp_path / "emails" / "triage.jsonl").write_text("".join(json.dumps(r) + "\n" for r in [bank, *newer]))
    assert [r["id"] for r in inbox.needs_person(store, tmp_path)] == ["fraud01"]
    assert inbox.mark_done(store, tmp_path, "fraud01")  # and a person can still mark it done
    assert inbox.needs_person(store, tmp_path) == []


def test_a_path_in_the_log_never_leaves_the_invoices_folder(tmp_path):
    folder = tmp_path / "invoices"
    (folder / "emails").mkdir(parents=True)
    (tmp_path / "secret.txt").write_text("x")
    entry = {"id": "abc", "needs_person": True, "email": "../secret.txt", "documents": [{"file": "../secret.txt"}]}
    (folder / "emails" / "triage.jsonl").write_text(json.dumps(entry) + "\n")
    assert inbox.file_of(folder, "abc", -1) is None and inbox.file_of(folder, "abc", 0) is None


# --- the dashboard ------------------------------------------------------------------------------------------------


def test_the_process_page_lists_emails_that_need_a_person(tmp_path):
    store = Store(tmp_path / "private" / "ap_coder.db")
    load_sample_setup(store)
    client = client_for(store)
    from ap_coder.web.core import invoice_dir

    folder = invoice_dir(store)
    folder.mkdir(parents=True, exist_ok=True)
    _eml(folder / "bank.eml", "Updated bank details", "Please update our EFT details: new account number below.")
    _eml(folder / "soa.eml", "Statement", "attached", ("soa.csv", b"Statement of account\nInvoice,Amount\nINV-1,10\n"))
    for p in folder.glob("*.eml"):
        unpack(p, folder)
    page = client.get("/process").text
    assert "Emails that need a person" in page and "Bank-detail change (possible fraud)" in page
    (row,) = inbox.needs_person(store, folder)
    assert client.get(f"/process/inbox/{row['id']}/file/-1").status_code == 200
    assert client.get("/process/inbox/nope/file/0").status_code == 404
    assert "Marked as dealt with" in follow(client, post(client, f"/process/inbox/{row['id']}/done"))
    assert "Emails that need a person" not in client.get("/process").text
    statements = client.get("/statements").text
    assert "Statements received by email" in statements and "soa.csv" in statements
    (st,) = inbox.statements(store, folder)
    answer = post(client, f"/statements/emailed/{st['id']}/{st['n']}", {"vendor": ""})
    assert answer.status_code == 200 and "/statements?" in answer.json()["redirect"]


def _sheet_bomb() -> bytes:
    """A small .xlsx (under 200 KB) whose shared strings unzip to about 25 MB, more than triage lets openpyxl read."""
    import io
    import zipfile

    import openpyxl

    wb = openpyxl.Workbook()
    wb.active.append(["Invoice", "Amount"])
    good = io.BytesIO()
    wb.save(good)
    bomb = io.BytesIO()
    with zipfile.ZipFile(bomb, "w", zipfile.ZIP_DEFLATED) as z, zipfile.ZipFile(good) as src:
        for item in src.infolist():
            if item.filename != "xl/sharedStrings.xml":
                z.writestr(item, src.read(item))
        with z.open("xl/sharedStrings.xml", "w") as fh:
            for _ in range(25):
                fh.write(b"<si><t>a</t></si>" * 60_000)
    return bomb.getvalue()


def test_every_reader_of_an_uploaded_sheet_checks_its_unzipped_sizes(tmp_path, monkeypatch):
    """W2-01: triage would not open a spreadsheet that unzips to far more (B-63), but the email was still filed as a
    statement and Reconcile on the Statements page handed it to pandas unchecked. Every sheet that is uploaded or
    emailed (a statement, the JDE GL map, AP line history, the JDE worksheet) is checked before a parser opens it."""
    import openpyxl
    import pandas as pd

    store = Store(tmp_path / "private" / "ap_coder.db")
    load_sample_setup(store)
    client = client_for(store)
    from ap_coder.web.core import invoice_dir

    folder = invoice_dir(store)
    folder.mkdir(parents=True, exist_ok=True)
    bomb = _sheet_bomb()
    _eml(folder / "soa.eml", "Statement of account", "Please find our statement attached.", ("soa.xlsx", bomb))
    unpack(folder / "soa.eml", folder)
    (st,) = inbox.statements(store, folder)
    opened: list[str] = []

    def refuse_to_open(*args, **kwargs):
        opened.append("opened")
        raise ValueError("opened")

    monkeypatch.setattr(pd, "read_excel", refuse_to_open)
    monkeypatch.setattr(openpyxl, "load_workbook", refuse_to_open)
    page = {"X-AP-Coder": "1"}
    answers = [
        post(client, f"/statements/emailed/{st['id']}/{st['n']}", {"vendor": ""}),
        client.post("/statements/upload", data={"vendor": ""}, files={"file": ("s.xlsx", bomb)}, headers=page),
        client.post("/settings/jde/gl-map/import", files={"file": ("map.xlsx", bomb)}, headers=page),
        client.post("/learning/history/upload", files={"file": ("ap.xlsx", bomb)}, headers=page),
        client.post("/settings/jde/worksheet", files={"file": ("answers.xlsx", bomb)}, headers=page),
    ]
    assert opened == []
    assert [a.status_code for a in answers] == [400] * 5, [a.text for a in answers]
    assert all("unzips to" in a.json()["error"] for a in answers), [a.json()["error"] for a in answers]


def test_an_emailed_statement_larger_than_an_upload_is_not_read(tmp_path, monkeypatch):
    """W2-01: the Statements page read an emailed statement whole, with no size cap; an upload is capped at 50 MB."""
    from ap_coder.web.pages import statements as statements_page

    store = Store(tmp_path / "private" / "ap_coder.db")
    load_sample_setup(store)
    client = client_for(store)
    from ap_coder.web.core import invoice_dir

    folder = invoice_dir(store)
    folder.mkdir(parents=True, exist_ok=True)
    _eml(folder / "soa.eml", "Statement", "attached", ("soa.csv", b"Statement of account\nInvoice,Amount\nINV-1,10\n"))
    unpack(folder / "soa.eml", folder)
    (st,) = inbox.statements(store, folder)
    monkeypatch.setattr(statements_page, "_MAX_BYTES", 10)
    answer = post(client, f"/statements/emailed/{st['id']}/{st['n']}", {"vendor": ""})
    assert answer.status_code == 400 and "larger than" in answer.json()["error"]
