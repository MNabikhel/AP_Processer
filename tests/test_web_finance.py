"""The web dashboard's close & compliance pages: the Invoice directory, Exports, Month-end, Sales tax, Vendor
statements.

Each page draws on an empty and on a busy database (sample setup + demo invoices); every download is the file the
earlier dashboard gave (the same ``exports`` / ``jde`` / ``stamp`` / ``accruals`` / ``taxreturn`` / ``statements``
calls);
uploads go through the column check; destructive actions ask first."""

from __future__ import annotations

import csv
import datetime as dt
import io
import zipfile
from dataclasses import replace

import openpyxl
import pytest
from fastapi.testclient import TestClient

from ap_coder.coding import taxreturn
from ap_coder.core.config import Settings
from ap_coder.core.store import APPROVED, PARKED, REJECTED, REVIEW, Store, load_sample_setup
from ap_coder.erp import accruals, dupaudit, export_layout, exports, jde, stamp
from ap_coder.erp import statements as stm
from ap_coder.ops.demo import load_demo
from ap_coder.web.app import create_app
from ap_coder.web.pages import exports as export_page

PAGE = {"X-AP-Coder": "1"}
FINANCE_PAGES = [
    "/search",
    "/search?q=northwind",
    "/exports",
    "/exports?fmt=jde",
    "/exports?fmt=custom",
    "/exports?audit=1",
    "/month-end",
    "/sales-tax",
    "/sales-tax?show=all",
    "/statements",
]


@pytest.fixture
def store(tmp_path):
    return Store(tmp_path / "private" / "ap_coder.db")


@pytest.fixture
def client(store):
    s = Settings()
    app = create_app(replace(s, llm=replace(s.llm, provider="off")), store, start_reader=False)
    return TestClient(app, base_url="http://localhost")


@pytest.fixture
def busy(store):
    load_sample_setup(store)
    load_demo(store, Settings())
    return store


def follow(client, response):
    """The page a POST's answer leads to (its message included)."""
    assert response.status_code == 200, response.text
    body = response.json()
    assert body["ok"], body
    return client.get(body["redirect"])


def attachment(response, name_part: str, mime: str) -> bytes:
    assert response.status_code == 200, response.text[:500]
    assert response.headers["content-type"].startswith(mime)
    assert (
        "attachment" in response.headers["content-disposition"] and name_part in response.headers["content-disposition"]
    )
    assert response.content
    return response.content


# --- Every page draws ------------------------------------------------------------------------------------------


@pytest.mark.parametrize("url", FINANCE_PAGES)
def test_pages_draw_on_an_empty_database(client, url):
    response = client.get(url)
    assert response.status_code == 200
    assert "is on its way" not in response.text  # not the placeholder any more
    assert "/static/data.css" in response.text


@pytest.mark.parametrize("url", FINANCE_PAGES)
def test_pages_draw_on_a_busy_database(client, busy, url):
    response = client.get(url)
    assert response.status_code == 200
    assert "is on its way" not in response.text


def test_empty_pages_say_so(client):
    assert "Nothing waiting" in client.get("/exports").text
    assert "No exports yet" in client.get("/exports").text
    assert "Nothing to accrue" in client.get("/month-end").text
    assert "Nothing to claim" in client.get("/sales-tax").text
    assert "No vendors yet" in client.get("/statements").text
    assert "Not found" in client.get("/search?q=zzz-nothing").text


# --- Invoice directory -------------------------------------------------------------------------------------------


def test_search_finds_an_invoice_and_says_where_it_is(client, busy):
    approved = busy.list_invoices(APPROVED)[0]
    text = client.get("/search", params={"q": approved["invoice_number"]}).text
    assert approved["vendor_name"].replace("&", "&amp;") in text and "Approved" in text
    assert f'href="/search/{approved["id"]}/approved.pdf"' in text
    pending = busy.list_invoices(REVIEW)[0]
    text = client.get("/search", params={"q": pending["invoice_number"]}).text
    assert f'href="/review/{pending["id"]}"' in text


def test_search_downloads_the_stamped_pdf_only_when_fully_approved(client, busy):
    approved = busy.list_invoices(APPROVED)[0]
    data = attachment(client.get(f"/search/{approved['id']}/approved.pdf"), "approved.pdf", "application/pdf")
    assert data.startswith(b"%PDF")
    assert (
        stamp.file_name(busy.get_invoice(approved["id"])).split(" ")[0]
        in client.get(f"/search/{approved['id']}/approved.pdf").headers["content-disposition"]
    )
    pending = busy.list_invoices(REVIEW)[0]
    assert client.get(f"/search/{pending['id']}/approved.pdf").status_code == 409
    assert client.get("/search/9999/approved.pdf").status_code == 404


def test_search_brings_a_parked_invoice_back(client, busy):
    inv = busy.list_invoices(REVIEW)[0]
    busy.park_invoice(inv["id"], "tester", "Waiting for the PO")
    text = client.get("/search", params={"q": inv["invoice_number"]}).text
    assert "Parked" in text and f'action="/search/{inv["id"]}/unpark"' in text
    assert client.post(f"/search/{inv['id']}/unpark").status_code == 403  # not from AP Coder's page
    page = follow(client, client.post(f"/search/{inv['id']}/unpark", headers=PAGE))
    assert page.url.path == f"/review/{inv['id']}"
    assert busy.get_invoice(inv["id"])["status"] == REVIEW
    assert busy.get_invoice(inv["id"])["status"] != PARKED


def test_bring_back_of_an_invoice_rejected_meanwhile_says_it_changed(client, busy):
    """B-194: "Bring back" on the directory swallowed the store's refusal and said "Back in the review queue." for
    an invoice rejected in another tab after the directory was drawn."""
    inv = busy.list_invoices(REVIEW)[0]
    busy.park_invoice(inv["id"], "Pat", "waiting for the buyer")
    busy.reject_invoice(inv["id"], "Sam", "not ours")
    page = follow(client, client.post(f"/search/{inv['id']}/unpark", headers=PAGE)).text
    assert "Back in the review queue." not in page and "changed meanwhile" in page and "rejected" in page
    assert busy.get_invoice(inv["id"])["status"] == REJECTED
    other = busy.list_invoices(REVIEW)[0]  # brought back meanwhile by someone else: it is in the queue, as asked
    busy.park_invoice(other["id"], "Pat", "waiting")
    busy.unpark_invoice(other["id"], "Sam")
    page = follow(client, client.post(f"/search/{other['id']}/unpark", headers=PAGE)).text
    assert "Back in the review queue." in page


def test_search_escapes_what_an_invoice_says(client, store, tmp_path):
    pdf = tmp_path / "a.pdf"
    pdf.write_bytes(b"%PDF-1.4")
    store.add_invoice(
        pdf, {"vendor_name": "<script>alert(1)</script> Ltd", "invoice_number": "X-1", "grand_total": 5}, {"issues": []}
    )
    text = client.get("/search?q=ltd").text
    assert "<script>alert(1)</script>" not in text and "&lt;script&gt;alert(1)&lt;/script&gt; Ltd" in text


# --- Exports ----------------------------------------------------------------------------------------------------


def test_the_ready_list_and_a_new_batch(client, busy):
    ready = busy.unexported_approved()
    text = client.get("/exports").text
    assert "Ready to export" in text and f"Export {len(ready)} invoices" in text
    for inv in ready:
        assert f'name="ids" value="{inv["id"]}"' in text
    assert client.post("/exports/create", data={"fmt": "xlsx", "ids": [ready[0]["id"]]}).status_code == 403
    assert client.post("/exports/create", data={"fmt": "nope", "ids": [1]}, headers=PAGE).status_code == 400
    assert client.post("/exports/create", data={"fmt": "xlsx"}, headers=PAGE).status_code == 400
    page = follow(client, client.post("/exports/create", data={"fmt": "xlsx", "ids": [ready[0]["id"]]}, headers=PAGE))
    batch = busy.export_batches()[0]["id"]
    assert f"Batch {batch} created with 1 invoice" in page.text and f"Batch {batch} is ready" in page.text
    assert f'href="/exports/batches/{batch}/download?fmt=xlsx"' in page.text
    assert busy.batch_invoice_ids(batch) == [ready[0]["id"]]
    # the same invoice cannot go out twice
    again = client.post("/exports/create", data={"fmt": "xlsx", "ids": [ready[0]["id"]]}, headers=PAGE)
    assert again.status_code == 200 and "exported meanwhile" in follow(client, again).text
    assert len(busy.export_batches()) == 1


def test_an_export_refused_as_exported_meanwhile_refreshes_the_page(client, busy):
    """It says "The page is refreshed": the answer goes to the Exports page again (in the same format), where the
    invoices exported meanwhile are no longer listed."""
    ready = [r["id"] for r in busy.unexported_approved()]
    busy.create_export_batch(ready, "csv", actor="Other tab")
    answer = client.post("/exports/create", data={"fmt": "csv", "ids": ready}, headers=PAGE).json()
    assert answer["ok"] and answer["redirect"].startswith("/exports?fmt=csv")
    page = client.get(answer["redirect"]).text
    assert "exported meanwhile" in page and "The page is refreshed" in page
    assert not any(f'name="ids" value="{i}"' in page for i in ready)


def _batch(client, store, fmt="xlsx"):
    ids = [i["id"] for i in store.unexported_approved()]
    follow(client, client.post("/exports/create", data={"fmt": fmt, "ids": ids}, headers=PAGE))
    return store.export_batches()[0]["id"]


def test_batch_downloads_are_the_export_functions_files(client, busy):
    batch = _batch(client, busy)
    invoices = [busy.get_invoice(i) for i in busy.batch_invoice_ids(batch)]
    names, ids = export_page.gl_names(busy), busy.vendor_ids()
    # CSV and custom CSV: byte for byte
    data, name, mime = exports.build("csv", invoices, names, batch, ids)
    assert attachment(client.get(f"/exports/batches/{batch}/download?fmt=csv"), name, "text/csv") == data
    data, name, _ = exports.build("custom", invoices, names, batch, ids, layout=export_page.layout(busy))
    assert attachment(client.get(f"/exports/batches/{batch}/download?fmt=custom"), name, "text/csv") == data
    # Excel: the same cells (the workbook's own creation time differs by the second)
    data, name, mime = exports.build("xlsx", invoices, names, batch, ids)
    got = attachment(client.get(f"/exports/batches/{batch}/download"), name, mime)

    def cells(raw: bytes) -> dict[str, list]:
        book = openpyxl.load_workbook(io.BytesIO(raw))
        return {ws.title: [list(r) for r in ws.iter_rows(values_only=True)] for ws in book.worksheets}

    assert cells(got) == cells(data)
    assert client.get("/exports/batches/99/download").status_code == 404
    assert client.get(f"/exports/batches/{batch}/download?fmt=pdf").status_code == 400


def test_jde_download_is_the_zip_jde_builds(client, busy):
    batch = _batch(client, busy)
    data, name, mime = export_page.jde_file(busy, batch)
    got = attachment(client.get(f"/exports/batches/{batch}/download?fmt=jde"), name, mime)
    with zipfile.ZipFile(io.BytesIO(got)) as a, zipfile.ZipFile(io.BytesIO(data)) as b:
        assert a.namelist() == b.namelist() and "README.txt" in a.namelist()
        assert all(a.read(n) == b.read(n) for n in a.namelist())
    text = client.get(f"/exports?batch={batch}&bfmt=jde").text
    invoices, left_out = export_page.jde_left_out(busy, batch)
    if left_out:
        assert "cannot go to JD Edwards" in text


def test_the_jde_check_leaves_out_what_cannot_go(client, busy):
    ok, problems, blocking = export_page.jde_check(busy, [i["id"] for i in busy.unexported_approved()])
    text = client.get("/exports?fmt=jde").text
    assert jde.LABEL.split(" (")[0] in text or "JD Edwards E1" in text
    for p in blocking:
        assert p.split(".")[0] in text
    if blocking or not ok:
        answer = client.post(
            "/exports/create", data={"fmt": "jde", "ids": [i["id"] for i in busy.unexported_approved()]}, headers=PAGE
        )
        assert answer.status_code == 400
    for inv_id in problems:
        assert f'name="ids" value="{inv_id}" class="export-pick"' in text
        assert "Why not" in text


def test_approved_pdfs_zip_and_undo_with_a_confirmation(client, busy):
    batch = _batch(client, busy)
    text = client.get("/exports").text
    assert f'action="/exports/batches/{batch}/undo" data-confirm="Undo batch {batch}?' in text
    data = attachment(client.get(f"/exports/batches/{batch}/approved-pdfs.zip"), "approved_pdfs.zip", "application/zip")
    expected = stamp.batch_zip([busy.get_invoice(i) for i in busy.batch_invoice_ids(batch)], export_page.gl_names(busy))
    with zipfile.ZipFile(io.BytesIO(data)) as a, zipfile.ZipFile(io.BytesIO(expected)) as b:
        assert a.namelist() == b.namelist() and a.namelist()
    count = len(busy.batch_invoice_ids(batch))
    page = follow(client, client.post(f"/exports/batches/{batch}/undo", headers=PAGE))
    assert f"Batch {batch} undone: {count} invoice" in page.text
    assert busy.batch_invoice_ids(batch) == [] and len(busy.unexported_approved()) == count
    assert "Undone" in page.text
    assert client.post(f"/exports/batches/{batch}/undo", headers=PAGE).status_code == 400  # not twice


def test_two_undos_of_a_batch_at_once_record_one_undo(client, busy, monkeypatch):
    """Two clicks on Undo at the same moment both found the batch not undone yet: the second wrote a second
    "export undone" event (0 invoices) to the audit trail and took over the batch's "undone by"."""
    batch = _batch(client, busy)
    before = busy.export_batches()  # as both requests read it, before either undid the batch
    follow(client, client.post(f"/exports/batches/{batch}/undo", headers=PAGE))
    monkeypatch.setattr(busy, "export_batches", lambda *a, **k: before)
    late = client.post(f"/exports/batches/{batch}/undo", headers=PAGE)
    assert late.status_code == 400 and "not an export that can be undone" in late.json()["error"]
    assert len(busy.events(actions=["export_undone"])) == 1


def test_an_undone_batch_cannot_be_downloaded(client, busy):
    batch = _batch(client, busy)
    follow(client, client.post(f"/exports/batches/{batch}/undo", headers=PAGE))
    for url in (f"/exports/batches/{batch}/download", f"/exports/batches/{batch}/download?fmt=jde",
                f"/exports/batches/{batch}/approved-pdfs.zip"):  # fmt: skip
        answer = client.get(url)
        assert answer.status_code == 410 and "undone" in answer.text


def test_the_custom_layout_is_saved_and_previewed(client, busy):
    form = {
        "header": ["Vendor", "Amount", "Company", "Half"],
        "field": ["vendor_name", "amount", export_layout.FIXED, ""],
        "text": ["", "", "ACME01", ""],
        "date_format": "YYYYMMDD",
        "delimiter": ";",
        "decimal_comma": "1",
        "header_row": "1",
    }
    preview = client.post("/exports/layout/preview", data=form, headers=PAGE).json()
    assert preview["ok"] and preview["preview"].startswith("Vendor;Amount;Company")
    assert any("row 4 (Half)" in w for w in preview["warnings"])
    page = follow(client, client.post("/exports/layout", data=form, headers=PAGE))
    assert "Export layout saved" in page.text and "row 4 (Half)" in page.text
    layout = export_page.layout(busy)
    assert [c.header for c in layout.columns] == ["Vendor", "Amount", "Company"]
    assert layout.delimiter == ";" and layout.decimal_comma and layout.date_format == "YYYYMMDD"
    custom = client.get("/exports?fmt=custom").text
    assert 'value="ACME01"' in custom and "Macro" not in custom
    assert '<pre class="code" id="layout-preview">Vendor;Amount;Company' in custom
    empty = client.post("/exports/layout", data={"header": [""], "field": [""], "text": [""]}, headers=PAGE)
    assert empty.status_code == 400


def test_the_custom_layout_keeps_or_leaves_out_the_byte_order_mark(client, busy):
    """B-224: a checkbox on the layout, ticked unless it was unticked, says whether the file starts with a BOM."""
    assert 'name="byte_order_mark" value="1" checked' in client.get("/exports?fmt=custom").text
    form = {"header": ["Vendor"], "field": ["vendor_name"], "text": [""], "header_row": "1"}
    follow(client, client.post("/exports/layout", data=form, headers=PAGE))
    assert export_page.layout(busy).byte_order_mark is False
    assert 'name="byte_order_mark" value="1">' in client.get("/exports?fmt=custom").text
    follow(client, client.post("/exports/layout", data={**form, "byte_order_mark": "1"}, headers=PAGE))
    assert export_page.layout(busy).byte_order_mark is True


def test_the_erp_register_is_uploaded_checked_imported_and_cleared(client, busy):
    sheet = "Supplier,Doc No,Doc Date,Amount\nAcme Ltd,A-1,2026-01-05,100.00\nBeta Inc,B-7,2026-02-01,55.10\n"
    bad = client.post(
        "/exports/register/upload", files={"file": ("list.pdf", b"%PDF", "application/pdf")}, headers=PAGE
    )
    assert bad.status_code == 400 and "CSV or Excel" in bad.json()["error"]
    page = follow(
        client,
        client.post("/exports/register/upload", files={"file": ("erp.csv", sheet.encode(), "text/csv")}, headers=PAGE),
    )
    assert "erp.csv" in page.text and "Doc No" in page.text and "A-1" in page.text
    token = page.url.params["register"]
    missing = client.post("/exports/register/import", data={"file": token, "col_vendor_name": "Supplier"}, headers=PAGE)
    assert missing.status_code == 400 and "Invoice number" in missing.json()["error"]
    page = follow(
        client,
        client.post(
            "/exports/register/import",
            data={
                "file": token,
                "col_vendor_name": "Supplier",
                "col_invoice_number": "Doc No",
                "col_invoice_date": "Doc Date",
                "col_total": "Amount",
            },
            headers=PAGE,
        ),
    )
    assert "2 ERP invoices known" in page.text and busy.erp_register_count() == 2
    assert 'action="/exports/register/clear" data-confirm=' in page.text
    follow(client, client.post("/exports/register/clear", headers=PAGE))
    assert busy.erp_register_count() == 0
    gone = client.post("/exports/register/import", data={"file": token}, headers=PAGE)
    assert gone.status_code == 400 and "Upload the file again" in gone.json()["error"]


def test_a_bill_entered_twice_in_the_erp_is_named_when_the_register_is_imported(client, busy):
    sheet = "Supplier,Invoice,Date,Amount\nAcme Ltd,A-1,2026-01-05,100.00\nAcme Ltd,A-1,2026-02-05,100.00\n"
    page = follow(
        client,
        client.post("/exports/register/upload", files={"file": ("erp.csv", sheet.encode(), "text/csv")}, headers=PAGE),
    )
    data = {"file": page.url.params["register"], "col_vendor_name": "Supplier", "col_invoice_number": "Invoice",
            "col_invoice_date": "Date", "col_total": "Amount"}  # fmt: skip
    page = follow(client, client.post("/exports/register/import", data=data, headers=PAGE))
    assert "1 ERP invoices known" in page.text
    assert "1 bill entered more than once in the ERP: see the duplicate payment audit." in page.text


def test_the_duplicate_audit_and_its_csv(client, busy):
    text = client.get("/exports?audit=1").text
    pairs = dupaudit.find(busy)
    assert "Duplicate payment audit" in text
    assert ("No likely duplicates found" in text) == (not pairs)
    data = attachment(client.get("/exports/duplicates.csv"), "duplicate_audit_", "text/csv")
    assert data == dupaudit.to_csv(pairs)


# --- Month-end and sales tax -----------------------------------------------------------------------------------


def test_month_end_shows_the_accruals_and_downloads_them(client, busy):
    end = dt.date.today() + dt.timedelta(days=60)
    items = accruals.build(busy, end)
    text = client.get(f"/month-end?end={end}").text
    assert f'value="{end}"' in text
    if items:
        assert "Accruals schedule" in text and "By GL account" in text
    data = attachment(client.get(f"/month-end/accruals.csv?end={end}"), f"accruals_{end}.csv", "text/csv")
    assert data == accruals.to_csv(items, end)
    assert client.get("/month-end?end=not-a-date").status_code == 200  # the default period instead


def test_sales_tax_claims_and_self_assessment_downloads(client, busy):
    start, end = dt.date(2020, 1, 1), dt.date.today()
    report = taxreturn.build(busy, start, end)
    text = client.get(f"/sales-tax?start={start}&end={end}&show=all").text
    assert "GST/HST to claim" in text
    if report.claims:
        assert "By tax and rate" in text
    data = attachment(
        client.get(f"/sales-tax/claims.csv?start={start}&end={end}"), f"sales_tax_{start}_{end}.csv", "text/csv"
    )
    assert data == taxreturn.to_csv(report)
    items = taxreturn.self_assessment(busy, start, end)
    data = attachment(
        client.get(f"/sales-tax/self-assessment.csv?start={start}&end={end}"), "self_assessment_", "text/csv"
    )
    assert data == taxreturn.self_assessment_csv(items, start, end)
    backwards = client.get(f"/sales-tax?start={end}&end={start}").text
    assert "The start date is after the end date." in backwards
    assert client.get(f"/sales-tax/claims.csv?start={end}&end={start}").status_code == 400


# --- Vendor statements -----------------------------------------------------------------------------------------


def _statement(store: Store) -> tuple[str, bytes]:
    vendor = store.vendor_summaries()[0]
    invoices = store.vendor_invoices(vendor["vendor_key"])
    out = io.StringIO()
    w = csv.writer(out)
    w.writerow(["Invoice No", "Date", "Amount"])
    for inv in invoices:
        w.writerow([inv["invoice_number"], inv["invoice_date"], f"{inv['grand_total']:.2f}"])
    w.writerow(["MISSING-<b>1</b>", "2026-09-01", "123.45"])
    return vendor["vendor_key"], out.getvalue().encode()


def test_a_statement_is_reconciled_and_downloaded(client, busy):
    key, content = _statement(busy)
    page = follow(
        client,
        client.post(
            "/statements/upload",
            data={"vendor": key},
            files={"file": ("statement.csv", content, "text/csv")},
            headers=PAGE,
        ),
    )
    text = page.text
    assert "statement.csv" in text and "Matched" in text and "Not received" in text
    assert "MISSING-&lt;b&gt;1&lt;/b&gt;" in text and "MISSING-<b>1</b>" not in text  # the vendor's text is escaped
    assert "Ask the vendor for the missing invoices" in text and 'href="mailto:' in text
    query = page.url.query.decode()
    data = attachment(client.get(f"/statements/reconciliation.csv?{query}"), "statement_reconciliation.csv", "text/csv")
    records = list(csv.DictReader(io.StringIO(content.decode())))
    columns = stm.map_columns(["Invoice No", "Date", "Amount"])
    assert data == stm.to_csv(stm.reconcile(records, columns, busy.vendor_invoices(key)))
    french = client.get(f"/statements?{query}&lang=fr").text
    assert "Votre relevé de compte" in french


def test_statement_columns_can_be_picked(client, busy):
    key, content = _statement(busy)
    page = follow(
        client,
        client.post(
            "/statements/upload", data={"vendor": key}, files={"file": ("s.csv", content, "text/csv")}, headers=PAGE
        ),
    )
    token = page.url.params["file"]
    unpicked = client.get("/statements", params={"vendor": key, "file": token, "col_number": "", "col_amount": ""})
    assert "Pick the invoice number column" in unpicked.text and "Matched" not in unpicked.text
    gone = client.get("/statements", params={"vendor": key, "file": "nope"})
    assert "Upload the file again" in gone.text


def test_a_statement_from_excel_with_two_sheets(client, busy):
    key, content = _statement(busy)
    book = openpyxl.Workbook()
    book.active.title = "Notes"
    book.active.append(["Nothing here"])
    sheet = book.create_sheet("Statement")
    for row in csv.reader(io.StringIO(content.decode())):
        sheet.append(row)
    raw = io.BytesIO()
    book.save(raw)
    page = follow(
        client,
        client.post(
            "/statements/upload",
            data={"vendor": key},
            files={"file": ("s.xlsx", raw.getvalue(), "application/octet-stream")},
            headers=PAGE,
        ),
    )
    assert 'name="sheet" data-autosubmit' in page.text and "Nothing here" in page.text
    token = page.url.params["file"]
    text = client.get("/statements", params={"vendor": key, "file": token, "sheet": "Statement"}).text
    assert "Matched" in text
