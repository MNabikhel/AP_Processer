"""Edge cases found in code review: real-world files, repeated line numbers, large charts of accounts."""

import copy
import datetime as dt
import json
from dataclasses import replace

from openpyxl import load_workbook

from ap_coder import cli, doctor
from ap_coder.config import DocumentIntelligenceSettings, OpenAISettings, Settings
from ap_coder.evaluation import _is_prediction_file
from ap_coder.labels import _prediction_files, export_labels
from ap_coder.memory import CORRECTED, pair_lines
from ap_coder.pipeline import discover_inputs
from ap_coder.reference_data import load_reference_data
from ap_coder.schema import (
    FIXED_ENUM_VALUES,
    MAX_ENUM_VALUES_TOTAL,
    build_json_schema,
    plan_code_enums,
)
from ap_coder.store import Store
from ap_coder.tax import TaxRateTable, load_tax_mapping

from .conftest import DATA


def _line_props(schema):
    return schema["properties"]["line_items"]["items"]["properties"]


def _enum_count(node) -> int:
    if isinstance(node, dict):
        return len(node.get("enum", [])) + sum(_enum_count(v) for v in node.values())
    if isinstance(node, list):
        return sum(_enum_count(v) for v in node)
    return 0


# --- Learning: repeated line numbers --------------------------------------------------------------


def test_pair_lines_matches_repeated_line_numbers_in_order():
    before = [{"line_number": 1, "x": "a"}, {"line_number": 2, "x": "b"}, {"line_number": 2, "x": "c"}]
    after = [{"line_number": 2, "x": "B"}, {"line_number": 2, "x": "C"}, {"line_number": 3, "x": "new"}]
    pairs = pair_lines(before, after)
    assert [(b and b["x"], a["x"]) for b, a in pairs] == [("b", "B"), ("c", "C"), (None, "new")]


def test_approve_with_repeated_line_numbers_compares_the_right_lines(tmp_path, ground_truth):
    ai = copy.deepcopy(ground_truth)
    for li in ai["line_items"][:2]:
        li["line_number"] = 1  # invoice prints "1" twice
    store = Store(tmp_path / "ap.db")
    inv_id = store.add_invoice(tmp_path / "x.pdf", ai, {"requires_review": True})
    final = copy.deepcopy(ai)
    final["line_items"][1]["predicted_gl_code"] = "6010"  # only the second "1" is corrected
    store.approve_invoice(inv_id, final, reviewer="ap.clerk")
    corrected = [r for r in store.feedback_rows() if r["outcome"] == CORRECTED]
    assert [(r["suggested_gl"], r["final_gl"]) for r in corrected] == [
        (ai["line_items"][1]["predicted_gl_code"], "6010")
    ]


# --- Structured output enum budget ------------------------------------------------------------------


def test_enum_budget_is_shared_across_the_schema():
    budget = MAX_ENUM_VALUES_TOTAL - FIXED_ENUM_VALUES
    gl = [str(5000 + i) for i in range(300)]
    cc = [f"CC{i}" for i in range(300)]
    gl_enum, cc_enum = plan_code_enums(gl, cc)
    assert gl_enum == [*gl, "UNASSIGNED"] and cc_enum is None  # GL wins, cost centers become free text
    assert plan_code_enums(gl[:100], cc[:100]) == ([*gl[:100], "UNASSIGNED"], [*cc[:100], "UNASSIGNED"])
    huge = [str(i) for i in range(budget + 1)]
    assert plan_code_enums(huge, None) == (None, None)


def test_schema_never_exceeds_the_enum_limit(reference):
    rows = [{"gl_code": str(5000 + i), "description": f"Account {i}"} for i in range(400)]
    big = replace(
        reference,
        chart_of_accounts=replace(reference.chart_of_accounts, rows=rows),
        cost_centers=replace(reference.cost_centers, rows=[{"cost_center": f"CC{i}"} for i in range(200)]),
    )
    schema = build_json_schema(big)
    assert _enum_count(schema) <= MAX_ENUM_VALUES_TOTAL
    assert "enum" in _line_props(schema)["predicted_gl_code"]
    assert "enum" not in _line_props(schema)["predicted_cost_center"]


# --- Doctor report redaction ------------------------------------------------------------------------


def test_doctor_scrubs_endpoint_hostnames_and_resource_names():
    settings = Settings(
        document_intelligence=DocumentIntelligenceSettings(endpoint="https://contoso-di.cognitiveservices.azure.com/"),
        openai=OpenAISettings(endpoint="https://contoso-aoai.openai.azure.com/", api_key="sk-secret-value-123"),
    )
    msg = (
        "Failed to resolve CONTOSO-AOAI.openai.azure.com; resource contoso-di not found; "
        "key sk-secret-value-123 rejected"
    )
    scrubbed = doctor._scrub(msg, settings)
    for secret in ("contoso", "sk-secret"):
        assert secret not in scrubbed.lower()
    assert "<redacted>" in scrubbed


# --- Files exported from Excel / ERPs ------------------------------------------------------------------


def test_csv_saved_as_windows_1252_loads(tmp_path):
    coa = tmp_path / "coa.csv"
    coa.write_bytes("Account,Description\n6100,Fournitures de bureau – générales\n".encode("cp1252"))
    mapping = tmp_path / "tax.csv"
    mapping.write_bytes("tax_type,treatment,gl_code\nQST,recoverable,2320\n".encode("cp1252"))
    ref = load_reference_data(coa, tax_mapping=mapping)
    assert "Fournitures de bureau – générales" in ref.chart_of_accounts.rows[0].values()
    assert load_tax_mapping(mapping)["QST"].gl_code == "2320"


def test_labels_workbook_without_cost_centers(tmp_path, ground_truth):
    out = tmp_path / "out"
    out.mkdir()
    (out / "inv.json").write_text(json.dumps(ground_truth))
    reference = load_reference_data(DATA / "chart_of_accounts.csv")
    assert export_labels(out, tmp_path / "l.xlsx", reference=reference) == 1
    assert "Cost Centers" not in load_workbook(tmp_path / "l.xlsx").sheetnames


def test_extraction_dumps_are_not_predictions(tmp_path, ground_truth):
    out = tmp_path / "out"
    out.mkdir()
    for name in ("inv.json", "inv.di.json", "inv.validation.json", "batch_summary.json"):
        (out / name).write_text(json.dumps(ground_truth))
    assert [p.name for p in _prediction_files(out)] == ["inv.json"]
    assert [p.name for p in sorted(out.iterdir()) if _is_prediction_file(p)] == ["inv.json"]


def test_text_copy_next_to_a_pdf_is_not_processed_twice(tmp_path):
    for name in ("a.pdf", "a.md", "b.md", "c.PNG", "c.txt", "notes.docx", "README.md"):
        (tmp_path / name).write_bytes(b"x")
    assert [p.name for p in discover_inputs([tmp_path])] == ["a.pdf", "b.md", "c.PNG"]


# --- Tax rates over time -------------------------------------------------------------------------------


def test_historic_hst_rates():
    rates = TaxRateTable.load()
    assert rates.rate_for("HST", "NB", dt.date(2009, 5, 1)) == 0.13
    assert rates.rate_for("HST", "NL", dt.date(2012, 5, 1)) == 0.13
    assert rates.rate_for("HST", "NS", dt.date(2009, 5, 1)) == 0.13
    assert rates.rate_for("HST", "NS", dt.date(2011, 5, 1)) == 0.15
    assert rates.rate_for("HST", "NS", dt.date(2025, 5, 1)) == 0.14


# --- Command line: explicit files on top of the dashboard database -----------------------------------------


def test_cli_explicit_files_override_the_database(tmp_path, capsys):
    db = tmp_path / "ap.db"
    Store(db).import_accounts("gl_accounts", [{"code": "6100", "name": "Office"}], "code", "name")
    cc = tmp_path / "cc.csv"
    cc.write_text("cost_center,name\nZ1,Zed\n")
    assert cli.main(["--db", str(db), "schema", "--cost-centers", str(cc)]) == 0
    line = _line_props(json.loads(capsys.readouterr().out))
    assert line["predicted_gl_code"]["enum"] == ["6100", "UNASSIGNED"]
    assert line["predicted_cost_center"]["enum"] == ["Z1", "UNASSIGNED"]


# --- Approvals never learn "UNASSIGNED" -------------------------------------------------------------------


def test_unassigned_lines_are_not_learned(tmp_path, ground_truth):
    store = Store(tmp_path / "ap.db")
    inv_id = store.add_invoice(tmp_path / "x.pdf", ground_truth, {"requires_review": True})
    final = copy.deepcopy(ground_truth)
    final["line_items"][0]["predicted_gl_code"] = "UNASSIGNED"
    store.approve_invoice(inv_id, final, reviewer="ap.clerk")
    learned = {r["line_number"] for r in store.feedback_rows()}
    assert 1 not in learned and len(learned) == len(final["line_items"]) - 1


def test_added_rows_get_unassigned_cost_center_when_configured(ground_truth):
    import pandas as pd

    from ap_coder.review import coding_from_inputs

    header = {k: ground_truth[k] for k in ground_truth if k not in ("line_items", "tax_lines", "confidence_score")}
    lines = pd.concat(
        [pd.DataFrame(ground_truth["line_items"]), pd.DataFrame([{"description": "Eco fee", "amount": 5}])],
        ignore_index=True,
    )
    coding, problems = coding_from_inputs(
        header, lines, pd.DataFrame(ground_truth["tax_lines"]), ground_truth, "UNASSIGNED"
    )
    assert problems == []
    assert coding.line_items[-1].predicted_cost_center == "UNASSIGNED"
    assert coding.line_items[0].predicted_cost_center == ground_truth["line_items"][0]["predicted_cost_center"]
