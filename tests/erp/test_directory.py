"""The invoice directory (directory.py, /search): filters over each column, sort, pages, columns, layouts, export."""

from __future__ import annotations

import csv
import io
import json
from urllib.parse import parse_qsl

import pytest
from openpyxl import load_workbook

from ap_coder.core.store import APPROVED, PARKED, REVIEW, Store
from ap_coder.erp import directory as d
from tests.conftest import SAMPLE_STEM, SAMPLES, copy_database
from tests.webclient import PAGE, client_for, follow


@pytest.fixture
def store(tmp_path, setup_demo_template):
    return copy_database(setup_demo_template, tmp_path / "private" / "ap.db")  # sample setup, then demo


def ids(store, **filters):
    return sorted(r["id"] for r in d.run(store, d.Query(filters=filters)).rows)


def every(store, column):
    return [d.cell(r, column) for r in d.run(store, d.Query(), everything=True).rows]


# --- Filters ----------------------------------------------------------------------------------------------------------


def test_text_filters_contain_start_end_exact_and_not(store):
    vendors = every(store, "vendor")
    assert len(ids(store, vendor="north")) == sum("north" in v.lower() for v in vendors)
    assert len(ids(store, vendor="north*")) == sum(v.lower().startswith("north") for v in vendors)
    assert len(ids(store, vendor="*inc.")) == sum(v.lower().endswith("inc.") for v in vendors)
    assert ids(store, vendor="=northwind it solutions inc.") == ids(store, vendor="north*")
    assert len(ids(store, vendor="!north")) == len(vendors) - len(ids(store, vendor="north"))
    assert ids(store, invoice_number="nw-2026-091?") == ids(store, invoice_number="=NW-2026-0912")


def test_wildcards_never_reach_sql_as_its_own(store, tmp_path):
    pdf = tmp_path / "x.pdf"
    pdf.write_bytes(b"%PDF-1.4")
    a = store.add_invoice(pdf, {"vendor_name": "100% Pure_Water Ltd", "invoice_number": "A_1", "grand_total": 5},
                          {"issues": []})  # fmt: skip
    store.add_invoice(pdf, {"vendor_name": "100X PureXWater Ltd", "invoice_number": "AB1", "grand_total": 5},
                      {"issues": []})  # fmt: skip
    assert ids(store, vendor="100%") == [a]  # % and _ are letters here, not SQL's wildcards
    assert ids(store, invoice_number="=a_1") == [a]
    assert ids(store, vendor="'; DROP TABLE invoices; --") == []
    assert store.get_invoice(a)


def test_text_filters_ignore_accents_and_case_beyond_ascii(store, tmp_path):
    """SQLite's lower() and LIKE fold only A to Z: "énergir" never found "Énergir S.E.C."."""
    pdf = tmp_path / "x.pdf"
    pdf.write_bytes(b"%PDF-1.4")
    a = store.add_invoice(pdf, {"vendor_name": "Énergir S.E.C.", "invoice_number": "É-100", "grand_total": 5},
                          {"issues": []})  # fmt: skip
    for typed in ("énergir", "ÉNERGIR", "energir", "=énergir s.e.c.", "=ENERGIR S.E.C.", "*gir s.e.c."):
        assert ids(store, vendor=typed) == [a], typed
    assert ids(store, invoice_number="=é-100") == [a]
    assert a not in ids(store, vendor="!énergir")


def test_amount_filters(store):
    totals = every(store, "total")
    assert len(ids(store, total=">10000")) == sum(t > 10000 for t in totals)
    assert len(ids(store, total="<=3053.12")) == sum(t <= 3053.12 for t in totals)
    assert len(ids(store, total="2000..12000")) == sum(2000 <= t <= 12000 for t in totals)
    assert len(ids(store, total="12000..2000")) == len(ids(store, total="2000..12000"))  # either order
    assert len(ids(store, total="18,017.85")) == 1 and len(ids(store, total="$18017.85")) == 1
    assert len(ids(store, total="!18017.85")) == len(totals) - 1
    assert len(ids(store, total="<0")) == 1  # the credit note


def test_date_filters_by_day_month_year_and_range(store):
    dates = every(store, "invoice_date")
    assert len(ids(store, invoice_date="2026-09")) == sum(x.startswith("2026-09") for x in dates)
    assert len(ids(store, invoice_date="2026")) == len(dates)
    assert len(ids(store, invoice_date="2026-09-14")) == sum(x == "2026-09-14" for x in dates)
    assert len(ids(store, invoice_date=">2026-09-14")) == sum(x > "2026-09-14" for x in dates)
    assert len(ids(store, invoice_date=">=2026-09-14")) == sum(x >= "2026-09-14" for x in dates)
    assert len(ids(store, invoice_date="<2026-09")) == sum(x < "2026-09-01" for x in dates)
    assert len(ids(store, invoice_date="2026-07..2026-09")) == sum("2026-07-01" <= x <= "2026-09-30" for x in dates)
    assert len(ids(store, invoice_date="!2026-09")) == sum(not x.startswith("2026-09") for x in dates)


def test_po_gl_status_and_has_or_has_not(store):
    pos = every(store, "po")
    assert len(ids(store, po="*")) == sum(bool(p) for p in pos)
    assert len(ids(store, po="!*")) == sum(not p for p in pos)
    gls = every(store, "gl")
    assert len(ids(store, gl="1550")) == sum("1550" in (g or "") for g in gls)
    assert len(ids(store, gl="15*")) == sum(any(c.startswith("15") for c in (g or "").split(", ")) for g in gls)
    assert len(ids(store, gl="!6010")) == sum("6010" not in (g or "") for g in gls)
    assert all("UNASSIGNED" not in (g or "") for g in gls)
    assert ids(store, status="approved") == sorted(i["id"] for i in store.list_invoices(APPROVED))
    assert len(ids(store, status="approved,review")) == len(gls)
    assert ids(store, batch="*") == [] and len(ids(store, batch="!*")) == len(gls)


def test_filters_that_cannot_be_read_are_explained_and_left_out(store):
    result = d.run(store, d.Query(filters={"total": "lots", "invoice_date": "yesterday", "vendor": ">north",
                                           "id": "abc", "where": "x"}))  # fmt: skip
    assert set(result.problems) == {"total", "invoice_date", "vendor", "id", "where"}
    assert "not an amount" in result.problems["total"] and "not a date" in result.problems["invoice_date"]
    assert result.total == 10  # nothing filtered


def test_quick_search_on_top_of_the_filters(store):
    assert d.run(store, d.Query(q="northwind")).total == 2
    assert d.run(store, d.Query(q="northwind", filters={"total": "<0"})).total == 1
    assert d.run(store, d.Query(q="zzz-nothing")).total == 0


def test_quick_search_with_more_hits_than_sqlite_takes_variables(store, monkeypatch, few_sql_variables):
    """Up to EXPORT_LIMIT hits: more than a stock SQLite build takes as bound variables (32,766, or 999)."""
    from ap_coder.erp import search

    real = {i["id"] for i in store.list_invoices()}
    many = [search.Hit(1, {"id": n}) for n in [*real, *range(10_000, 12_000)]]
    monkeypatch.setattr(search, "find", lambda *a, **kw: many)
    found = d.run(store, d.Query(q="anything"), everything=True)
    assert found.total == len(real) and {r["id"] for r in found.rows} == real


# --- Sort, pages, columns -------------------------------------------------------------------------------------------


def test_sorting_puts_empty_values_last_either_way(store):
    for descending in (False, True):
        pos = [d.cell(r, "po") for r in d.run(store, d.Query(sort="po", descending=descending)).rows]
        filled = [p for p in pos if p]
        assert pos[: len(filled)] == sorted(filled, key=str.lower, reverse=descending) and not any(pos[len(filled) :])
    totals = [d.cell(r, "total") for r in d.run(store, d.Query(sort="total", descending=False)).rows]
    assert totals == sorted(totals)


def test_pages_of_a_hundred(tmp_path):
    store = Store(tmp_path / "a.db")
    gt = json.loads((SAMPLES / "ground_truth" / f"{SAMPLE_STEM}.json").read_text())
    for i in range(205):
        store.add_invoice(tmp_path / f"{i}.pdf", {**gt, "invoice_number": f"N-{i}"}, {"issues": []})
    first, last = d.run(store, d.Query()), d.run(store, d.Query(page=99))
    assert (first.total, first.pages, len(first.rows), first.first, first.last) == (205, 3, 100, 1, 100)
    assert (last.page, len(last.rows), last.first, last.last) == (3, 5, 201, 205)
    seen = [r["id"] for p in (1, 2, 3) for r in d.run(store, d.Query(page=p)).rows]
    assert sorted(seen) == sorted(set(seen)) and len(seen) == 205  # no invoice twice, none missing
    assert len(d.run(store, d.Query(), everything=True).rows) == 205


def test_the_query_is_its_url():
    q = d.Query(columns=["gl", "vendor", "po"], filters={"vendor": "n*", "total": ">5"}, sort="total",
                descending=False, page=2, q="0912")  # fmt: skip
    assert d.Query.from_params(parse_qsl(q.query_string())) == q
    odd = d.Query.from_params([("col", "nope"), ("col", "vendor"), ("col", "vendor"), ("f_nope", "x"),
                               ("sort", "where"), ("page", "-4"), ("f_po", "  ")])  # fmt: skip
    assert (odd.columns, odd.filters, odd.sort, odd.page) == (["vendor"], {}, d.DEFAULT_SORT, 1)
    assert d.Query.from_params({}).columns == list(d.DEFAULT_COLUMNS)


# --- Layouts ----------------------------------------------------------------------------------------------------------


def test_layouts_saved_replaced_defaulted_and_deleted(store):
    assert d.presets(store) == [] and d.default_preset(store) is None
    with pytest.raises(ValueError):
        d.save_preset(store, "   ", "")
    d.save_preset(store, "Approved", "col=vendor&f_status=approved&page=3&evil=1", make_default=True)
    d.save_preset(store, "approved", "col=total")  # the same name, another case: replaced
    assert [(p["name"], p["query"]) for p in d.presets(store)] == [("approved", "col=total&sort=received&dir=desc")]
    assert d.default_preset(store) is None  # its default was the old spelling
    assert d.set_default(store, "approved") and d.default_preset(store)["name"] == "approved"
    assert not d.set_default(store, "nope")
    assert d.delete_preset(store, "approved") and d.presets(store) == [] and d.default_preset(store) is None
    assert not d.delete_preset(store, "approved")
    store.set_setting(d.PRESETS_SETTING, "not json")
    assert d.presets(store) == []


# --- The page ---------------------------------------------------------------------------------------------------------


def test_the_page_filters_sorts_and_shows_the_columns_chosen(store):
    client = client_for(store)
    page = client.get("/search").text
    assert "Invoice directory" in page and page.count("<tr>") >= 10 and 'name="f_vendor"' in page
    page = client.get("/search?col=gl&col=vendor&col=total&f_total=%3C0&sort=total&dir=asc").text
    assert page.index(">GL accounts<") < page.index(">Vendor<") and "Rows 1–1 of 1" in page
    assert 'aria-sort="ascending"' in page
    assert "is not a date" in client.get("/search?f_due_date=soon").text


def test_the_page_opens_with_the_default_layout(store):
    client = client_for(store)
    answer = client.post("/search/presets", data={"name": "Credits", "query": "col=vendor&col=total&f_total=<0",
                                                  "make_default": "1"}, headers=PAGE)  # fmt: skip
    assert "Layout “Credits” saved." in follow(client, answer)
    page = client.get("/search").text
    assert "Rows 1–1 of 1" in page and "Opens first" in page
    assert "Rows 1–10 of 10" in client.get("/search?sort=received").text  # its own query wins
    assert client.post("/search/presets", data={"name": "Credits"}).status_code == 403  # not from AP Coder's page
    follow(client, client.post("/search/presets/default", data={"name": ""}, headers=PAGE))
    assert "Rows 1–10 of 10" in client.get("/search").text
    assert client.post("/search/presets/default", data={"name": "gone"}, headers=PAGE).status_code == 404
    assert "deleted" in follow(client, client.post("/search/presets/delete", data={"name": "Credits"}, headers=PAGE))
    assert client.post("/search/presets/delete", data={"name": "Credits"}, headers=PAGE).status_code == 404
    assert client.post("/search/presets", data={"name": ""}, headers=PAGE).status_code == 400


def test_export_has_the_columns_shown_and_every_row_kept(store):
    client = client_for(store)
    query = "col=vendor&col=total&col=status&f_status=approved&sort=total&dir=asc"
    response = client.get(f"/search/export.csv?{query}")
    assert response.headers["content-type"].startswith("text/csv")
    rows = list(csv.reader(io.StringIO(response.content.decode("utf-8-sig"))))
    approved = sorted(store.list_invoices(APPROVED), key=lambda r: r["grand_total"])
    assert rows == [["Vendor", "Total", "Status"]] + [
        [r["vendor_name"], f"{r['grand_total']:.2f}", "Approved"] for r in approved
    ]
    book = load_workbook(io.BytesIO(client.get(f"/search/export.xlsx?{query}").content))
    sheet = book.active
    assert [c.value for c in sheet[1]] == ["Vendor", "Total", "Status"] and sheet.max_row == len(approved) + 1


def test_export_never_writes_a_formula(store, tmp_path):
    pdf = tmp_path / "x.pdf"
    pdf.write_bytes(b"%PDF-1.4")
    store.add_invoice(pdf, {"vendor_name": '=HYPERLINK("x")', "invoice_number": "1", "grand_total": 1}, {"issues": []})
    text = client_for(store).get("/search/export.csv?col=vendor&f_invoice_number==1").content.decode("utf-8-sig")
    assert "\n=HYPERLINK" not in text and "'=HYPERLINK" in text


def test_row_actions_follow_where_the_invoice_is(store):
    client = client_for(store)
    review = store.list_invoices(REVIEW)[0]
    store.park_invoice(review["id"], "tester", "Waiting for the PO")
    page = client.get("/search?col=vendor").text
    assert f'action="/search/{review["id"]}/unpark"' in page
    approved = store.list_invoices(APPROVED)[0]
    assert f'href="/search/{approved["id"]}/approved.pdf"' in page
    assert store.get_invoice(review["id"])["status"] == PARKED


def test_a_status_filter_with_more_statuses_than_sqlite_takes_variables(store, few_sql_variables):
    """The Status filter comes from the address: a link or a script can list a status a thousand times. One "?" per
    status typed failed past SQLite's variable limit (999 on older builds), and a status that doesn't exist just
    matched nothing without a word (B-126)."""
    assert ids(store, status=",".join(["review"] * 1200)) == ids(store, status="review")
    with pytest.raises(d.FilterError):
        d.condition("status", "reveiw")
