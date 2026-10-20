"""Sweep 4 (real-world testing): the findings Sweep 3 left, fixed."""

import datetime as dt
import io
import json

import pymupdf
from openpyxl import load_workbook

from ap_coder.core.store import Store, load_sample_setup
from ap_coder.erp import accruals, spend, stamp

from .conftest import SAMPLES

NORTHWIND = "northwind_ON_HST_NW-2026-0912"


def _approved_by_an_older_version(tmp_path):
    """An approved invoice whose final coding has no stored gl_distribution (approved by an older version)."""
    store = Store(tmp_path / "s.db")
    load_sample_setup(store)
    doc = json.loads((SAMPLES / "ground_truth" / f"{NORTHWIND}.json").read_text())
    assert "gl_distribution" not in doc
    invoice_id = store.add_invoice(SAMPLES / f"{NORTHWIND}.pdf", doc, {})
    store.approve_invoice(invoice_id, doc, "Jordan")
    assert not (store.get_invoice(invoice_id)["final_output"] or {}).get("gl_distribution")
    return store, invoice_id, doc


def test_an_invoice_approved_without_a_stored_distribution_is_in_the_spend_report(tmp_path):
    """B-217: spend.lines and the full data workbook build the posting lines again, as the exports do (B-151)."""
    store, invoice_id, doc = _approved_by_an_older_version(tmp_path)
    rows = spend.lines(store, dt.date(2020, 1, 1), dt.date(2030, 12, 31))
    assert {r.invoice_id for r in rows} == {invoice_id}
    assert round(sum(r.amount for r in rows), 2) == doc["subtotal"]  # the expense lines, net of the HST
    posting = load_workbook(io.BytesIO(spend.workbook(store, store.reference_data())))["Posting"]
    assert posting.max_row == len(doc["line_items"]) + 2  # the header, each line and the HST
    assert "2310" in {c.value for c in posting["F"]}  # the HST's own account (the store's tax setup)


def test_an_invoice_approved_without_a_stored_distribution_is_in_the_accrual_detail(tmp_path):
    """B-217: an approved invoice not yet exported is accrued by GL account even with no stored distribution."""
    store, _, doc = _approved_by_an_older_version(tmp_path)
    items = [a for a in accruals.build(store, dt.date(2026, 12, 31)) if a.source == accruals.NOT_IN_ERP]
    assert items and round(sum(a.amount for a in items), 2) == doc["subtotal"]
    assert all(a.gl_code for a in items)


def test_an_invoice_approved_without_a_stored_distribution_has_its_posting_on_the_stamp(tmp_path):
    """B-217: the stamped PDF's coding page lists the posting lines built again, with the tax setup's accounts."""
    store, invoice_id, doc = _approved_by_an_older_version(tmp_path)
    pdf = stamp.stamped_pdf(store.get_invoice(invoice_id), {}, store.tax_treatments())
    with pymupdf.open("pdf", pdf) as opened:
        coding = opened[opened.page_count - 1].get_text()
    assert f"{doc['grand_total']:,.2f}" in coding and "2310" in coding
