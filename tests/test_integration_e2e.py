"""End to end, offline: a folder of made-up emails from the AP mailbox goes in; a checked JD Edwards batch and its dry
run come out.

The emails are what Power Automate or a person dragging from Outlook would leave in the invoices folder (.eml and
.msg, an invoice forwarded inside an email, a statement, a bank-detail change, a quote). The folder watcher sorts and
reads them; a person approves the invoices; the JDE team's worksheet sets the supplier numbers; the dry run says what
the batch will load; the batch is made and downloaded with its files and DRY_RUN.txt."""

from __future__ import annotations

import csv
import io
import json
import os
import time
import zipfile
from email.message import EmailMessage

import openpyxl

from ap_coder.bench.msgwriter import Attachment, Message, build_msg
from ap_coder.coding.tax import build_gl_distribution
from ap_coder.core.schema import InvoiceCoding
from ap_coder.core.store import APPROVED, REVIEW, Store, load_sample_setup
from ap_coder.erp import jde, jde_worksheet
from ap_coder.intake import inbox, mailbox, triage
from ap_coder.ops import cli
from ap_coder.web.core import invoice_dir

from .conftest import SAMPLES
from .test_msg_and_triage import QUOTE, STATEMENT
from .webclient import client_for, post

STEMS = {"NW-2026-0912": "northwind_ON_HST_NW-2026-0912", "SIL-4471": "laurentides_QC_TPS_TVQ_SIL-4471",
         "PO-77120": "pacific_BC_GST_PST_PO-77120", "CCO-26-10418": "chinook_AB_GST_CCO-26-10418"}  # fmt: skip
INVOICES = ["northwind_ON_HST_NW-2026-0912", "laurentides_QC_TPS_TVQ_SIL-4471", "pacific_BC_GST_PST_PO-77120",
            "chinook_AB_GST_CCO-26-10418"]  # fmt: skip


def _eml(subject: str, body: str, attachments: list[tuple[str, bytes]] = (), forwarded: EmailMessage | None = None):
    m = EmailMessage()
    m["Subject"], m["From"], m["To"] = subject, "ar@vendor.example", "ap@pcl-made-up.example"
    m.set_content(body)
    for name, content in attachments:
        m.add_attachment(content, maintype="application", subtype="pdf", filename=name)
    if forwarded is not None:
        m.add_attachment(forwarded)
    return m.as_bytes()


def _mailbox(folder) -> None:
    pdf = {stem: (SAMPLES / f"{stem}.pdf").read_bytes() for stem in INVOICES}
    inner = EmailMessage()
    inner["Subject"] = "Facture SIL-4471"
    inner.set_content("Veuillez trouver ci-joint notre facture.")
    inner.add_attachment(pdf[INVOICES[1]], maintype="application", subtype="pdf", filename="SIL-4471.pdf")
    emails = {
        "01 invoice.eml": _eml("Invoice NW-2026-0912", "Attached.", [("NW-2026-0912.pdf", pdf[INVOICES[0]])]),
        "02 forwarded.eml": _eml("Fwd: facture", "See the forwarded email.", forwarded=inner),
        "03 from outlook.msg": build_msg(Message(subject="Invoice PO-77120", sender_email="ar@pacific.example",
                                                 body="Invoice attached.", attachments=[
                                                     Attachment("PO-77120.pdf", pdf[INVOICES[2]]),
                                                     Attachment("image001.png", b"\x89PNG" + b"0" * 3000, "image/png",
                                                                "logo")])),  # fmt: skip
        "04 past due.eml": _eml("PAST DUE: invoice CCO-26-10418", "Copy attached.", [("CCO.pdf", pdf[INVOICES[3]])]),
        "05 statement.eml": _eml("Statement of account", "Attached.", [("statement.pdf", STATEMENT)]),
        "06 bank.msg": build_msg(Message(subject="Updated bank details", body="Please update your records with our "
                                         "new banking information for all future payments.")),  # fmt: skip
        "07 quote.eml": _eml("Quote for your project", "See attached.", [("quote.pdf", QUOTE)]),
    }  # fmt: skip
    folder.mkdir(parents=True, exist_ok=True)
    old = time.time() - 600  # finished copying a while ago
    for name, data in emails.items():
        (folder / name).write_bytes(data)
        os.utime(folder / name, (old, old))


def test_a_folder_of_emails_becomes_a_checked_jde_batch(tmp_path, monkeypatch):
    db = tmp_path / "private" / "ap_coder.db"
    store = Store(db)
    load_sample_setup(store)
    folder = invoice_dir(store)
    _mailbox(folder)
    monkeypatch.setattr(cli, "SETTLE_SECONDS", 0)  # the attachments just saved are read in the same check
    args = [
        "--env-file",
        str(tmp_path / "none.env"),
        "--db",
        str(db),
        "watch",
        str(folder),
        "--once",
        "--cache-dir",
        "",
    ]

    # 1. Sorted and read: the four invoices (one forwarded, one in a .msg, one with a reminder) reach the queue.
    assert cli.main(args) == 0
    queue = store.list_invoices(status=REVIEW)
    assert sorted(i["invoice_number"] for i in queue) == ["CCO-26-10418", "NW-2026-0912", "PO-77120", "SIL-4471"]
    assert len(list((folder / mailbox.STATEMENTS_FOLDER).iterdir())) == 1
    set_aside = inbox.needs_person(store, folder)
    assert [r["kind"] for r in set_aside] == [triage.BANK_CHANGE, triage.NOT_FOR_AP]  # possible fraud first
    assert any(e["action"] == inbox.FRAUD_ACTION for e in store.events())
    assert cli.main(args) == 0 and len(store.list_invoices()) == 4  # a second check reads nothing twice

    # 2. A person approves each invoice, coding first any line AP Coder left for them (nothing learned yet).
    tax = store.reference_data().tax
    left_for_person = 0
    for inv in queue:
        full = store.get_invoice(inv["id"])
        doc = dict(full["final_output"] or full["ai_output"])
        truth = json.loads((SAMPLES / "ground_truth" / f"{STEMS[doc['invoice_number']]}.json").read_text())
        for line, right in zip(doc["line_items"], truth["line_items"], strict=True):
            for key in ("predicted_gl_code", "predicted_cost_center"):
                if str(line.get(key) or "UNASSIGNED").upper() == "UNASSIGNED":
                    line[key] = right[key]
                    left_for_person += 1
        coding = InvoiceCoding.model_validate({k: v for k, v in doc.items() if k != "gl_distribution"})
        doc["gl_distribution"] = build_gl_distribution(coding, tax)
        store.approve_invoice(inv["id"], doc, "tester")
    assert {i["status"] for i in store.list_invoices()} == {APPROVED}
    assert left_for_person < 20  # most lines arrive coded

    # 3. Before the JDE team answers, the dry run holds every invoice back and says why.
    client = client_for(store)
    before = client.get("/exports/jde/dry-run.txt").text
    assert "Will load: 0 voucher(s)" in before and before.count("HELD BACK") == 4 and "AN8" in before

    # 4. The JDE team fills in the worksheet (supplier numbers, a company) and AP imports it.
    sheet = client.get("/settings/jde/worksheet.xlsx")
    assert sheet.status_code == 200
    wb = openpyxl.load_workbook(io.BytesIO(sheet.content))
    for n, row in enumerate(wb["Suppliers"].iter_rows(min_row=2), start=10001):
        row[2].value = str(n)
    questions = wb["Questions"]
    company = next(r for r in questions.iter_rows(min_row=2) if r[0].value == "company")
    company[5].value = "00042"
    buf = io.BytesIO()
    wb.save(buf)
    answer = post(client, "/settings/jde/worksheet", files={"file": ("answers.xlsx", buf.getvalue())})
    assert answer.status_code == 200 and answer.json()["ok"], answer.text
    assert jde.load(store).company == "00042" and len(jde.load(store).an8_overrides) == 4

    # 5. The dry run now shows what the batch will load, line by line, and that each voucher balances.
    dry = client.get("/exports/jde/dry-run.txt").text
    assert "Will load: 4 voucher(s)" in dry and "HELD BACK" not in dry and dry.count("distribution ") == 4

    # 6. The batch: its files hold exactly what the dry run said.
    ids = [i["id"] for i in store.list_invoices()]
    made = post(client, "/exports/create", {"fmt": jde.FORMAT, "ids": [str(i) for i in ids]})
    assert made.status_code == 200 and made.json()["ok"], made.text
    (batch,) = [b["id"] for b in store.export_batches()]
    zipped = client.get(f"/exports/batches/{batch}/download")
    with zipfile.ZipFile(io.BytesIO(zipped.content)) as zf:
        assert sorted(zf.namelist()) == ["DRY_RUN.txt", "F0411Z1.csv", "F0911Z1.csv", "README.txt"]
        headers = list(csv.DictReader(io.StringIO(zf.read("F0411Z1.csv").decode())))
        lines = list(csv.DictReader(io.StringIO(zf.read("F0911Z1.csv").decode())))
        in_zip = zf.read("DRY_RUN.txt").decode()
    assert len(headers) == 4 and {h["VLCO"] for h in headers} == {"00042"}
    assert sorted(h["VLAN8"] for h in headers) == ["10001", "10002", "10003", "10004"]
    assert f"Rows: F0411Z1.csv 4, F0911Z1.csv {len(lines)}" in in_zip and "Will load: 4 voucher(s)" in in_zip
    assert jde_worksheet.QUESTIONS  # the worksheet the JDE team answered is the one AP Coder reads
