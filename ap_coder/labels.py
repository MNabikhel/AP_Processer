"""Ground-truth labelling via a spreadsheet the AP team can correct in Excel.

``export_labels`` turns a folder of engine outputs into a workbook with one row
per line item, pre-filled with the model's answers (or blank, in ``blind``
mode). An AP clerk corrects anything wrong and marks each row reviewed;
``load_labels`` reads it back as ground truth for ``evaluate``. Every cell
is stored as text, so Excel cannot strip leading zeros from GL codes or
reformat dates.
"""

from __future__ import annotations

import csv
import datetime as dt
import json
from pathlib import Path
from typing import Any

from .reference_data import UNASSIGNED, ReferenceData
from .safe import neutralise_sheet, parse_amount

SHEET = "Labels"
HEADER_FIELDS = (
    "vendor_name",
    "invoice_number",
    "invoice_date",
    "currency",
    "subtotal",
    "tax_total",
    "grand_total",
)
LINE_FIELDS = ("line_number", "description", "quantity", "unit_price", "amount", "gl_code", "cost_center")
COLUMNS = (
    "reviewed",
    "document",
    *HEADER_FIELDS,
    *LINE_FIELDS,
    "model_justification",
    "model_flagged_for_review",
    "notes",
)
_NUMERIC = {"subtotal", "tax_total", "grand_total", "quantity", "unit_price", "amount"}
_TRUTHY = {"y", "yes", "true", "1", "x", "ok", "✓"}

INSTRUCTIONS = [
    "HOW TO REVIEW",
    "1. Open each invoice next to this sheet. Each row is one invoice line.",
    "2. Invoice-level fields (vendor, number, date, totals) appear on the FIRST row of each invoice only.",
    "   Correct them there; blank header cells on later rows are expected.",
    "3. Correct any wrong value in place. gl_code and cost_center have dropdowns of valid codes.",
    "4. A line is missing? Insert a row under the invoice, fill line_number, description, amounts and codes.",
    "   The document column may be left blank on inserted rows (it is copied from the row above).",
    "5. A line should not exist (e.g. a page subtotal)? Delete the row.",
    "6. Set 'reviewed' to Y on every row once it is correct. Invoices with any row not marked Y are skipped.",
    "7. Dates must stay YYYY-MM-DD. Amounts are plain numbers without currency symbols.",
    "8. Save as .xlsx (keep the file name) and run: python -m ap_coder evaluate --ground-truth <this file>",
    "",
    "Yellow rows were flagged by the engine for review. model_justification shows the engine's reasoning.",
    "This file contains your enterprise data: keep it out of Git and do not share it externally.",
]


def _prediction_files(output_dir: Path) -> list[Path]:
    return sorted(
        p
        for p in output_dir.glob("*.json")
        if not p.name.endswith((".validation.json", ".di.json")) and p.name != "batch_summary.json"
    )


def _fmt_num(value: Any) -> str:
    if value is None or value == "":
        return ""
    return f"{float(value):.2f}" if not float(value).is_integer() else f"{float(value):.0f}"


def _rows_for_document(stem: str, doc: dict[str, Any], blind: bool, flagged: bool) -> list[dict[str, str]]:
    header = {f: (_fmt_num(doc.get(f)) if f in _NUMERIC else str(doc.get(f, ""))) for f in HEADER_FIELDS}
    lines = doc.get("line_items") or [{}]
    rows = []
    for i, li in enumerate(lines):
        row = {c: "" for c in COLUMNS}
        row["document"] = stem
        if i == 0:
            row.update(header)
        if li:
            row["line_number"] = str(li.get("line_number", ""))
            row["description"] = str(li.get("description", ""))
            row["quantity"] = _fmt_num(li.get("quantity"))
            row["unit_price"] = _fmt_num(li.get("unit_price"))
            row["amount"] = _fmt_num(li.get("amount"))
            if not blind:
                row["gl_code"] = str(li.get("predicted_gl_code", ""))
                row["cost_center"] = str(li.get("predicted_cost_center", ""))
                row["model_justification"] = str(li.get("reasoning_justification", ""))
        row["model_flagged_for_review"] = "YES" if flagged else ""
        rows.append(row)
    return rows


def export_labels(
    output_dir: str | Path,
    dest: str | Path,
    reference: ReferenceData | None = None,
    blind: bool = False,
) -> int:
    """Write the labelling workbook (.xlsx) or CSV. Returns the number of invoices exported."""
    output_dir, dest = Path(output_dir), Path(dest)
    rows: list[dict[str, str]] = []
    files = _prediction_files(output_dir)
    for path in files:
        doc = json.loads(path.read_text(encoding="utf-8"))
        meta_path = output_dir / f"{path.stem}.validation.json"
        flagged = False
        if meta_path.exists():
            meta = json.loads(meta_path.read_text(encoding="utf-8"))
            flagged = bool((meta.get("validation") or {}).get("requires_review"))
        rows += _rows_for_document(path.stem, doc, blind, flagged)

    dest.parent.mkdir(parents=True, exist_ok=True)
    if dest.suffix.lower() == ".csv":
        with dest.open("w", newline="", encoding="utf-8-sig") as fh:
            writer = csv.DictWriter(fh, fieldnames=COLUMNS)
            writer.writeheader()
            writer.writerows(rows)
    elif dest.suffix.lower() == ".xlsx":
        _write_xlsx(dest, rows, reference)
    else:
        raise ValueError("labels file must end in .xlsx or .csv")
    return len(files)


def _write_xlsx(dest: Path, rows: list[dict[str, str]], reference: ReferenceData | None) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation

    wb = Workbook()
    ws = wb.active
    ws.title = SHEET
    ws.append(list(COLUMNS))
    for cell in ws[1]:
        cell.font = Font(bold=True)
    flagged_fill = PatternFill("solid", fgColor="FFF4CC")
    for row in rows:
        ws.append([row[c] for c in COLUMNS])
        if row["model_flagged_for_review"]:
            for cell in ws[ws.max_row]:
                cell.fill = flagged_fill
    for col in ws.iter_cols(min_row=2, max_row=max(ws.max_row, 2)):
        for cell in col:
            cell.number_format = "@"  # text: keeps leading zeros and ISO dates intact
    neutralise_sheet(ws)
    widths = {"description": 50, "model_justification": 60, "vendor_name": 28, "notes": 30, "document": 30}
    for idx, name in enumerate(COLUMNS, start=1):
        ws.column_dimensions[ws.cell(1, idx).column_letter].width = widths.get(name, 14)
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = ws.dimensions

    if reference is not None:
        last_row = max(ws.max_row + 500, 1000)  # room for inserted rows
        for sheet_name, table, column in (
            ("GL Accounts", reference.chart_of_accounts, "gl_code"),
            ("Cost Centers", reference.cost_centers, "cost_center"),
        ):
            if table is None:
                continue  # cost centers are optional
            ref_ws = wb.create_sheet(sheet_name)
            name_col = next((c for c in table.rows[0] if c != table.key_column), None)
            ref_ws.append([table.key_column, name_col or ""])
            for r in [*table.rows, {table.key_column: UNASSIGNED, name_col: "No suitable code"}]:
                ref_ws.append([r[table.key_column], r.get(name_col, "") if name_col else ""])
            for cell in ref_ws["A"]:
                cell.number_format = "@"
            ref_ws.column_dimensions["A"].width = 16
            ref_ws.column_dimensions["B"].width = 45
            dv = DataValidation(
                type="list",
                formula1=f"='{sheet_name}'!$A$2:$A${ref_ws.max_row}",
                allow_blank=True,
                errorStyle="warning",
                showErrorMessage=True,
                error="Not a code in the reference list",
            )
            ws.add_data_validation(dv)
            letter = ws.cell(1, COLUMNS.index(column) + 1).column_letter
            dv.add(f"{letter}2:{letter}{last_row}")

    info = wb.create_sheet("Instructions")
    for line in INSTRUCTIONS:
        info.append([line])
    info.column_dimensions["A"].width = 110
    wb.save(dest)


def _read_rows(path: Path) -> list[tuple[int, dict[str, Any]]]:
    """Return ``(sheet_row_number, row_dict)`` pairs."""
    if path.suffix.lower() == ".csv":
        for encoding in ("utf-8-sig", "cp1252"):
            try:
                with path.open(newline="", encoding=encoding) as fh:
                    return [(i, row) for i, row in enumerate(csv.DictReader(fh), start=2)]
            except UnicodeDecodeError:
                continue
        raise ValueError(f"{path}: could not decode CSV (save it as UTF-8)")

    from openpyxl import load_workbook

    wb = load_workbook(path, data_only=True, read_only=True)
    ws = wb[SHEET] if SHEET in wb.sheetnames else wb.worksheets[0]
    it = ws.iter_rows(values_only=True)
    header = [str(h).strip() if h is not None else "" for h in next(it)]
    rows = [(i, dict(zip(header, values, strict=False))) for i, values in enumerate(it, start=2)]
    wb.close()
    return rows


def _text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, dt.datetime):
        return value.date().isoformat()
    if isinstance(value, dt.date):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _number(value: Any, where: str) -> float:
    number = parse_amount(value)  # "(250.00)" is a credit, "150,00" a decimal comma
    if number is None:
        raise ValueError(f"{where}: expected a number, got {value!r}")
    return number


def load_labels(path: str | Path) -> tuple[dict[str, dict[str, Any]], list[str]]:
    """Read a reviewed labels file. Returns ``({stem: target-schema dict}, [unreviewed stems])``."""
    path = Path(path)
    docs: dict[str, dict[str, Any]] = {}
    unreviewed: set[str] = set()
    current = ""
    for row_no, raw in _read_rows(path):
        row = {k: _text(v) for k, v in raw.items() if k}
        if not any(row.values()):
            continue
        current = row.get("document") or current
        if not current:
            raise ValueError(f"{path}: row {row_no} has no document")
        where = f"{path.name} row {row_no}"

        if current not in docs:
            header = {f: row.get(f, "") for f in HEADER_FIELDS}
            docs[current] = {**header, "_where": where, "confidence_score": 1.0, "line_items": []}
        if row.get("reviewed", "").lower() not in _TRUTHY:
            unreviewed.add(current)
        if row.get("line_number"):
            docs[current]["line_items"].append((where, row))

    reviewed = {stem: _finalise(doc) for stem, doc in docs.items() if stem not in unreviewed}
    return reviewed, sorted(unreviewed)


def _finalise(doc: dict[str, Any]) -> dict[str, Any]:
    """Convert a reviewed document's text cells to target-schema types."""
    where = doc.pop("_where")
    for f in ("subtotal", "tax_total", "grand_total"):
        doc[f] = _number(doc[f], f"{where} {f}")
    lines = []
    for line_where, row in doc["line_items"]:
        lines.append(
            {
                "line_number": int(_number(row["line_number"], f"{line_where} line_number")),
                "description": row.get("description", ""),
                "quantity": _number(row.get("quantity") or 1, f"{line_where} quantity"),
                "unit_price": _number(row.get("unit_price") or row.get("amount"), f"{line_where} unit_price"),
                "amount": _number(row.get("amount"), f"{line_where} amount"),
                "predicted_gl_code": row.get("gl_code", ""),
                "predicted_cost_center": row.get("cost_center", ""),
                "reasoning_justification": row.get("notes", ""),
            }
        )
    doc["line_items"] = lines
    return doc
