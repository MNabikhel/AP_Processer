"""ERP export files for approved invoices.

An export batch is a set of approved invoices handed to the ERP together. Each invoice is exported
once (the store remembers its batch), and a batch's file can be downloaded again at any time: it is
rebuilt from the approved data: the same invoices every time, with today's settings and master data (vendor IDs,
GL names, exchange rates).

Formats
-------
* ``xlsx``: a workbook with an *Invoices* sheet (one row per invoice), a *GL lines* sheet (the posting:
  one row per expense line and per recoverable-tax account) and a *By GL account* summary.
* ``csv``: one row per GL line, the shape most ERP AP-invoice imports accept.
* ``custom``: a CSV with the columns of your ERP's import (``export_layout``).

The two CSV files are for the ERP's loader, like the JD Edwards files: their text is written exactly as it
is, with no spreadsheet apostrophe (an invoice number ``'+A-100`` would miss the ERP's duplicate check).
The xlsx workbook is opened by people, so its formula-looking text is stored as plain text.
* ``jde``: JD Edwards E1 Z-table files (F0411Z1 / F0911Z1), built by ``jde.build_zip``.
"""

from __future__ import annotations

import csv
import datetime as dt
import io
from collections import defaultdict
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from ..coding.tax import PLACE_OF_SUPPLY, TaxRateTable, TaxSetup, TaxTreatment, build_gl_distribution
from ..coding.vendors import master_vendor_id
from ..core.safe import neutralise_sheet, xlsx_row

FORMATS = {
    "xlsx": "Excel workbook: invoices, GL lines and a GL summary",
    "csv": "CSV: one row per GL line (for ERP import)",
    "custom": "Custom CSV: your ERP's columns (set up below)",
    "jde": "JD Edwards E1 (F0411Z1 / F0911Z1): ZIP of Z-table files (Settings → JD Edwards E1)",
}

INVOICE_COLUMNS = [
    ("batch", "Batch"), ("invoice_id", "AP Coder #"), ("vendor_id", "Vendor ID"), ("vendor_name", "Vendor"),
    ("gst_hst_registration_number", "Vendor GST/HST #"), ("qst_registration_number", "Vendor QST #"),
    ("invoice_number", "Invoice #"), ("invoice_date", "Invoice date"), ("po_number", "PO #"),
    ("payment_terms", "Terms"), ("due_date", "Due date"),
    ("currency", "Currency"), ("subtotal", "Subtotal"), ("tax_total", "Tax"), ("grand_total", "Total"),
    ("supplier_province", "Supplier province"), ("ship_to_province", "Place of supply"),
    ("reviewer", "Approved by"), ("reviewed_at", "Approved at"), ("second_reviewer", "Second approver"),
    ("file_name", "File"), ("pay_from", "Pay from"), ("pay_from_gl", "Pay from GL"),
]  # fmt: skip
LINE_COLUMNS = [
    ("batch", "Batch"), ("invoice_id", "AP Coder #"), ("vendor_id", "Vendor ID"), ("vendor_name", "Vendor"),
    ("invoice_number", "Invoice #"),
    ("invoice_date", "Invoice date"), ("currency", "Currency"), ("line", "Line"), ("kind", "Type"),
    ("gl_code", "GL account"), ("gl_name", "GL name"), ("cost_center", "Cost center"),
    ("description", "Description"), ("tax_code", "Tax code"), ("net_amount", "Net"),
    ("non_recoverable_tax", "Non-recoverable tax"), ("amount", "Amount"), ("pay_from", "Pay from"),
    ("pay_from_gl", "Pay from GL"),
]  # fmt: skip
MONEY_HEADERS = {"Subtotal", "Tax", "Total", "Net", "Non-recoverable tax", "Amount"}
# Real dates in the workbook (Excel sorts and filters them as dates), not text (B-470); the CSV keeps YYYY-MM-DD.
DATE_HEADERS = {"Invoice date": "yyyy-mm-dd", "Due date": "yyyy-mm-dd", "Approved at": "yyyy-mm-dd hh:mm"}


def _excel_date(value: Any) -> Any:
    """'2026-10-01' -> date, '2026-10-01 14:05:09' -> datetime (seconds kept); anything else as it is."""
    if not isinstance(value, str) or not value.strip():
        return value
    text = value.strip().replace("T", " ")
    try:
        return dt.datetime.fromisoformat(text).replace(tzinfo=None) if " " in text else dt.date.fromisoformat(text)
    except ValueError:
        return value


def _vendor_id(final: dict[str, Any], vendor_ids: dict[str, str]) -> str:
    """The ERP vendor ID: by vendor name, else by GST/HST number (``vendors.master_vendor_id``)."""
    return master_vendor_id(vendor_ids, final.get("vendor_name"), final.get("gst_hst_registration_number"))


def pay_from(final: dict[str, Any], bank_accounts: dict[str, dict[str, Any]] | None) -> dict[str, str]:
    """{"pay_from": the bank account's name, "pay_from_gl": its cash GL account} for the account an invoice is paid
    from (D-66), as the accounts are today (a batch is rebuilt with today's master data, D-29); blank when none."""
    code = str(final.get("pay_from_account") or "")
    account = (bank_accounts or {}).get(code) if code else None
    if account is None:
        return {"pay_from": code, "pay_from_gl": ""}
    return {"pay_from": str(account.get("name") or code), "pay_from_gl": str(account.get("gl_account") or "")}


def invoice_rows(
    invoices: list[dict[str, Any]], batch: int | str = "", vendor_ids: dict[str, str] | None = None,
    bank_accounts: dict[str, dict[str, Any]] | None = None,
) -> list[dict[str, Any]]:  # fmt: skip
    """``invoices``: store.get_invoice() dicts of approved invoices; ``vendor_ids``: ``Store.vendor_ids()``;
    ``bank_accounts``: our bank accounts by code (``Store.bank_accounts``), for Pay from."""
    vendor_ids = vendor_ids or {}
    rows = []
    for inv in invoices:
        final = inv.get("final_output") or {}
        rows.append(
            {
                **pay_from(final, bank_accounts),
                **{k: final.get(k) for k, _ in INVOICE_COLUMNS if k in final},
                "batch": batch,
                "invoice_id": inv["id"],
                "vendor_id": _vendor_id(final, vendor_ids),
                # the ship-to printed, else the place of supply worked out when it was coded (pipeline)
                "ship_to_province": final.get("ship_to_province") or final.get(PLACE_OF_SUPPLY) or "",
                # printed, else worked out from the terms / vendor master / default when it was approved
                "due_date": final.get("due_date") or inv.get("due_date") or "",
                "reviewer": inv.get("reviewer"),
                "reviewed_at": (inv.get("reviewed_at") or "").replace("T", " "),
                "second_reviewer": inv.get("second_reviewer") or "",
                "file_name": inv.get("file_name"),
            }  # fmt: skip
        )
    return rows


def distribution(final: dict[str, Any], tax_treatments: dict[str, TaxTreatment] | None = None) -> list[dict[str, Any]]:
    """An approved invoice's posting lines: as stored when it was coded, else (an invoice approved by an older
    version) built again from its coding, with ``tax_treatments`` (``Store.tax_treatments()``; without them the
    default treatments, and a recoverable tax line has no GL account)."""
    dist = final.get("gl_distribution")
    if dist:
        return list(dist)
    from ..core.schema import InvoiceCoding

    try:
        coding = InvoiceCoding.model_validate({k: v for k, v in final.items() if k != "gl_distribution"})
    except ValueError:
        return []
    return build_gl_distribution(coding, TaxSetup(TaxRateTable(()), dict(tax_treatments or {})))


def line_rows(
    invoices: list[dict[str, Any]], gl_names: dict[str, str] | None = None, batch: int | str = "",
    vendor_ids: dict[str, str] | None = None, tax_treatments: dict[str, TaxTreatment] | None = None,
    bank_accounts: dict[str, dict[str, Any]] | None = None,
) -> list[dict[str, Any]]:  # fmt: skip
    gl_names, vendor_ids = gl_names or {}, vendor_ids or {}
    rows = []
    for inv in invoices:
        final = inv.get("final_output") or {}
        paid_from = pay_from(final, bank_accounts)
        dist = distribution(final, tax_treatments)
        codes = line_tax_codes(dist, final.get("line_items") or [])
        for e, code in zip(dist, codes, strict=True):
            rows.append(
                {
                    **paid_from,
                    "batch": batch,
                    "invoice_id": inv["id"],
                    "vendor_id": _vendor_id(final, vendor_ids),
                    "vendor_name": final.get("vendor_name"),
                    "invoice_number": final.get("invoice_number"),
                    "invoice_date": final.get("invoice_date"),
                    "currency": final.get("currency"),
                    "line": e.get("line_number") or "",
                    "kind": "Expense"
                    if e.get("kind") == "expense"
                    else ("Non-recoverable tax" if e.get("non_recoverable_tax") else "Recoverable tax"),
                    "gl_code": e.get("gl_code"),
                    "gl_name": gl_names.get(e.get("gl_code") or "", ""),
                    "cost_center": e.get("cost_center") or "",
                    "description": e.get("description") or "",
                    "tax_code": code,
                    "net_amount": e.get("net_amount"),
                    "non_recoverable_tax": e.get("non_recoverable_tax"),
                    "amount": e.get("amount"),
                }  # fmt: skip
            )
    return rows


def line_tax_codes(dist: list[dict[str, Any]], items: list[dict[str, Any]]) -> list[str]:
    """The tax code of the invoice line behind each posting line ("GST+PST", "Exempt", D-64), "" for a tax line: the
    expense lines follow the invoice lines in order (``build_gl_distribution``); matched by line number when the
    counts differ."""
    from ..coding.line_tax import export_code

    expenses = [e for e in dist if e.get("kind") == "expense"]
    if len(expenses) == len(items):
        it = iter(items)
        return [export_code(next(it)) if e.get("kind") == "expense" else "" for e in dist]
    by_number: dict[Any, str] = {}
    for li in items:
        by_number.setdefault(li.get("line_number"), export_code(li))
    return [by_number.get(e.get("line_number"), "") if e.get("kind") == "expense" else "" for e in dist]


def gl_summary(lines: list[dict[str, Any]]) -> list[dict[str, Any]]:
    totals: dict[tuple[str, str, str], float] = defaultdict(float)
    for r in lines:
        totals[(r["gl_code"] or "", r["gl_name"] or "", r["currency"] or "")] += float(r["amount"] or 0)
    return [
        {"gl_code": gl, "gl_name": name, "currency": cur, "amount": round(amount, 2)}
        for (gl, name, cur), amount in sorted(totals.items())
    ]


def build_xlsx(
    invoices: list[dict[str, Any]], gl_names: dict[str, str] | None = None, batch: int | str = "",
    vendor_ids: dict[str, str] | None = None, tax_treatments: dict[str, TaxTreatment] | None = None,
    bank_accounts: dict[str, dict[str, Any]] | None = None,
) -> bytes:  # fmt: skip
    lines = line_rows(invoices, gl_names, batch, vendor_ids, tax_treatments, bank_accounts)
    wb = Workbook()
    _sheet(wb.active, "Invoices", INVOICE_COLUMNS, invoice_rows(invoices, batch, vendor_ids, bank_accounts))
    _sheet(wb.create_sheet(), "GL lines", LINE_COLUMNS, lines)
    _sheet(
        wb.create_sheet(), "By GL account",
        [("gl_code", "GL account"), ("gl_name", "GL name"), ("currency", "Currency"), ("amount", "Amount")],
        gl_summary(lines),
    )  # fmt: skip
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _sheet(ws: Any, title: str, columns: list[tuple[str, str]], rows: list[dict[str, Any]]) -> None:
    ws.title = title
    ws.append([label for _, label in columns])
    dated = {key for key, label in columns if label in DATE_HEADERS}
    for r in rows:
        ws.append(xlsx_row([_excel_date(r.get(key)) if key in dated else r.get(key) for key, _ in columns]))
    header_fill = PatternFill("solid", fgColor="E9EEF6")
    for cell in ws[1]:
        cell.font, cell.fill, cell.alignment = Font(bold=True), header_fill, Alignment(vertical="center")
    for i, (_, label) in enumerate(columns, start=1):
        width = max([len(label)] + [len(str(r.get(columns[i - 1][0]) or "")) for r in rows[:500]])
        ws.column_dimensions[get_column_letter(i)].width = min(max(10, width + 2), 60)
        if label in MONEY_HEADERS:
            for row in ws.iter_rows(min_row=2, min_col=i, max_col=i):
                row[0].number_format = "#,##0.00"
        if label in DATE_HEADERS:
            for row in ws.iter_rows(min_row=2, min_col=i, max_col=i):
                if isinstance(row[0].value, (dt.date, dt.datetime)):
                    row[0].number_format = DATE_HEADERS[label]
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    neutralise_sheet(ws)  # vendor names and descriptions come from documents: never formulas


def build_csv(
    invoices: list[dict[str, Any]], gl_names: dict[str, str] | None = None, batch: int | str = "",
    vendor_ids: dict[str, str] | None = None, tax_treatments: dict[str, TaxTreatment] | None = None,
    bank_accounts: dict[str, dict[str, Any]] | None = None,
) -> bytes:  # fmt: skip
    buf = io.StringIO()
    writer = csv.writer(buf)
    writer.writerow([label for _, label in LINE_COLUMNS])
    for r in line_rows(invoices, gl_names, batch, vendor_ids, tax_treatments, bank_accounts):
        # for the ERP's loader: text as it is, no spreadsheet apostrophe (see the module docstring)
        writer.writerow([r.get(key) if r.get(key) is not None else "" for key, _ in LINE_COLUMNS])
    return buf.getvalue().encode("utf-8-sig")  # BOM: Excel shows accents correctly


def custom_rows(
    invoices: list[dict[str, Any]], gl_names: dict[str, str] | None = None, batch: int | str = "",
    vendor_ids: dict[str, str] | None = None, tax_treatments: dict[str, TaxTreatment] | None = None,
    bank_accounts: dict[str, dict[str, Any]] | None = None,
) -> list[dict[str, Any]]:  # fmt: skip
    """Posting lines, each with its invoice's fields too (for a custom layout)."""
    heads = {r["invoice_id"]: r for r in invoice_rows(invoices, batch, vendor_ids, bank_accounts)}
    lines = line_rows(invoices, gl_names, batch, vendor_ids, tax_treatments, bank_accounts)
    return [{**heads.get(r["invoice_id"], {}), **r} for r in lines]


def build(
    fmt: str, invoices: list[dict[str, Any]], gl_names: dict[str, str] | None = None, batch: int | str = "",
    vendor_ids: dict[str, str] | None = None, layout: Any = None,
    tax_treatments: dict[str, TaxTreatment] | None = None, bank_accounts: dict[str, dict[str, Any]] | None = None,
) -> tuple[bytes, str, str]:  # fmt: skip
    """(file bytes, file name, mime type). ``tax_treatments``: ``Store.tax_treatments()``, for the posting lines of
    an invoice approved without them stored (``distribution``); ``bank_accounts``: our bank accounts by code, for the
    Pay from columns (D-66)."""
    name = f"ap_coder_export_{batch}" if batch != "" else "ap_coder_export"
    if fmt == "jde":  # needs its settings and checks: see ap_coder.erp.jde.build_zip
        raise ValueError("JD Edwards files are built with ap_coder.erp.jde.build_zip")
    if fmt == "csv":
        return (
            build_csv(invoices, gl_names, batch, vendor_ids, tax_treatments, bank_accounts),
            f"{name}.csv",
            "text/csv",
        )
    if fmt == "custom":
        from . import export_layout

        rows = custom_rows(invoices, gl_names, batch, vendor_ids, tax_treatments, bank_accounts)
        return (
            export_layout.build(rows, layout or export_layout.Layout(export_layout.default_columns())),
            (f"{name}_custom.csv"),
            "text/csv",
        )
    mime = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    return build_xlsx(invoices, gl_names, batch, vendor_ids, tax_treatments, bank_accounts), f"{name}.xlsx", mime
