"""Spend analysis from approved invoices, and a full data download for Excel / Power BI.

Spend is the cost to the company: each expense line's amount plus any sales tax that is not recovered
(recoverable GST/HST/QST is left out, as it is claimed back). Credit notes reduce it. Amounts stay in
each invoice's currency; currencies are never added together.
"""

from __future__ import annotations

import datetime as dt
import io
import json
from collections import defaultdict
from dataclasses import dataclass
from typing import Any

from ..coding.tax import PLACE_OF_SUPPLY
from ..core.reference_data import ReferenceData
from ..core.safe import neutralise_sheet, xlsx_row
from ..core.store import APPROVED, FAILED, Store
from .exports import distribution


@dataclass
class SpendLine:
    invoice_id: int
    month: str  # YYYY-MM of the invoice date
    vendor: str
    gl_code: str
    cost_center: str
    amount: float
    currency: str
    # A row of ``parts`` sums one invoice's lines on one GL account and cost center (in the order of the first of
    # them): ``sizes`` sums the sizes of their amounts (``currencies``), ``cents`` their amounts each rounded to the
    # cent and ``amounts`` lists them, for ``in_cad`` (None: a single line, ``amount``).
    sizes: float | None = None
    cents: float | None = None
    amounts: tuple[float, ...] | None = None


def spend_lines(doc: dict[str, Any], treatments: dict[str, Any]) -> list[tuple[str, str, float]]:
    """(GL account, cost center, amount) of each cost line of an approved invoice's coding ``doc``: its expense
    lines, and non-recoverable tax posted to its own GL account (a cost too)."""
    out = []
    for e in distribution(doc, treatments):  # an invoice approved by an older version: built again
        if e.get("kind") != "expense" and not e.get("non_recoverable_tax"):
            continue
        out.append((str(e.get("gl_code") or ""), str(e.get("cost_center") or ""), float(e.get("amount") or 0)))
    return out


def spend_parts(doc: dict[str, Any], treatments: dict[str, Any]) -> list[tuple[Any, ...]]:
    """``spend_lines`` by GL account and cost center, in the order of the first line of each, as the summary keeps
    them (``summary_spend``): (GL account, cost center, amount, sizes, cents, amounts as JSON or None). The amounts
    are summed line by line as the report sums them; each line's amount is listed for an invoice not in CAD, which
    ``in_cad`` converts line by line."""
    parts: dict[tuple[str, str], list[float]] = {}
    for gl, cc, amount in spend_lines(doc, treatments):
        parts.setdefault((gl, cc), []).append(amount)
    listed = str(doc.get("currency") or "CAD") != "CAD"
    out = []
    for (gl, cc), amounts in parts.items():
        total = sizes = cents = 0.0
        for a in amounts:
            total += a
            sizes += abs(a)
            cents += round(a, 2)
        out.append((gl, cc, total, sizes, cents, json.dumps(amounts) if listed else None))
    return out


def parts(store: Store, start: dt.date, end: dt.date) -> list[SpendLine]:
    """The cost lines of approved invoices dated from ``start`` to ``end`` (inclusive), by invoice, GL account and
    cost center (``spend_parts``), from the invoices' summary (D-110): what the Spend page totals. ``total_by``,
    ``by_month_and_category``, ``currencies`` and ``in_cad`` give what they give for ``lines``."""
    first, last = start.isoformat(), end.isoformat()
    rows = store.summary_query(
        "SELECT s.id, substr(s.f_date, 1, 7), s.p_vendor, s.p_currency, p.gl_code, p.cost_center, p.amount, p.sizes, "
        "p.cents, p.amounts FROM invoice_summary s JOIN summary_spend p ON p.invoice_id = s.id "
        "WHERE s.f_date BETWEEN ? AND ? AND +s.status = ? AND s.broken = 0 ORDER BY s.id, p.part",
        (first, last, APPROVED),
    )
    out = [
        SpendLine(r[0], r[1], r[2], r[4], r[5], r[6], r[3], r[7], r[8], tuple(json.loads(r[9])) if r[9] else None)
        for r in rows
    ]
    out += [line for line in _broken_lines(store, first, last)]
    return out


def _broken_lines(store: Store, first: str, last: str) -> list[SpendLine]:
    """``lines`` of the approved invoices the summary could not read (their JSON): read again from the invoice."""
    ids = [r[0] for r in store.summary_rows("broken = 1 AND +status = ?", (APPROVED,), "id")]
    return [line for line in _lines_of(store.invoice_columns(("id", "final_output"), ids=ids), store, first, last)
            ] if ids else []  # fmt: skip


def lines(store: Store, start: dt.date, end: dt.date) -> list[SpendLine]:
    """Expense postings of approved invoices dated from ``start`` to ``end`` (inclusive)."""
    return _lines_of(store.invoice_columns(("id", "final_output"), APPROVED), store, start.isoformat(),
                     end.isoformat())  # fmt: skip


def _lines_of(rows: list[dict[str, Any]], store: Store, first: str, last: str) -> list[SpendLine]:
    treatments = store.tax_treatments()
    out = []
    for r in rows:
        doc = r["final_output"] or {}
        date = str(doc.get("invoice_date") or "")
        if not (first <= date <= last):
            continue
        vendor, currency = str(doc.get("vendor_name") or ""), str(doc.get("currency") or "CAD")
        out += [SpendLine(r["id"], date[:7], vendor, gl, cc, amount, currency)
                for gl, cc, amount in spend_lines(doc, treatments)]  # fmt: skip
    return out


def currencies(rows: list[SpendLine]) -> list[str]:
    """Currencies by total spend, largest first."""
    totals: dict[str, float] = defaultdict(float)
    for r in rows:
        totals[r.currency] += abs(r.amount) if r.sizes is None else r.sizes
    return sorted(totals, key=lambda c: -totals[c])


def in_cad(rows: list[SpendLine], rates: dict[str, float]) -> list[SpendLine]:
    """The lines in CAD at ``rates`` (CAD per unit); lines in a currency without a rate are left out."""
    return [
        SpendLine(r.invoice_id, r.month, r.vendor, r.gl_code, r.cost_center, _in_cad(r, rates[r.currency]), "CAD")
        for r in rows
        if r.currency in rates
    ]


def _in_cad(r: SpendLine, rate: float) -> float:
    """``r``'s amount in CAD, line by line rounded to the cent (a row of ``parts``: each of its lines)."""
    if r.amounts is not None:
        return sum(round(a * rate, 2) for a in r.amounts)
    if r.cents is not None and rate == 1.0:  # a * 1.0 is a: each line rounded to the cent, summed
        return r.cents
    return round(r.amount * rate, 2)


def total_by(rows: list[SpendLine], key: str, currency: str) -> list[tuple[str, float, int]]:
    """(value of ``key``, total, invoices) for one currency, largest total first."""
    totals: dict[str, float] = defaultdict(float)
    invoices: dict[str, set[int]] = defaultdict(set)
    for r in rows:
        if r.currency == currency:
            value = getattr(r, key)
            totals[value] += r.amount
            invoices[value].add(r.invoice_id)
    return sorted(((k, round(v, 2), len(invoices[k])) for k, v in totals.items()), key=lambda x: -x[1])


def by_month_and_category(rows: list[SpendLine], currency: str, category: dict[str, str]) -> list[dict[str, Any]]:
    """[{month, category, amount}] for a stacked chart."""
    totals: dict[tuple[str, str], float] = defaultdict(float)
    for r in rows:
        if r.currency == currency:
            totals[(r.month, category.get(r.gl_code) or "Uncategorised")] += r.amount
    return [{"month": m, "category": c, "amount": round(v, 2)} for (m, c), v in sorted(totals.items())]


def default_period(today: dt.date | None = None) -> tuple[dt.date, dt.date]:
    """The last 12 months, this one included."""
    today = today or dt.date.today()
    start = dt.date(today.year - 1, today.month, 1)
    start = dt.date(start.year + (start.month == 12), start.month % 12 + 1, 1)
    return start, today


def accounts_by_code(reference: ReferenceData | None, cost_centers: bool = False) -> dict[str, dict[str, str]]:
    table = (reference.cost_centers if cost_centers else reference.chart_of_accounts) if reference else None
    return {row[table.key_column]: row for row in table.rows} if table else {}


# --- Full data download -------------------------------------------------------------------------------

INVOICE_COLUMNS = [
    ("id", "AP Coder #"), ("status", "Status"), ("vendor_name", "Vendor"), ("invoice_number", "Invoice #"),
    ("invoice_date", "Invoice date"), ("due_date", "Due date"), ("po_number", "PO #"),
    ("payment_terms", "Terms"), ("currency", "Currency"), ("subtotal", "Subtotal"), ("tax_total", "Tax"),
    ("grand_total", "Total"), ("supplier_province", "Supplier province"), ("ship_to_province", "Place of supply"),
    ("gst_hst_registration_number", "GST/HST #"), ("qst_registration_number", "QST #"),
    ("reviewer", "Approved by"), ("reviewed_at", "Approved at"), ("second_reviewer", "Second approval by"),
    ("second_reviewed_at", "Second approval at"), ("export_batch", "Export batch"), ("created_at", "Processed at"),
]  # fmt: skip


def workbook(store: Store, reference: ReferenceData | None) -> bytes:
    """Every invoice (not the unreadable ones), its lines and its posting, as an Excel workbook with three
    tables that join on "AP Coder #"."""
    from openpyxl import Workbook
    from openpyxl.worksheet.table import Table, TableStyleInfo

    accounts = accounts_by_code(reference)

    def gl_name(code: str) -> str:
        return accounts.get(code, {}).get("description", "")

    def category(code: str) -> str:
        return accounts.get(code, {}).get("category", "")

    columns = ("id", "status", "reviewer", "reviewed_at", "second_reviewer", "second_reviewed_at", "export_batch",
               "created_at", "due_date", "ai_output", "final_output")  # fmt: skip
    rows = [r for r in store.invoice_columns(columns) if r["status"] != FAILED]
    treatments = store.tax_treatments()
    wb = Workbook()
    sheets = {
        "Invoices": [label for _, label in INVOICE_COLUMNS],
        "Lines": ["AP Coder #", "Vendor", "Invoice #", "Invoice date", "Line", "Description", "Quantity",
                  "Unit price", "Amount", "GL account", "GL name", "GL category", "Cost center", "Currency"],
        "Posting": ["AP Coder #", "Vendor", "Invoice #", "Invoice date", "Kind", "GL account", "GL name",
                    "Cost center", "Description", "Net", "Non-recoverable tax", "Amount", "Currency"],
    }  # fmt: skip
    ws_by_name = {}
    for n, (name, header) in enumerate(sheets.items()):
        ws = wb.active if n == 0 else wb.create_sheet()
        ws.title = name
        ws.append(header)
        ws_by_name[name] = ws
    for r in rows:
        doc = r["final_output"] or r["ai_output"] or {}
        merged = {**doc, **{k: r[k] for k in r if k not in ("ai_output", "final_output")}}
        # the ship-to printed, else the place of supply worked out when it was coded (pipeline)
        merged["ship_to_province"] = doc.get("ship_to_province") or doc.get(PLACE_OF_SUPPLY) or ""
        ws_by_name["Invoices"].append(xlsx_row([merged.get(k) for k, _ in INVOICE_COLUMNS]))
        head = [r["id"], doc.get("vendor_name"), doc.get("invoice_number"), doc.get("invoice_date")]
        for li in doc.get("line_items") or []:
            gl = li.get("predicted_gl_code") or ""
            ws_by_name["Lines"].append(
                xlsx_row(
                    [
                        *head,
                        li.get("line_number"),
                        li.get("description"),
                        li.get("quantity"),
                        li.get("unit_price"),
                        li.get("amount"),
                        gl,
                        gl_name(gl),
                        category(gl),
                        li.get("predicted_cost_center"),
                        doc.get("currency"),
                    ]
                )
            )
        for e in distribution(doc, treatments):
            gl = e.get("gl_code") or ""
            ws_by_name["Posting"].append(
                xlsx_row(
                    [
                        *head,
                        e.get("kind"),
                        gl,
                        gl_name(gl),
                        e.get("cost_center"),
                        e.get("description"),
                        e.get("net_amount"),
                        e.get("non_recoverable_tax"),
                        e.get("amount"),
                        doc.get("currency"),
                    ]
                )
            )
    for name, ws in ws_by_name.items():
        neutralise_sheet(ws)
        last = ws.cell(row=1, column=ws.max_column).column_letter
        if ws.max_row > 1:  # an Excel table (filters, and a named source for Power BI)
            table = Table(displayName=f"{name}Table", ref=f"A1:{last}{ws.max_row}")
            table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
            ws.add_table(table)
        for col in ws.columns:
            width = max(len(str(c.value or "")) for c in list(col)[:200])
            ws.column_dimensions[col[0].column_letter].width = min(max(width + 2, 9), 48)
        ws.freeze_panes = "A2"
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
