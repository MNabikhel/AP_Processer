"""The JD Edwards settings worksheet: one Excel file AP sends the JDE team, and AP Coder reads back.

``build(settings, gl_accounts, vendors)`` writes a workbook with every question AP Coder needs answered before the
first live load (the open questions of ``docs/JDE_E1.md``), what AP Coder uses now, and a column for the answer:

* **Questions**: company, document types, batch and user, amounts, line numbering, dates, PO path, currencies, and
  the questions that are not settings (the R04110ZA version, its duplicate edit, how the AAIs post tax ...), whose
  answers are kept with the settings so nothing the JDE team wrote is lost;
* **Tax areas**: province → tax area (VLTXA1) and explanation code (VLEXR1);
* **Account mapping**: each AP Coder GL account (and cost centre) → business unit, object, subsidiary, subledger;
* **Suppliers**: each vendor → its JDE address number (AN8).

``read(content, settings)`` applies the answers to a copy of the settings and says what changed and what could not
be used (a value of the wrong form is reported, never guessed). A blank answer keeps what AP Coder has.
"""

from __future__ import annotations

import io
import re
from dataclasses import dataclass, field
from typing import Any

from . import jde
from .memory import vendor_key
from .safe import neutralise_sheet
from .tax import OUTSIDE_CANADA, PROVINCES, province_label

TEXT, BOOL, CHOICE, LIST, EXTRAS, INFO = "text", "bool", "choice", "list", "extras", "info"


@dataclass(frozen=True)
class Question:
    id: str  # a JdeSettings field, or an "info:" question that is not a setting
    topic: str
    question: str
    jde_field: str = ""
    kind: str = TEXT
    choices: dict[str, str] = field(default_factory=dict)
    width: int = 0  # most characters allowed (0: no limit)


QUESTIONS = [
    Question("company", "Company", "Which company number do AP Coder's vouchers belong to?", "VLCO / VNCO", width=5),
    Question(
        "header_bu", "Company", "Which business unit goes on the voucher header (blank if none)?", "VLMCU", width=12
    ),
    Question(
        "default_bu", "Company", "Default business unit for G/L lines with no cost centre mapped?", "VNMCU", width=12
    ),
    Question(
        "bu_from_cost_center",
        "Company",
        "Is a line's business unit its AP Coder cost centre? (yes / no)",
        "VNMCU",
        BOOL,
    ),
    Question(
        "po_company",
        "Company",
        "Company of the purchase orders, if different (blank: the company above)",
        "VLPKCO",
        width=5,
    ),
    Question(
        "document_type",
        "Document types",
        "Document type for imported vouchers (PV, or a custom one)?",
        "VLDCT",
        width=2,
    ),
    Question(
        "credit_document_type",
        "Document types",
        "Document type for credit notes (blank: the type above, with negative amounts)?",
        "VLDCT",
        width=2,
    ),
    Question(
        "info:p0400047",
        "Document types",
        "Which P0400047 version does R04110ZA run with? (Its processing options set defaults and edits.)",
        "R04110ZA",
        INFO,
    ),
    Question(
        "info:tax_recalc",
        "Document types",
        "Does that version recalculate tax, or take the tax amounts AP Coder sends?",
        "P0400047",
        INFO,
    ),
    Question(
        "info:duplicate_edit",
        "Document types",
        "Duplicate-invoice edit in the processing options: H (hard error) or I (warning)?",
        "P0400047",
        INFO,
    ),
    Question(
        "batch_prefix",
        "Batch",
        "Prefix for batch numbers (the batch is prefix + AP Coder's batch number)?",
        "VLEDBT",
        width=9,
    ),
    Question("user", "Batch", "User ID the batches are loaded and R04110ZA run for?", "VLEDUS", width=10),
    Question(
        "info:batch_size", "Batch", "Should a batch be limited in size (how many vouchers at most)?", "VLEDBT", INFO
    ),
    Question(
        "amount_mode",
        "Amounts",
        "Send the gross amount only (E1 calculates the tax), or the taxable / tax amounts printed on the invoice?",
        "VLAG or VLATXA / VLSTAM",
        CHOICE,
        jde.AMOUNT_MODES,
    ),
    Question(
        "decimals",
        "Amounts",
        "Are amounts loaded with implied decimals (5.50 as 550) or a decimal point?",
        "all amounts",
        CHOICE,
        jde.DECIMAL_MODES,
    ),
    Question(
        "line_numbering",
        "Amounts",
        "How are G/L lines numbered under the pay item?",
        "VNEDLN",
        CHOICE,
        jde.LINE_NUMBERING,
    ),
    Question(
        "dist_tax_columns",
        "Amounts",
        "Should the G/L lines carry the tax area and code too? (yes / no)",
        "VNTXA1 / VNEXR1",
        BOOL,
    ),
    Question(
        "account_mode",
        "Accounts",
        "Is the account sent as one field (BU.Object.Subsidiary) or as separate columns?",
        "VNANI or VNMCU / VNOBJ / VNSUB",
        CHOICE,
        jde.ACCOUNT_MODES,
    ),
    Question(
        "use_default_rule",
        "Accounts",
        "For a GL account not in the Account mapping sheet, use BU = cost centre and object = GL code? (yes / no)",
        "VNANI",
        BOOL,
    ),
    Question(
        "info:bu_justified",
        "Accounts",
        "Must business units be right-justified in 12 characters for your load method?",
        "VNMCU",
        INFO,
    ),
    Question(
        "info:aai",
        "Tax",
        "How do the tax AAIs (PT / GT) post GST, HST, QST and PST, in particular for codes C and B?",
        "AAIs",
        INFO,
    ),
    Question(
        "self_assessed_code",
        "Tax",
        "Explanation code for self-assessed PST (supplier charged GST only)?",
        "VLEXR1",
        width=2,
    ),
    Question("exempt_code", "Tax", "Explanation code when no tax is charged?", "VLEXR1", width=2),
    Question("gl_date", "Dates", "Which date is the G/L date?", "VLDGJ", CHOICE, jde.GL_DATES),
    Question(
        "send_due_date",
        "Dates",
        "Send the due date printed on the invoice? (no: E1 calculates it from the terms)",
        "VLDDJ",
        BOOL,
    ),
    Question(
        "payment_terms_code",
        "Payment",
        "Payment terms code to send (blank: E1 takes the supplier's terms)?",
        "VLPTC",
        width=3,
    ),
    Question("pay_status", "Payment", "Pay status of new vouchers (e.g. A approved, H held)?", "VLPST", width=1),
    Question(
        "po_matched",
        "Purchase orders",
        "Do PO-matched vouchers come in through the Z-tables? (yes / no)",
        "F0411Z1",
        BOOL,
    ),
    Question(
        "send_po_reference",
        "Purchase orders",
        "Send the PO number on vouchers that are not matched? (yes / no)",
        "VLPO",
        BOOL,
    ),
    Question("po_document_type", "Purchase orders", "Purchase order document type?", "VLPDCT", width=2),
    Question(
        "po_extra_columns",
        "Purchase orders",
        "Extra columns and values for PO-matched pay items (e.g. VLATFLG=1, VLPSTE=)",
        "F0411Z1",
        EXTRAS,
    ),
    Question(
        "info:po_path",
        "Purchase orders",
        "How are PO invoices matched to receipts: by hand in E1, a custom program reading a match file, "
        "or an orchestration?",
        "P4314",
        INFO,
    ),
    Question("domestic_currency", "Currency", "Domestic currency?", "VLCRCD", width=3),
    Question("currencies", "Currency", "Currencies AP Coder may send (comma-separated)?", "VLCRCD", LIST),
    Question(
        "info:exchange_rates",
        "Currency",
        "Exchange rates: from E1's table (VLCRR blank) or the rate AP Coder has?",
        "VLCRR",
        INFO,
    ),
]
BY_ID = {q.id: q for q in QUESTIONS}
UPPER = {"document_type", "credit_document_type", "po_document_type", "domestic_currency", "self_assessed_code",
         "exempt_code"}  # fmt: skip
YES, NO = {"yes", "y", "true", "1", "oui", "o", "x"}, {"no", "n", "false", "0", "non"}
SHEETS = ("Read me", "Questions", "Tax areas", "Account mapping", "Suppliers")


def _current(q: Question, settings: jde.JdeSettings) -> str:
    if q.kind == INFO:
        return settings.answers.get(q.id, "")
    value = getattr(settings, q.id)
    if q.kind == BOOL:
        return "yes" if value else "no"
    if q.kind == CHOICE:
        return q.choices.get(value, str(value))
    if q.kind == LIST:
        return ", ".join(value)
    if q.kind == EXTRAS:
        return ", ".join(f"{k}={v}" for k, v in value.items())
    return str(value)


def build(settings: jde.JdeSettings, gl_accounts: list[dict[str, Any]], vendors: list[dict[str, Any]],
          an8_of: Any = None) -> bytes:  # fmt: skip
    """The worksheet as .xlsx bytes. ``gl_accounts``: [{"code", "description"}]; ``vendors``: [{"vendor_name"}];
    ``an8_of(name)``: the address number AP Coder would send for a vendor now ("" if none)."""
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation

    wb = openpyxl.Workbook()
    bold, head_fill = Font(bold=True), PatternFill("solid", fgColor="DCE6F2")
    answer_fill = PatternFill("solid", fgColor="FFF2CC")
    wrap = Alignment(wrap_text=True, vertical="top")

    def table(ws: Any, widths: list[int], answer_cols: tuple[int, ...]) -> None:
        """Style a sheet whose first row is its header: widths, a bold header, the answer columns in yellow."""
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w
            cell = ws.cell(row=1, column=i)
            cell.font, cell.fill = bold, head_fill
        ws.freeze_panes = "A2"
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = wrap
                if cell.column in answer_cols:
                    cell.fill = answer_fill

    readme = wb.active
    readme.title = "Read me"
    for line in [
        "AP Coder: questions for the JD Edwards team",
        "",
        "AP Coder writes JD Edwards E1 voucher batches (F0411Z1 / F0911Z1 files for R04110ZA). Before the first load",
        "into a JDE test environment, it needs your answers to the questions in this workbook.",
        "",
        "How to fill it in: write your answer in the yellow column of each sheet. A blank answer keeps what AP Coder",
        "uses now (shown next to it). Do not change the ID column or the sheet names.",
        "",
        "Sheets: Questions (settings and questions about R04110ZA), Tax areas (per province), Account mapping (AP",
        "Coder GL account -> BU.Object.Subsidiary), Suppliers (JDE address number, AN8, for each vendor).",
        "",
        "Send it back to AP, who imports it in AP Coder: Settings -> JD Edwards E1 -> Import the JDE team's worksheet.",
        "AP Coder then lists what changed and anything it could not use. Every change is in its activity log.",
    ]:
        readme.append([line])
    readme["A1"].font = Font(bold=True, size=13)
    readme.column_dimensions["A"].width = 110

    ws = wb.create_sheet("Questions")
    ws.append(
        ["ID", "Topic", "Question for the JDE team", "JDE field", "AP Coder uses now", "Your answer", "Allowed answers"]
    )
    choice_dv: dict[str, DataValidation] = {}
    for q in QUESTIONS:
        allowed = {BOOL: "yes or no", CHOICE: " / ".join(q.choices.values()), LIST: "codes separated by commas",
                   EXTRAS: "COLUMN=value, separated by commas", INFO: "free text"}.get(q.kind, "")  # fmt: skip
        if q.kind == TEXT and q.width:
            allowed = f"up to {q.width} characters"
        ws.append([q.id, q.topic, q.question, q.jde_field, _current(q, settings), "", allowed])
        if q.kind in (BOOL, CHOICE):
            values = ["yes", "no"] if q.kind == BOOL else list(q.choices.values())
            formula = '"' + ",".join(v.replace('"', "'").replace(",", ";") for v in values) + '"'
            if len(formula) <= 255:  # Excel's limit for a list typed in the rule
                dv = choice_dv.setdefault(formula, DataValidation(type="list", formula1=formula, allow_blank=True))
                dv.add(f"F{ws.max_row}")
    for dv in choice_dv.values():
        ws.add_data_validation(dv)
    table(ws, [18, 14, 60, 18, 34, 34, 34], (6,))

    tax = wb.create_sheet("Tax areas")
    tax.append(["Province", "Name", "Tax area now", "Code now", "Your tax area (VLTXA1)", "Your code (VLEXR1)"])
    for p in [*PROVINCES, OUTSIDE_CANADA]:
        area = settings.tax_map.get(p) or jde.TaxArea("", "")
        tax.append([p, province_label(p), area.tax_area, area.code, "", ""])
    table(tax, [10, 26, 16, 10, 24, 18], (5, 6))

    acc = wb.create_sheet("Account mapping")
    acc.append(["GL code", "Cost centre", "Description", "BU", "Object", "Subsidiary", "Subledger", "Subledger type"])
    mapped = {(m.gl_code, m.cost_center): m for m in settings.gl_map}
    rows = {(m.gl_code, m.cost_center): "" for m in settings.gl_map}
    rows.update({(str(a.get("code") or ""), ""): str(a.get("description") or "") for a in gl_accounts if a.get("code")})
    for (gl, cc), desc in sorted(rows.items()):
        m = mapped.get((gl, cc)) or jde.GlMap(gl, cc)
        acc.append([gl, cc, desc, m.bu, m.obj, m.sub, m.sbl, m.sblt])
    table(acc, [12, 12, 40, 12, 10, 12, 12, 14], (4, 5, 6, 7, 8))

    sup = wb.create_sheet("Suppliers")
    sup.append(["Vendor", "AN8 AP Coder sends now", "Your AN8"])
    for v in sorted(vendors, key=lambda r: str(r.get("vendor_name") or "")):
        name = str(v.get("vendor_name") or "")
        if name:
            sup.append([name, (an8_of(name) if an8_of else "") or "", ""])
    table(sup, [44, 22, 16], (3,))
    for sheet in wb.worksheets:
        neutralise_sheet(sheet)  # a vendor name or GL description read from a document is text, never a formula
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


@dataclass
class Result:
    settings: jde.JdeSettings
    changes: list[str] = field(default_factory=list)
    problems: list[str] = field(default_factory=list)


def _rows(wb: Any, name: str) -> list[dict[str, str]]:
    if name not in wb.sheetnames:
        return []
    rows = list(wb[name].iter_rows(values_only=True))
    if not rows:
        return []
    header = [str(h or "").strip() for h in rows[0]]
    return [{h: "" if c is None else str(c).strip() for h, c in zip(header, r, strict=False)} for r in rows[1:]]


def _choice(q: Question, answer: str) -> str | None:
    """The choice an answer names: its code, its label (as the drop-down list writes it: commas as semicolons),
    or the start of its label."""
    folded = answer.strip().lower().replace(";", ",").replace("'", '"')
    for key, label in q.choices.items():
        full = label.lower()
        if folded in (key.lower(), full) or (len(folded) >= 4 and full.startswith(folded)):
            return key
    return None


def read(content: bytes, settings: jde.JdeSettings) -> Result:
    """The settings with the worksheet's answers applied (a copy), what changed and what could not be used."""
    import openpyxl

    out = Result(jde.JdeSettings.from_json(settings.to_json()))
    s = out.settings
    try:
        wb = openpyxl.load_workbook(io.BytesIO(content), data_only=True, read_only=True)
    except Exception as exc:  # noqa: BLE001 - not an Excel file, or damaged
        out.problems.append(f"The file could not be opened as an Excel workbook ({type(exc).__name__}).")
        return out
    if "Questions" not in wb.sheetnames:
        out.problems.append("This is not AP Coder's JDE worksheet (it has no Questions sheet).")
        return out
    for row in _rows(wb, "Questions"):
        q, answer = BY_ID.get(row.get("ID", "")), row.get("Your answer", "")
        if q is None or not answer:
            continue
        label = f'"{q.question[:60]}"'
        if q.kind == INFO:
            if s.answers.get(q.id) != answer:
                s.answers[q.id] = answer
                out.changes.append(f"{q.topic}: {q.question} → {answer}")
            continue
        value: Any
        if q.kind == BOOL:
            if answer.lower() not in YES | NO:
                out.problems.append(f"{label}: answer yes or no, not “{answer}”.")
                continue
            value = answer.lower() in YES
        elif q.kind == CHOICE:
            value = _choice(q, answer)
            if value is None:
                out.problems.append(f"{label}: “{answer}” is not one of: {' / '.join(q.choices.values())}.")
                continue
        elif q.kind == LIST:
            value = [c.strip().upper() for c in re.split(r"[,;\s]+", answer) if c.strip()]
        elif q.kind == EXTRAS:
            value = {}
            for part in re.split(r"[,;]", answer):
                if part.strip():
                    key, _, val = part.partition("=")
                    value[key.strip().upper()] = val.strip()
        else:
            value = answer.upper() if q.id in UPPER else answer
            if q.width and len(value) > q.width:
                out.problems.append(f"{label}: “{answer}” is longer than {q.width} characters.")
                continue
        if getattr(s, q.id) != value:
            setattr(s, q.id, value)
            out.changes.append(f"{q.topic}: {q.question} → {_current(q, s)}")
    for row in _rows(wb, "Tax areas"):
        prov = row.get("Province", "")
        area, code = row.get("Your tax area (VLTXA1)", ""), row.get("Your code (VLEXR1)", "").upper()
        if prov not in s.tax_map or not (area or code):
            continue
        if code and code not in jde.RECOVERABLE_BY_CODE:
            out.problems.append(f"Tax areas, {prov}: code “{code}” is not one AP Coder knows (V, C, B or E).")
            continue
        new = jde.TaxArea(area or s.tax_map[prov].tax_area, code or s.tax_map[prov].code)
        if new != s.tax_map[prov]:
            s.tax_map[prov] = new
            out.changes.append(f"Tax area for {prov}: {new.tax_area}, code {new.code}")
    filled = [
        r for r in _rows(wb, "Account mapping") if any(r.get(k) for k in ("BU", "Object", "Subsidiary", "Subledger"))
    ]
    if filled:
        new_map = jde.gl_map_from_records(filled)
        by_key = {(m.gl_code, m.cost_center): m for m in s.gl_map}
        for m in new_map:
            if by_key.get((m.gl_code, m.cost_center)) != m:
                by_key[(m.gl_code, m.cost_center)] = m
                out.changes.append(
                    f"Account {m.gl_code}{' / ' + m.cost_center if m.cost_center else ''} → {jde.ani(m)}"
                )
        s.gl_map = list(by_key.values())
    for row in _rows(wb, "Suppliers"):
        name, an8 = row.get("Vendor", ""), row.get("Your AN8", "")
        if not (name and an8):
            continue
        if not an8.isdigit() or len(an8) > 8:
            out.problems.append(f"Suppliers, {name}: the address number “{an8}” should be up to 8 digits.")
            continue
        key = vendor_key(name)
        if s.an8_overrides.get(key) != an8:
            s.an8_overrides[key] = an8
            out.changes.append(f"Supplier {name} → AN8 {an8}")
    out.problems += [f"After the answers: {p}" for p in jde.settings_problems(s)]
    return out
