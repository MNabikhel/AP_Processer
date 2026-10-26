"""The invoice directory: every invoice in one table, filtered column by column (``directory.py``).

``GET /search`` shows it: a filter box over each column (query by example), sort by a heading, 100 rows a page, the
columns chosen and ordered, saved layouts ("presets"), and the filtered rows exported (CSV or Excel). The quick search
(``q``: vendor, invoice or PO number, amount; ``search.find``) narrows it further and also looks in the ERP register.
An invoice in the queue opens, a parked one can be brought back, and an approved one can be downloaded with its
APPROVED stamp (``GET /search/<id>/approved.pdf``).
"""

from __future__ import annotations

from typing import Any

from fastapi import APIRouter, Form, HTTPException, Request
from fastapi.responses import HTMLResponse, JSONResponse, Response

from ap_coder.core.safe import csv_row
from ap_coder.core.store import APPROVED, FAILED, PARKED, PENDING, REJECTED, REVIEW
from ap_coder.erp import directory, search, stamp
from ap_coder.web.access import is_admin
from ap_coder.web.core import done, get_store, local_url, reference_or_none, refuse, render, reviewer, server_mode
from ap_coder.web.files import download
from ap_coder.web.identity import person
from ap_coder.web.pages.invoice import changed_meanwhile

router = APIRouter()

STATUS = {
    REVIEW: ("To review", "info", "inbox"),
    PARKED: ("Parked", "warn", "pause_circle"),
    PENDING: ("Second approval", "violet", "how_to_reg"),
    APPROVED: ("Approved", "ok", "task_alt"),
    REJECTED: ("Rejected", "err", "block"),
    FAILED: ("Unreadable", "err", "error"),
}


def hit_view(hit: search.Hit) -> dict[str, Any]:
    """One search result as the template draws it (plain values: the template escapes them)."""
    r = hit.row
    erp = hit.source == "ERP"
    if erp:
        pill = ("In the ERP register", "gray", "database")
        where = "Found in the ERP's invoice register (imported on the Exports page), not in AP Coder."
        status = ""
    else:
        status = r["status"]
        pill = STATUS.get(status, (status, "gray", "help"))
        where = search.where(r)
    show_due = r.get("due_date") and status not in (REJECTED, FAILED) and (r.get("grand_total") or 0) > 0
    return {
        "id": r.get("id"),
        "erp": erp,
        "status": status,
        "vendor": r.get("vendor_name") or "",
        "number": r.get("invoice_number") or "—",
        "date": r.get("invoice_date") or "no date",
        "due": r.get("due_date") if show_due else "",
        "total": r.get("grand_total") or 0,
        "currency": r.get("currency") or "",
        "pill": pill,
        "where": where,
    }


def _owner(request: Request) -> str:
    """Whose saved layouts a person changes: on the server, their own (the shared ones are an Admin's, D-109); on a
    laptop, and for an Admin, the shared ones ("")."""
    me = person()
    if not server_mode(request) or me is None or is_admin(request):
        return ""
    return f"user:{me.user_id}"


def _query(request: Request) -> directory.Query:
    """The directory's query: the page's own, else (nothing asked) the default saved layout."""
    params = request.query_params
    if not params:
        preset = directory.default_preset(get_store(request), _owner(request))
        if preset:
            from urllib.parse import parse_qsl

            return directory.Query.from_params(parse_qsl(preset["query"]))
    return directory.Query.from_params(params)


def row_view(row: dict[str, Any], columns: list[str]) -> dict[str, Any]:
    """One invoice of the directory as the template draws it (plain values: the template escapes them)."""
    status = row["status"]
    return {
        "id": row["id"],
        "status": status,
        "pill": STATUS.get(status, (status, "gray", "help")),
        "cells": {key: directory.cell(row, key) for key in columns},
        "currency": row.get("currency") or "",
    }


@router.get("/search", response_class=HTMLResponse)
def search_page(request: Request) -> HTMLResponse:
    store = get_store(request)
    query = _query(request)
    result = directory.run(store, query)
    chosen = [directory.COLUMNS[k] for k in query.columns]
    others = [c for k, c in directory.COLUMNS.items() if k not in query.columns]
    return render(
        request,
        "search.html",
        "search",
        query=query,
        result=result,
        columns=chosen,
        others=others,
        rows=[row_view(r, query.columns) for r in result.rows],
        erp=[hit_view(h) for h in result.erp],
        statuses=STATUS,
        presets=directory.shown_presets(store, _owner(request)),
        default_preset=(directory.default_preset(store, _owner(request)) or {}).get("name", ""),
        here=query.query_string(),
        filtered=bool(query.filters or query.q),
        kinds={"money": directory.MONEY, "date": directory.DATE, "number": directory.NUMBER},
        all_columns=directory.COLUMNS,
        default_columns=list(directory.DEFAULT_COLUMNS),
        export_limit=directory.EXPORT_LIMIT,
    )


def _export_rows(request: Request) -> tuple[list[tuple[str, str]], list[dict[str, Any]], directory.Result]:
    store = get_store(request)
    query = _query(request)
    result = directory.run(store, query, everything=True)
    columns = [(k, directory.COLUMNS[k].label) for k in query.columns]
    rows = [{k: directory.cell(r, k) for k, _ in columns} for r in result.rows]
    for row in rows:  # the status in plain words, as the page says it (B-1186)
        if row.get("status") in STATUS:
            row["status"] = STATUS[row["status"]][0]
    return columns, rows, result


@router.get("/search/export.csv")
def export_csv(request: Request) -> Response:
    """The filtered rows (all pages, up to 50,000), with the columns shown, as CSV."""
    import csv
    import io

    from ap_coder.erp.exports import csv_amount

    columns, rows, _ = _export_rows(request)
    buf = io.StringIO()
    writer = csv.writer(buf)
    writer.writerow([label for _, label in columns])
    for r in rows:
        writer.writerow(csv_row([csv_amount(r.get(k)) if directory.COLUMNS[k].kind == directory.MONEY else r.get(k)
                                 if r.get(k) is not None else "" for k, _ in columns]))  # fmt: skip
    return download(buf.getvalue().encode("utf-8-sig"), "invoice_directory.csv", "text/csv")


@router.get("/search/export.xlsx")
def export_xlsx(request: Request) -> Response:
    """The filtered rows (all pages, up to 50,000), with the columns shown, as an Excel workbook."""
    import io

    from openpyxl import Workbook

    from ap_coder.erp.exports import _sheet

    columns, rows, _ = _export_rows(request)
    wb = Workbook()
    _sheet(wb.active, "Invoices", columns, rows)
    buf = io.BytesIO()
    wb.save(buf)
    mime = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    return download(buf.getvalue(), "invoice_directory.xlsx", mime)


@router.post("/search/presets")
def save_preset(
    request: Request, name: str = Form(""), query: str = Form(""), make_default: str = Form("")
) -> JSONResponse:
    """Save the layout shown (its columns, filters and sort) by name."""
    store = get_store(request)
    try:
        kept = directory.save_preset(store, name, query, make_default=bool(make_default), owner=_owner(request))
    except ValueError as exc:
        return refuse(str(exc))
    store.log_event("directory_preset_saved", None, reviewer(), {"name": kept})
    preset = next(p for p in directory.presets(store, _owner(request)) if p["name"] == kept)
    return done(request, local_url(f"/search?{preset['query']}"), f"Layout “{kept}” saved.")


@router.post("/search/presets/default")
def default_preset(request: Request, name: str = Form("")) -> JSONResponse:
    """The saved layout the directory opens with ("" for the standard one)."""
    store = get_store(request)
    if not directory.set_default(store, name, _owner(request)):
        return refuse("That layout is not saved any more.", 404)
    text = f"The directory now opens with “{name}”." if name else "The directory now opens with the standard layout."
    return done(request, "/search", text)


@router.post("/search/presets/delete")
def delete_preset(request: Request, name: str = Form("")) -> JSONResponse:
    store = get_store(request)
    if not directory.delete_preset(store, name, _owner(request)):
        if _owner(request) and any(p["name"] == name for p in directory.presets(store)):
            return refuse("That layout is shared with everyone: an Admin deletes it.", 403)
        return refuse("That layout is not saved any more.", 404)
    store.log_event("directory_preset_deleted", None, reviewer(), {"name": name})
    return done(request, "/search", f"Layout “{name}” deleted.")


@router.get("/search/{invoice_id:int}/approved.pdf")
def approved_pdf(request: Request, invoice_id: int) -> Response:
    """The approved invoice with its APPROVED stamp and coding page (never before the second approval)."""
    store = get_store(request)
    inv = store.get_invoice(invoice_id)
    if inv is None:
        raise HTTPException(status_code=404, detail="That invoice is not in AP Coder.")
    if not stamp.can_stamp(inv):
        raise HTTPException(status_code=409, detail="No longer fully approved: no approved PDF.")
    reference = reference_or_none(store)
    names = (
        {row["gl_code"]: row.get("description", "") for row in reference.chart_of_accounts.rows} if reference else {}
    )
    pdf = stamp.stamped_pdf(inv, names, store.tax_treatments())
    return download(pdf, stamp.file_name(inv), "application/pdf")


@router.post("/search/{invoice_id:int}/unpark")
def bring_back(request: Request, invoice_id: int) -> JSONResponse:
    """A parked invoice back to the review queue, and open it."""
    store = get_store(request)
    inv = store.get_invoice(invoice_id)
    if inv is None:
        return refuse("That invoice is not in AP Coder.", 404)
    try:
        store.unpark_invoice(invoice_id, reviewer())
    except ValueError as exc:
        if (store.get_invoice(invoice_id) or {}).get("status") != REVIEW:  # rejected, approved, deleted meanwhile
            return changed_meanwhile(request, invoice_id, exc)
        # else someone else brought it back meanwhile: it is opened all the same
    return done(request, f"/review/{invoice_id}", "Back in the review queue.")
